# -*- coding: utf-8 -*-

# =============================================================================
# empirical_percentile_sensitivity_analysis.py
#
# Reproducible analysis for Taiwan GSAT (2015–2024) and AST Mathematics A
#
# Analysis purpose:
# Recalculate human-cohort percentile ranks directly from CEEC empirical
# grade-band counts and cumulative distributions, using an explicit
# within-band interpolation rule, and compare the empirical results with
# the original normal-approximation approach as a sensitivity analysis.
#
# 目的：
#   1. 讀取 6 科 AI 逐題原始得分
#   2. 加總成 Year × Subject × Model 的 AI score
#   3. 讀取 CEEC 官方 human grade-band distributions
#   4. 以 empirical cumulative counts + within-band linear interpolation
#      重新計算 percentile
#   5. 重新計算原 normal approximation percentile
#   6. Normal vs Empirical sensitivity analysis
#   7. QC
#   8. 輸出 Excel
#   9. 輸出 Word 摘要報告
#  10. 直接產出 manuscript-ready Table 8 / Table 9（Excel + Word）
#
# 本程式完全獨立執行：
# 不需要先執行任何 Jupyter Step
# =============================================================================


# =============================================================================
# Fixed score-scale definitions for this study
# =============================================================================
#
# IMPORTANT:
#   AI Accuracy and the human-percentile comparison scale are related but
#   conceptually distinct. All AI-human percentile comparisons must first
#   harmonize AI scores and CEEC human score distributions to the same scale.
#
# Chinese 2015–2017:
#   AI objectively scoreable maximum = 54
#   AI Accuracy = AI raw / 54 × 100
#   AI comparison score = AI raw / 54 × 100
#   CEEC human distribution uses the full 108-point scale
#   Human comparison boundary = official raw boundary / 108 × 100
#
# Chinese 2018–2024:
#   AI and Human are already on the 0–100 scale.
#
# English 2015–2024:
#   AI objectively scoreable maximum = 72
#   AI Accuracy = AI raw / 72 × 100
#   AI comparison score = AI raw / 72 × 100
#   CEEC human distribution is already on the 0–100 scale.
#
# GSAT Mathematics human baseline:
#   2015–2021 -> Math
#   2022–2024 -> MathA
#   MathB is outside the study definition.
#
# AST Mathematics A:
#   2015–2021 -> raw-score comparison
#   2022–2024 -> AI raw score converted to official CEEC 0–60 level;
#                Human uses official 0–60 level distribution.
#
# Primary percentile:
#   Official CEEC empirical cumulative counts with within-band interpolation.
#   Single-score bands use mid-rank (0.5).
#
# Normal approximation:
#   Sensitivity analysis only; not the primary percentile method.
# =============================================================================


# =============================================================================
# Step 0. Import
# =============================================================================

from pathlib import Path
import pandas as pd
import numpy as np
import math
import re
from datetime import datetime

# Excel manuscript-ready 表格格式
try:
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
except ImportError as exc:
    raise ImportError(
        "缺少 openpyxl。請先執行：\n"
        "pip install openpyxl"
    ) from exc

# Word 輸出需要 python-docx
try:
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt
except ImportError as exc:
    raise ImportError(
        "缺少 python-docx。請先執行：\n"
        "pip install python-docx"
    ) from exc


# =============================================================================
# Step 1. Project paths
# =============================================================================
#
# Repository structure:
#
# GSAT-Scoring/
# ├─ analysis/
# │  └─ empirical_percentile_sensitivity_analysis.py
# │
# ├─ supplementary-materials/
# │  └─ analysis-ready-data/
# │     ├─ GSAT_Chinese.xlsx
# │     ├─ GSAT_English.xlsx
# │     ├─ GSAT_Mathematics.xlsx
# │     ├─ GSAT_Science.xlsx
# │     ├─ GSAT_Social_Studies.xlsx
# │     ├─ MathA.xlsx
# │     ├─ Human Baseline_GSAT_2015_2024.xlsx
# │     └─ Human Baseline_MathA_2014_2024.xlsx
# │
# └─ output/
#
# Important:
#   No computer-specific absolute path is used.
#   The project root is resolved from this .py file itself.
# =============================================================================

SCRIPT_DIR = Path(__file__).resolve().parent
PROJECT_DIR = SCRIPT_DIR.parent

DATA_DIR = (
    PROJECT_DIR
    / "supplementary-materials"
    / "analysis-ready-data"
)

AI_DIR = DATA_DIR

OUTPUT_DIR = (
    PROJECT_DIR
    / "output"
)

OUTPUT_DIR.mkdir(
    parents=True,
    exist_ok=True
)


GSAT_HUMAN_FILE = (
    DATA_DIR
    / "Human Baseline_GSAT_2015_2024.xlsx"
)


MATHA_HUMAN_FILE = (
    DATA_DIR
    / "Human Baseline_MathA_2014_2024.xlsx"
)


OUTPUT_FILE = (
    OUTPUT_DIR
    / "empirical_percentile_sensitivity_analysis.xlsx"
)


WORD_OUTPUT_FILE = (
    OUTPUT_DIR
    / "empirical_percentile_sensitivity_report.docx"
)


# Manuscript-ready Table 8 + Table 9
MANUSCRIPT_TABLE_EXCEL = (
    OUTPUT_DIR
    / "table8_table9_manuscript_ready.xlsx"
)


MANUSCRIPT_TABLE_WORD = (
    OUTPUT_DIR
    / "table8_table9_manuscript_ready.docx"
)


print("=" * 90)
print("Project paths")
print("=" * 90)

print(
    "Project directory:",
    PROJECT_DIR
)

print(
    "Data directory   :",
    DATA_DIR
)

print(
    "Output directory :",
    OUTPUT_DIR
)


# =============================================================================
# Step 2. AI 檔案與 Sheet
# =============================================================================

AI_FILES = {

    "Chinese":
        AI_DIR / "GSAT_Chinese.xlsx",

    "English":
        AI_DIR / "GSAT_English.xlsx",

    "Mathematics":
        AI_DIR / "GSAT_Mathematics.xlsx",

    "Science":
        AI_DIR / "GSAT_Science.xlsx",

    "SocialStudies":
        AI_DIR / "GSAT_Social_Studies.xlsx",

    "MathematicsA_AST":
        AI_DIR / "MathA.xlsx",
}


AI_SHEETS = {

    "Chinese":
        "SPSS_Use",

    "English":
        "SPSS_Use",

    "Mathematics":
        "SPSS_Use",

    "Science":
        "SPSS_Use",

    "SocialStudies":
        "SPSS_Use",

    "MathematicsA_AST":
        "SPSS_Use_This",
}


# =============================================================================
# Step 3. AI Model 欄位
# =============================================================================

MODEL_COLUMNS = {

    "GPT-4o":
        "GPT4o_Score",

    "o1-preview":
        "o1_preview_Score",

    "o1-mini":
        "o1_mini_Score",

    "GPT-5.4 Thinking":
        "GPT54Thinking_Score",

    "Gemini 3 Flash Thinking":
        "Gemini3FlashThinking_Score",
}


# =============================================================================
# Step 4. 檢查所有輸入檔案
# =============================================================================

print("=" * 90)
print("Step 4. 檢查輸入檔案")
print("=" * 90)


required_files = [
    GSAT_HUMAN_FILE,
    MATHA_HUMAN_FILE,
    *AI_FILES.values()
]


missing_files = []


for file_path in required_files:

    if file_path.exists():

        print(
            "OK   :",
            file_path
        )

    else:

        print(
            "找不到:",
            file_path
        )

        missing_files.append(
            file_path
        )


if missing_files:

    raise FileNotFoundError(
        "\n有輸入檔案不存在，程式停止。"
    )


print(
    "\n所有輸入檔案均存在。"
)


# =============================================================================
# Step 5. 讀取 Human baseline
# =============================================================================

print("\n")
print("=" * 90)
print("Step 5. 讀取 Human baseline")
print("=" * 90)


# -----------------------------------------------------------------------------
# GSAT
# -----------------------------------------------------------------------------

human_gsat_raw = pd.read_excel(
    GSAT_HUMAN_FILE,
    sheet_name="All_Long"
)


# -----------------------------------------------------------------------------
# AST Mathematics A
# -----------------------------------------------------------------------------

human_matha_raw = pd.read_excel(
    MATHA_HUMAN_FILE,
    sheet_name="MathA_2015_2024"
)


print(
    "GSAT Human rows:",
    len(human_gsat_raw)
)

print(
    "AST Math A Human rows:",
    len(human_matha_raw)
)


# =============================================================================
# Step 6. 讀取 AI 原始逐題資料
# =============================================================================

print("\n")
print("=" * 90)
print("Step 6. 讀取 AI 原始逐題資料")
print("=" * 90)


ai_data = {}


for subject_key, file_path in AI_FILES.items():

    sheet_name = AI_SHEETS[
        subject_key
    ]

    print(
        f"{subject_key}"
        f" → {sheet_name}"
    )

    df = pd.read_excel(
        file_path,
        sheet_name=sheet_name
    )


    # -------------------------------------------------------------------------
    # 基本必要欄位
    # -------------------------------------------------------------------------

    required_columns = [
        "Year",
        *MODEL_COLUMNS.values()
    ]


    missing_columns = [
        column
        for column in required_columns
        if column not in df.columns
    ]


    if missing_columns:

        raise KeyError(
            f"\n{subject_key} 缺少欄位："
            f"{missing_columns}"
            f"\n目前欄位："
            f"{df.columns.tolist()}"
        )


    ai_data[
        subject_key
    ] = df


print(
    "\n6 個 AI 檔案讀取完成。"
)


# =============================================================================
# Step 7. AI 逐題分數 → 每年 × 每科 × 每模型總分
# =============================================================================

print("\n")
print("=" * 90)
print("Step 7. 計算 AI yearly raw score")
print("=" * 90)


ai_summary_list = []


reverse_model_map = {

    column: model

    for model, column
    in MODEL_COLUMNS.items()
}


for subject_key, df in ai_data.items():

    data = df.copy()


    # -------------------------------------------------------------------------
    # Year
    # -------------------------------------------------------------------------

    data["Year"] = pd.to_numeric(
        data["Year"],
        errors="coerce"
    )


    # -------------------------------------------------------------------------
    # Model scores
    # -------------------------------------------------------------------------

    for model_column in MODEL_COLUMNS.values():

        data[
            model_column
        ] = pd.to_numeric(
            data[
                model_column
            ],
            errors="coerce"
        )


    # -------------------------------------------------------------------------
    # 依 Year 加總
    # -------------------------------------------------------------------------

    yearly = (

        data
        .groupby(
            "Year",
            as_index=False
        )[
            list(
                MODEL_COLUMNS.values()
            )
        ]
        .sum(
            min_count=1
        )
    )


    # -------------------------------------------------------------------------
    # Wide → Long
    # -------------------------------------------------------------------------

    yearly_long = yearly.melt(

        id_vars=[
            "Year"
        ],

        value_vars=list(
            MODEL_COLUMNS.values()
        ),

        var_name=
            "ModelColumn",

        value_name=
            "AI_RawScore"
    )


    yearly_long[
        "Model"
    ] = (

        yearly_long[
            "ModelColumn"
        ]

        .map(
            reverse_model_map
        )
    )


    yearly_long[
        "Analysis_Subject"
    ] = subject_key


    ai_summary_list.append(
        yearly_long
    )


ai_rawscore_summary = pd.concat(

    ai_summary_list,

    ignore_index=True
)


ai_rawscore_summary[
    "Year"
] = (

    ai_rawscore_summary[
        "Year"
    ]
    .astype(
        "Int64"
    )
)

# =============================================================================
# Step 8. Score scale adjustment
# =============================================================================
#
# 【Chinese】
#
# 2015–2017:
#   AI objectively scoreable maximum = 54
#
#       AI comparison score
#       = AI raw / 54 × 100
#
#   CEEC human distribution uses the full 108-point scale.
#   Step 10 therefore converts official human score-band boundaries to:
#
#       Human comparison boundary
#       = Human raw / 108 × 100
#
#   AI and Human are thus compared on the same 0–100 scale.
#
# 2018–2024:
#   AI and Human are already on the 0–100 scale.
#
#
# 【English】
#
# 2015–2024:
#   AI objectively scoreable maximum = 72
#
#       AI comparison score
#       = AI raw / 72 × 100
#
#   CEEC human distribution is already on the 0–100 scale.
#   Human scores are not transformed.
#
#
# 【AST Mathematics A】
#
# 2015–2021:
#   AI and Human use the raw-score scale.
#
# 2022–2024:
#   AI raw score is converted to the official CEEC 0–60 level scale,
#   then compared with the official human 0–60 level distribution.
#
#   Official Mathematics A intervals:
#       2022 = 1.43700
#       2023 = 1.36100
#       2024 = 1.56867
#
# CEEC level conversion:
#   raw score = 0 -> Level 0
#   0 < X <= 1 × interval -> Level 1
#   (n-1) × interval < X <= n × interval -> Level n
#   X > 59 × interval -> Level 60
#
# CEEC score boundaries are rounded to 2 decimal places.
# =============================================================================


from decimal import Decimal, ROUND_HALF_UP


# -----------------------------------------------------------------------------
# AST Mathematics A 2022–2024 官方級距
# -----------------------------------------------------------------------------

MATHA_AST_GRADE_INTERVAL = {

    2022:
        Decimal("1.43700"),

    2023:
        Decimal("1.36100"),

    2024:
        Decimal("1.56867"),
}


# -----------------------------------------------------------------------------
# CEEC 分數邊界：四捨五入至小數第 2 位
# -----------------------------------------------------------------------------

def round_ceec_2(value):

    return value.quantize(
        Decimal("0.01"),
        rounding=ROUND_HALF_UP
    )


# -----------------------------------------------------------------------------
# AST Mathematics A raw score → 60-level score
# -----------------------------------------------------------------------------

def convert_matha_ast_raw_to_level(
    year,
    raw_score
):

    if pd.isna(
        raw_score
    ):

        return np.nan


    year = int(
        year
    )


    raw_score_float = float(
        raw_score
    )


    # -------------------------------------------------------------------------
    # 2015–2021：
    # 仍使用 raw-score distribution
    # -------------------------------------------------------------------------

    if year <= 2021:

        return raw_score_float


    # -------------------------------------------------------------------------
    # 2022–2024：
    # 必須轉為 0–60 級分
    # -------------------------------------------------------------------------

    if year not in MATHA_AST_GRADE_INTERVAL:

        raise ValueError(
            "\nAST Mathematics A 找不到該年度級距設定。"
            f"\nYear: {year}"
            f"\nRaw score: {raw_score}"
        )


    # -------------------------------------------------------------------------
    # QC：原得總分應為 0–100
    # -------------------------------------------------------------------------

    if (
        raw_score_float < 0
        or
        raw_score_float > 100
    ):

        raise ValueError(
            "\nAST Mathematics A AI raw score 超出 0–100。"
            f"\nYear: {year}"
            f"\nRaw score: {raw_score_float}"
        )


    # -------------------------------------------------------------------------
    # 0 分 = 0 級分
    # -------------------------------------------------------------------------

    if np.isclose(
        raw_score_float,
        0.0,
        atol=1e-12
    ):

        return 0.0


    grade_interval = (
        MATHA_AST_GRADE_INTERVAL[
            year
        ]
    )


    score_decimal = Decimal(
        str(
            raw_score_float
        )
    )


    # -------------------------------------------------------------------------
    # Level 1–59
    #
    # 逐級建立 CEEC 公布的 upper boundary：
    #
    #   upper = round(level × grade_interval, 2)
    #
    # 並依：
    #
    #   X <= upper
    #
    # 找到第一個符合的 level。
    # -------------------------------------------------------------------------

    for level in range(
        1,
        60
    ):

        upper_boundary = round_ceec_2(

            Decimal(
                level
            )
            *
            grade_interval
        )


        if score_decimal <= upper_boundary:

            return float(
                level
            )


    # -------------------------------------------------------------------------
    # 超過 59 級分上界 → 60 級分
    # -------------------------------------------------------------------------

    return 60.0


# -----------------------------------------------------------------------------
# 所有 AI score → comparison scale
# -----------------------------------------------------------------------------

def convert_ai_score_to_comparison_scale(
    subject,
    year,
    raw_score
):

    if pd.isna(
        raw_score
    ):

        return np.nan


    year = int(
        year
    )


    raw_score = float(
        raw_score
    )


    # -------------------------------------------------------------------------
    # Chinese 2015–2017
    # objectively scoreable maximum = 54
    # -------------------------------------------------------------------------

    if (
        subject == "Chinese"
        and
        year in [
            2015,
            2016,
            2017
        ]
    ):

        if (
            raw_score < 0
            or
            raw_score > 54
        ):

            raise ValueError(
                "\nChinese AI raw score is outside the valid 0–54 range."
                f"\nYear: {year}"
                f"\nRaw score: {raw_score}"
            )

        return (
            raw_score
            /
            54.0
            *
            100.0
        )


    # -------------------------------------------------------------------------
    # English 2015–2024
    # objectively scoreable maximum = 72
    # -------------------------------------------------------------------------

    if (
        subject == "English"
        and
        year in range(
            2015,
            2025
        )
    ):

        if (
            raw_score < 0
            or
            raw_score > 72
        ):

            raise ValueError(
                "\nEnglish AI raw score is outside the valid 0–72 range."
                f"\nYear: {year}"
                f"\nRaw score: {raw_score}"
            )

        return (
            raw_score
            /
            72.0
            *
            100.0
        )


    # -------------------------------------------------------------------------
    # AST Mathematics A
    # 2015–2021 -> raw score
    # 2022–2024 -> CEEC 0–60 level
    # -------------------------------------------------------------------------

    if (
        subject
        ==
        "MathematicsA_AST"
    ):

        return convert_matha_ast_raw_to_level(

            year=
                year,

            raw_score=
                raw_score,
        )


    # -------------------------------------------------------------------------
    # Other subjects already use the comparison scale directly.
    # -------------------------------------------------------------------------

    return raw_score


# -----------------------------------------------------------------------------
# 建立 AI comparison score
# -----------------------------------------------------------------------------

ai_rawscore_summary[
    "AI_ComparisonScore"
] = ai_rawscore_summary.apply(

    lambda row:
        convert_ai_score_to_comparison_scale(

            row[
                "Analysis_Subject"
            ],

            int(
                row[
                    "Year"
                ]
            ),

            row[
                "AI_RawScore"
            ]
        ),

    axis=1
)


# -----------------------------------------------------------------------------
# 額外建立 score-scale 說明欄位，方便 QC
# -----------------------------------------------------------------------------

def get_ai_comparison_scale(
    subject,
    year
):

    year = int(
        year
    )


    if (
        subject == "Chinese"
        and
        year in [
            2015,
            2016,
            2017
        ]
    ):

        return "Normalized_0_100"


    if (
        subject == "English"
        and
        year in range(
            2015,
            2025
        )
    ):

        return "Normalized_0_100"


    if (
        subject == "MathematicsA_AST"
        and
        year >= 2022
    ):

        return "AST_Level_0_60"


    return "Raw_Score"


ai_rawscore_summary[
    "AI_ComparisonScale"
] = ai_rawscore_summary.apply(

    lambda row:
        get_ai_comparison_scale(

            row[
                "Analysis_Subject"
            ],

            int(
                row[
                    "Year"
                ]
            )
        ),

    axis=1
)


# -----------------------------------------------------------------------------
# AI completeness QC
# -----------------------------------------------------------------------------

EXPECTED_ANALYSIS_SUBJECTS = {
    "Chinese",
    "English",
    "Mathematics",
    "Science",
    "SocialStudies",
    "MathematicsA_AST",
}

EXPECTED_YEARS = set(
    range(
        2015,
        2025
    )
)

EXPECTED_MODELS = set(
    MODEL_COLUMNS.keys()
)


actual_subjects = set(
    ai_rawscore_summary[
        "Analysis_Subject"
    ]
    .dropna()
    .astype(str)
    .unique()
    .tolist()
)


if actual_subjects != EXPECTED_ANALYSIS_SUBJECTS:

    raise ValueError(
        "\nAI subject coverage does not match the study definition."
        f"\nExpected: {sorted(EXPECTED_ANALYSIS_SUBJECTS)}"
        f"\nActual: {sorted(actual_subjects)}"
    )


_ai_duplicate = ai_rawscore_summary.duplicated(
    subset=[
        "Analysis_Subject",
        "Year",
        "Model",
    ],
    keep=False
)


if _ai_duplicate.any():

    raise ValueError(
        "\nDuplicate AI Subject × Year × Model rows were found."
        f"\nDuplicate rows: {int(_ai_duplicate.sum())}"
    )


for qc_subject, qc_group in ai_rawscore_summary.groupby(
    "Analysis_Subject"
):

    qc_years = set(
        qc_group[
            "Year"
        ]
        .dropna()
        .astype(int)
        .unique()
        .tolist()
    )

    if qc_years != EXPECTED_YEARS:

        raise ValueError(
            "\nAI year coverage is incomplete."
            f"\nSubject: {qc_subject}"
            f"\nExpected: {sorted(EXPECTED_YEARS)}"
            f"\nActual: {sorted(qc_years)}"
        )


for (
    qc_subject,
    qc_year
), qc_group in ai_rawscore_summary.groupby(
    [
        "Analysis_Subject",
        "Year",
    ]
):

    qc_models = set(
        qc_group[
            "Model"
        ]
        .dropna()
        .astype(str)
        .tolist()
    )

    if qc_models != EXPECTED_MODELS:

        raise ValueError(
            "\nAI model coverage is incomplete."
            f"\nSubject: {qc_subject}"
            f"\nYear: {qc_year}"
            f"\nExpected: {sorted(EXPECTED_MODELS)}"
            f"\nActual: {sorted(qc_models)}"
        )


EXPECTED_AI_COMPARISON_ROWS = (
    len(
        EXPECTED_ANALYSIS_SUBJECTS
    )
    *
    len(
        EXPECTED_YEARS
    )
    *
    len(
        EXPECTED_MODELS
    )
)


if len(
    ai_rawscore_summary
) != EXPECTED_AI_COMPARISON_ROWS:

    raise ValueError(
        "\nUnexpected number of AI comparison rows."
        f"\nExpected: {EXPECTED_AI_COMPARISON_ROWS}"
        f"\nActual: {len(ai_rawscore_summary)}"
    )


print(
    "\nAI completeness QC: PASS"
)


# -----------------------------------------------------------------------------
# Chinese / English AI comparison-scale QC
# -----------------------------------------------------------------------------

language_ai_qc = ai_rawscore_summary[

    (
        (
            ai_rawscore_summary[
                "Analysis_Subject"
            ]
            ==
            "Chinese"
        )
        &
        (
            ai_rawscore_summary[
                "Year"
            ].isin(
                [
                    2015,
                    2016,
                    2017
                ]
            )
        )
    )

    |

    (
        ai_rawscore_summary[
            "Analysis_Subject"
        ]
        ==
        "English"
    )

][
    [
        "Analysis_Subject",
        "Year",
        "Model",
        "AI_RawScore",
        "AI_ComparisonScore",
        "AI_ComparisonScale",
    ]
].copy()


invalid_language_ai = language_ai_qc[

    (
        language_ai_qc[
            "AI_ComparisonScore"
        ]
        <
        -1e-12
    )

    |

    (
        language_ai_qc[
            "AI_ComparisonScore"
        ]
        >
        100.0 + 1e-12
    )

].copy()


if not invalid_language_ai.empty:

    raise ValueError(
        "\nChinese / English AI comparison score is outside 0–100."
        f"\n{invalid_language_ai.to_string(index=False)}"
    )


if not (
    language_ai_qc[
        "AI_ComparisonScale"
    ]
    ==
    "Normalized_0_100"
).all():

    bad_scale_labels = language_ai_qc[
        language_ai_qc[
            "AI_ComparisonScale"
        ]
        !=
        "Normalized_0_100"
    ]

    raise ValueError(
        "\nChinese / English AI comparison-scale label is incorrect."
        f"\n{bad_scale_labels.to_string(index=False)}"
    )


print(
    "Chinese 2015–2017 / English 2015–2024 AI 0–100 scale QC: PASS"
)


# -----------------------------------------------------------------------------
# AST Mathematics A 2022–2024 conversion QC
# -----------------------------------------------------------------------------

print("\n")

print(
    "=" * 90
)

print(
    "Step 8 QC. AST Mathematics A 2022–2024 raw score → 60-level conversion"
)

print(
    "=" * 90
)


matha_level_qc = ai_rawscore_summary[

    (
        ai_rawscore_summary[
            "Analysis_Subject"
        ]
        ==
        "MathematicsA_AST"
    )

    &

    (
        ai_rawscore_summary[
            "Year"
        ]
        >=
        2022
    )

][
    [
        "Year",
        "Model",
        "AI_RawScore",
        "AI_ComparisonScore",
        "AI_ComparisonScale"
    ]
].copy()


print(

    matha_level_qc
    .sort_values(
        [
            "Year",
            "Model"
        ]
    )
    .to_string(
        index=False
    )
)

# =============================================================================
# Step 9. 建立 GSAT empirical human distribution
# =============================================================================
#
# 重要：
#
# CEEC 官方 score-band interval 並非所有年度都使用同一邊界規則。
#
# 主要型態：
#
#   1. Closed Range
#      例如 2015：
#
#          59.01 - 64.90
#
#      視為：
#
#          59.01 <= X <= 64.90
#
#
#   2. Left-Open Right-Closed
#      例如 2024：
#
#          30.84 < X <= 37.00
#
#
#   3. Single
#      Grade 0：
#
#          X = 0.00
#
#
# 原始資料中 2022–2024 的 Grade 0：
#
#   RawScore_Lower = NaN
#   RawScore_Upper = NaN
#
# 因此在 filter 之前必須先將 X = 0.00 補成：
#
#   RawScore_Lower = 0
#   RawScore_Upper = 0
#   IntervalType   = Single
#
# =============================================================================


print("\n")

print(
    "=" * 90
)

print(
    "Step 9. 建立 GSAT human empirical distribution"
)

print(
    "=" * 90
)


gsat = human_gsat_raw.copy()


# -----------------------------------------------------------------------------
# 必要欄位確認
# -----------------------------------------------------------------------------

required_gsat_columns = [

    "AcademicYear_CE",
    "Subject",
    "GradeLevel",
    "ScoreRange",
    "RawScore_Lower",
    "RawScore_Upper",
    "IntervalType",
    "Count",
    "CumCount_LowToHigh",
]


missing_gsat_columns = [

    column

    for column
    in required_gsat_columns

    if column not in gsat.columns
]


if missing_gsat_columns:

    raise KeyError(

        "\nGSAT Human baseline 缺少必要欄位："

        f"\n{missing_gsat_columns}"

        f"\n目前欄位："

        f"\n{gsat.columns.tolist()}"
    )


# -----------------------------------------------------------------------------
# Numeric
# -----------------------------------------------------------------------------

for column in [

    "AcademicYear_CE",
    "GradeLevel",
    "RawScore_Lower",
    "RawScore_Upper",
    "Count",
    "CumCount_LowToHigh",

]:

    gsat[
        column
    ] = pd.to_numeric(

        gsat[
            column
        ],

        errors="coerce"
    )


# -----------------------------------------------------------------------------
# Year
# -----------------------------------------------------------------------------

gsat[
    "Year"
] = (

    gsat[
        "AcademicYear_CE"
    ]

    .astype(
        "Int64"
    )
)


# -----------------------------------------------------------------------------
# 保留原始 IntervalType，方便 QC
# -----------------------------------------------------------------------------

gsat[
    "IntervalType_Original"
] = gsat[
    "IntervalType"
]


# =============================================================================
# Step 9A. 修復 Grade 0：X = 0.00
# =============================================================================
#
# 2022–2024 CEEC 表格：
#
#   GradeLevel = 0
#   ScoreRange = X = 0.00
#   RawScore_Lower / Upper = NaN
#
# 將其明確轉成 single-score band：
#
#   lower = 0
#   upper = 0
#
# =============================================================================


def is_zero_single_band(
    row
):

    grade_level = row[
        "GradeLevel"
    ]


    score_range = str(

        row[
            "ScoreRange"
        ]

    ).strip().lower()


    interval_type = str(

        row[
            "IntervalType"
        ]

    ).strip().lower()


    # -------------------------------------------------------------------------
    # GradeLevel 必須為 0
    # -------------------------------------------------------------------------

    if pd.isna(
        grade_level
    ):

        return False


    if int(
        grade_level
    ) != 0:

        return False


    # -------------------------------------------------------------------------
    # 已經是 0–0 的舊年度資料
    # -------------------------------------------------------------------------

    lower = row[
        "RawScore_Lower"
    ]


    upper = row[
        "RawScore_Upper"
    ]


    if (
        pd.notna(
            lower
        )
        and
        pd.notna(
            upper
        )
        and
        np.isclose(
            float(
                lower
            ),
            0.0
        )
        and
        np.isclose(
            float(
                upper
            ),
            0.0
        )
    ):

        return True


    # -------------------------------------------------------------------------
    # 例如：
    #   X = 0.00
    # -------------------------------------------------------------------------

    zero_text = (

        (
            "x" in score_range
            and
            "=" in score_range
            and
            "0" in score_range
        )

        or

        (
            "x" in interval_type
            and
            "=" in interval_type
            and
            "0" in interval_type
        )
    )


    return zero_text


zero_single_mask = gsat.apply(

    is_zero_single_band,

    axis=1
)


gsat.loc[

    zero_single_mask,

    "RawScore_Lower"

] = 0.0


gsat.loc[

    zero_single_mask,

    "RawScore_Upper"

] = 0.0


# =============================================================================
# Step 9B. IntervalType 標準化
# =============================================================================
#
# 最終只允許：
#
#   Closed Range
#   Left-Open Right-Closed
#   Single
#
# =============================================================================


def standardize_gsat_interval_type(
    row
):

    lower = row[
        "RawScore_Lower"
    ]


    upper = row[
        "RawScore_Upper"
    ]


    original = str(

        row[
            "IntervalType"
        ]

    ).strip()


    original_lower = original.lower()


    score_range = str(

        row[
            "ScoreRange"
        ]

    ).strip()


    score_range_lower = score_range.lower()


    # -------------------------------------------------------------------------
    # Single
    #
    # 包括：
    #   0–0
    #   X = 0.00
    # -------------------------------------------------------------------------

    if (
        pd.notna(
            lower
        )
        and
        pd.notna(
            upper
        )
        and
        np.isclose(
            float(
                lower
            ),
            float(
                upper
            )
        )
    ):

        return "Single"


    if (

        (
            "x" in original_lower
            and
            "=" in original_lower
            and
            "0" in original_lower
        )

        or

        (
            "x" in score_range_lower
            and
            "=" in score_range_lower
            and
            "0" in score_range_lower
        )

    ):

        return "Single"


    # -------------------------------------------------------------------------
    # Left-Open Right-Closed
    # -------------------------------------------------------------------------

    if (

        (
            "left-open" in original_lower
            and
            "right-closed" in original_lower
        )

        or

        (
            "＜x≦" in score_range_lower
        )

        or

        (
            "<x≤" in score_range_lower
        )

        or

        (
            "<x<=" in score_range_lower
        )

    ):

        return "Left-Open Right-Closed"


    # -------------------------------------------------------------------------
    # Closed Range
    # -------------------------------------------------------------------------

    if (
        "closed range"
        in
        original_lower
    ):

        return "Closed Range"


    # -------------------------------------------------------------------------
    # 若有合法 lower / upper，但來源型態沒有明確標示：
    # 不自行猜測，保留原文字，後續 QC 會停止。
    # -------------------------------------------------------------------------

    return original


gsat[
    "IntervalType"
] = gsat.apply(

    standardize_gsat_interval_type,

    axis=1
)


# =============================================================================
# Step 9C. 移除說明列
# =============================================================================
#
# Grade 0 已於前面補成 0–0，因此不會被誤刪。
#
# =============================================================================


gsat = gsat[

    gsat[
        "Year"
    ].notna()

    &

    gsat[
        "Subject"
    ].notna()

    &

    gsat[
        "RawScore_Lower"
    ].notna()

    &

    gsat[
        "RawScore_Upper"
    ].notna()

    &

    gsat[
        "Count"
    ].notna()

    &

    gsat[
        "CumCount_LowToHigh"
    ].notna()

].copy()


# -----------------------------------------------------------------------------
# Human source
# -----------------------------------------------------------------------------

gsat[
    "Human_Source"
] = (
    "GSAT"
)


# =============================================================================
# Step 9D. IntervalType QC
# =============================================================================


allowed_gsat_interval_types = {

    "Closed Range",
    "Left-Open Right-Closed",
    "Single",
}


found_gsat_interval_types = set(

    gsat[
        "IntervalType"
    ]
    .dropna()
    .unique()
    .tolist()
)


unknown_gsat_interval_types = (

    found_gsat_interval_types

    -

    allowed_gsat_interval_types
)


print(
    "\nGSAT IntervalType:"
)


print(

    gsat[
        "IntervalType"
    ]
    .value_counts(
        dropna=False
    )
)


if unknown_gsat_interval_types:

    raise ValueError(

        "\nGSAT 發現未定義的 IntervalType："

        f"\n{sorted(unknown_gsat_interval_types)}"

        "\n請先確認官方資料，不可自行推定。"
    )


# -----------------------------------------------------------------------------
# Grade 0 QC
# -----------------------------------------------------------------------------

grade0_qc = gsat[

    gsat[
        "GradeLevel"
    ]
    ==
    0

][
    [
        "Year",
        "Subject",
        "GradeLevel",
        "ScoreRange",
        "RawScore_Lower",
        "RawScore_Upper",
        "IntervalType",
        "Count",
        "CumCount_LowToHigh",
    ]
].copy()


print(
    "\nGSAT Grade 0 QC:"
)


print(

    grade0_qc
    .sort_values(
        [
            "Subject",
            "Year"
        ]
    )
    .to_string(
        index=False
    )
)

# =============================================================================
# Step 10. GSAT human score → comparison scale
# =============================================================================

def convert_gsat_human_boundary(
    subject,
    year,
    value
):

    if pd.isna(
        value
    ):

        return np.nan


    value = float(
        value
    )


    # -------------------------------------------------------------------------
    # Chinese 2015–2017:
    # official human full-score scale 108 -> 0–100 comparison scale
    # -------------------------------------------------------------------------

    if (
        subject == "Chinese"
        and
        year in [
            2015,
            2016,
            2017
        ]
    ):

        return (
            value
            /
            108.0
            *
            100.0
        )


    # -------------------------------------------------------------------------
    # English 2015–2024:
    # official human distribution is already on the 0–100 scale.
    # No transformation is applied.
    # -------------------------------------------------------------------------


    return value


gsat[
    "Comparison_Lower"
] = gsat.apply(

    lambda row:
        convert_gsat_human_boundary(

            row[
                "Subject"
            ],

            int(
                row[
                    "Year"
                ]
            ),

            row[
                "RawScore_Lower"
            ]
        ),

    axis=1
)


gsat[
    "Comparison_Upper"
] = gsat.apply(

    lambda row:
        convert_gsat_human_boundary(

            row[
                "Subject"
            ],

            int(
                row[
                    "Year"
                ]
            ),

            row[
                "RawScore_Upper"
            ]
        ),

    axis=1
)


# -----------------------------------------------------------------------------
# 官方 cumulative count 包含目前 band
#
# 所以：
#
# Count below band
# =
# CumCount_LowToHigh - Count
# -----------------------------------------------------------------------------

gsat[
    "Count_Below_Band"
] = (

    gsat[
        "CumCount_LowToHigh"
    ]

    -

    gsat[
        "Count"
    ]
)


gsat[
    "Total_N"
] = (

    gsat
    .groupby(
        [
            "Year",
            "Subject"
        ]
    )[
        "CumCount_LowToHigh"
    ]
    .transform(
        "max"
    )
)

# =============================================================================
# Step 11. 建立 AST Mathematics A empirical human distribution
# =============================================================================
#
# Human baseline：
#
#   2015–2021
#       BandType = Score
#       BandValue 例如：
#           100.00
#           99.00 - 99.99
#
#       → 使用 raw-score bands
#
#
#   2022–2024
#       BandType = Level
#       BandValue：
#           60
#           59
#           ...
#           0
#
#       → 使用官方 0–60 級分 distribution
#
#
# AI comparison scale 已於 Step 8 同步處理：
#
#   2015–2021 AI comparison score = raw score
#   2022–2024 AI comparison score = official 0–60 level
#
# =============================================================================


print("\n")

print(
    "=" * 90
)

print(
    "Step 11. 建立 AST Mathematics A human empirical distribution"
)

print(
    "=" * 90
)


matha = human_matha_raw.copy()


# -----------------------------------------------------------------------------
# 必要欄位確認
# -----------------------------------------------------------------------------

required_matha_columns = [

    "Year",
    "BandType",
    "BandValue",
    "Count",
    "CumCount_LowToHigh",
]


missing_matha_columns = [

    column

    for column
    in required_matha_columns

    if column not in matha.columns
]


if missing_matha_columns:

    raise KeyError(

        "\nAST Mathematics A human baseline 缺少必要欄位："

        f"\n{missing_matha_columns}"

        f"\n目前欄位："

        f"\n{matha.columns.tolist()}"
    )


# -----------------------------------------------------------------------------
# Numeric
# -----------------------------------------------------------------------------

matha[
    "Year"
] = pd.to_numeric(

    matha[
        "Year"
    ],

    errors="coerce"

).astype(
    "Int64"
)


matha[
    "Count"
] = pd.to_numeric(

    matha[
        "Count"
    ],

    errors="coerce"
)


matha[
    "CumCount_LowToHigh"
] = pd.to_numeric(

    matha[
        "CumCount_LowToHigh"
    ],

    errors="coerce"
)


# -----------------------------------------------------------------------------
# BandType 標準化
# -----------------------------------------------------------------------------

matha[
    "BandType"
] = (

    matha[
        "BandType"
    ]

    .astype(
        "string"
    )

    .str.strip()

    .str.lower()
)


# =============================================================================
# Step 11A. Score band parser
# =============================================================================
#
# 用於 2015–2021：
#
#   100.00
#   99.00 - 99.99
#
# 缺考等文字 → NaN
#
# =============================================================================


def parse_matha_score_band(
    value
):

    if pd.isna(
        value
    ):

        return (
            np.nan,
            np.nan
        )


    text = str(
        value
    ).strip()


    # -------------------------------------------------------------------------
    # Range
    # 支援一般 hyphen / en dash / em dash
    # -------------------------------------------------------------------------

    match_range = re.match(

        r"^\s*"
        r"([0-9]+(?:\.[0-9]+)?)"
        r"\s*[-–—]\s*"
        r"([0-9]+(?:\.[0-9]+)?)"
        r"\s*$",

        text
    )


    if match_range:

        lower = float(
            match_range.group(
                1
            )
        )

        upper = float(
            match_range.group(
                2
            )
        )


        return (
            lower,
            upper
        )


    # -------------------------------------------------------------------------
    # Single raw-score value
    # 例如 100.00
    # -------------------------------------------------------------------------

    match_single = re.match(

        r"^\s*"
        r"([0-9]+(?:\.[0-9]+)?)"
        r"\s*$",

        text
    )


    if match_single:

        number = float(
            match_single.group(
                1
            )
        )


        return (
            number,
            number
        )


    # -------------------------------------------------------------------------
    # 缺考 / 說明文字
    # -------------------------------------------------------------------------

    return (
        np.nan,
        np.nan
    )


# =============================================================================
# Step 11B. Level parser
# =============================================================================
#
# 用於 2022–2024：
#
#   BandValue = 60 ... 0
#
# 每個 level 是離散值，所以：
#
#   lower = level
#   upper = level
#
# Step 15 會套用 single-score band 的 mid-rank rule。
#
# =============================================================================


def parse_matha_level(
    value
):

    if pd.isna(
        value
    ):

        return (
            np.nan,
            np.nan
        )


    try:

        level = float(
            str(
                value
            ).strip()
        )

    except ValueError:

        return (
            np.nan,
            np.nan
        )


    # -------------------------------------------------------------------------
    # 只接受 0–60
    # -------------------------------------------------------------------------

    if (
        level < 0
        or
        level > 60
    ):

        return (
            np.nan,
            np.nan
        )


    # -------------------------------------------------------------------------
    # Level 必須是整數級分
    # -------------------------------------------------------------------------

    if not np.isclose(
        level,
        round(
            level
        ),
        atol=1e-12
    ):

        return (
            np.nan,
            np.nan
        )


    level = float(
        int(
            round(
                level
            )
        )
    )


    return (
        level,
        level
    )


# =============================================================================
# Step 11C. 依 BandType 決定解析方式
# =============================================================================


def parse_matha_band(
    row
):

    band_type = row[
        "BandType"
    ]


    band_value = row[
        "BandValue"
    ]


    # -------------------------------------------------------------------------
    # 2015–2021 raw-score distribution
    # -------------------------------------------------------------------------

    if band_type == "score":

        return parse_matha_score_band(
            band_value
        )


    # -------------------------------------------------------------------------
    # 2022–2024 level distribution
    # -------------------------------------------------------------------------

    if band_type == "level":

        return parse_matha_level(
            band_value
        )


    return (
        np.nan,
        np.nan
    )


parsed_bands = matha.apply(

    parse_matha_band,

    axis=1
)


matha[
    "RawScore_Lower"
] = [

    value[
        0
    ]

    for value
    in parsed_bands
]


matha[
    "RawScore_Upper"
] = [

    value[
        1
    ]

    for value
    in parsed_bands
]


# -----------------------------------------------------------------------------
# 移除無法使用的資料列
#
# 例如：
#   缺考
#   說明列
# -----------------------------------------------------------------------------

matha = matha[

    matha[
        "Year"
    ].notna()

    &

    matha[
        "RawScore_Lower"
    ].notna()

    &

    matha[
        "RawScore_Upper"
    ].notna()

    &

    matha[
        "Count"
    ].notna()

    &

    matha[
        "CumCount_LowToHigh"
    ].notna()

].copy()


# -----------------------------------------------------------------------------
# Subject / Source
# -----------------------------------------------------------------------------

matha[
    "Subject"
] = (
    "MathematicsA_AST"
)


matha[
    "Human_Source"
] = (
    "AST_MathA"
)


# =============================================================================
# Step 11D. Comparison scale
# =============================================================================
#
# 2015–2021：
#   RawScore_Lower / Upper = raw score
#
# 2022–2024：
#   RawScore_Lower / Upper 實際代表 Level 0–60
#
# 因 Step 8 已把 AI 同步轉到相同尺度，
# 所以 Comparison_Lower / Upper 可直接使用此值。
#
# =============================================================================


matha[
    "Comparison_Lower"
] = matha[
    "RawScore_Lower"
]


matha[
    "Comparison_Upper"
] = matha[
    "RawScore_Upper"
]


# -----------------------------------------------------------------------------
# Count below band
#
# CumCount_LowToHigh 包含目前 band，
# 因此：
#
#   Count_Below_Band
#   =
#   CumCount_LowToHigh - Count
# -----------------------------------------------------------------------------

matha[
    "Count_Below_Band"
] = (

    matha[
        "CumCount_LowToHigh"
    ]

    -

    matha[
        "Count"
    ]
)


# -----------------------------------------------------------------------------
# Total N
# -----------------------------------------------------------------------------

matha[
    "Total_N"
] = (

    matha

    .groupby(
        "Year"
    )[
        "CumCount_LowToHigh"
    ]

    .transform(
        "max"
    )
)


# =============================================================================
# Step 11E. MathA scale QC
# =============================================================================


matha_scale_qc = (

    matha

    .groupby(
        [
            "Year",
            "BandType"
        ],
        as_index=False
    )

    .agg(

        Band_N=(
            "BandValue",
            "count"
        ),

        Min_Comparison=(
            "Comparison_Lower",
            "min"
        ),

        Max_Comparison=(
            "Comparison_Upper",
            "max"
        ),

        Total_N=(
            "Total_N",
            "max"
        ),
    )
)


print(
    "\nAST Mathematics A Human scale QC:"
)


print(

    matha_scale_qc
    .to_string(
        index=False
    )
)


# -----------------------------------------------------------------------------
# 強制 QC
#
# 2022–2024 必須：
#   BandType = level
#   Min = 0
#   Max = 60
# -----------------------------------------------------------------------------

for check_year in [
    2022,
    2023,
    2024
]:

    check = matha[

        matha[
            "Year"
        ]
        ==
        check_year

    ].copy()


    if check.empty:

        raise ValueError(

            "\nAST Mathematics A Human baseline 找不到年度："

            f"{check_year}"
        )


    band_types = set(

        check[
            "BandType"
        ]
        .dropna()
        .tolist()
    )


    if band_types != {
        "level"
    }:

        raise ValueError(

            "\nAST Mathematics A 2022–2024 BandType 不符合預期。"

            f"\nYear: {check_year}"

            f"\nBandType: {band_types}"
        )


    min_level = float(

        check[
            "Comparison_Lower"
        ].min()
    )


    max_level = float(

        check[
            "Comparison_Upper"
        ].max()
    )


    if not (
        np.isclose(
            min_level,
            0.0
        )
        and
        np.isclose(
            max_level,
            60.0
        )
    ):

        raise ValueError(

            "\nAST Mathematics A Level 範圍錯誤。"

            f"\nYear: {check_year}"

            f"\nMin: {min_level}"

            f"\nMax: {max_level}"

            "\n預期為 0–60。"
        )


print(
    "\nAST Mathematics A 2022–2024 human Level QC：PASS"
)


# =============================================================================
# Step 12. GSAT Mathematics 對應規則
# =============================================================================
#
# 2015–2021：
#   official Human subject = Math
#
# 2022–2024：
#   GSAT 已分 Mathematics A / B
#   本研究 GSAT Mathematics 使用 Mathematics A
#
# =============================================================================


def get_human_subject(
    analysis_subject,
    year
):

    if analysis_subject == "Chinese":

        return "Chinese"


    if analysis_subject == "English":

        return "English"


    if analysis_subject == "Science":

        return "Science"


    if analysis_subject == "SocialStudies":

        return "SocialStudies"


    if analysis_subject == "Mathematics":

        if year <= 2021:

            return "Math"

        else:

            return "MathA"


    if analysis_subject == "MathematicsA_AST":

        return "MathematicsA_AST"


    return np.nan


ai_rawscore_summary[
    "Human_Subject"
] = ai_rawscore_summary.apply(

    lambda row:
        get_human_subject(

            row[
                "Analysis_Subject"
            ],

            int(
                row[
                    "Year"
                ]
            )
        ),

    axis=1
)

# =============================================================================
# Step 13. 整合 human distributions
# =============================================================================
#
# 目的：
#
# 將 GSAT 與 AST Mathematics A 整合成同一套 human empirical table，
# 並保留正式的 interval definition。
#
# 統一 IntervalType：
#
#   Closed Range
#   Left-Open Right-Closed
#   Single
#
# =============================================================================


print("\n")

print(
    "=" * 90
)

print(
    "Step 13. 整合 Human empirical distributions"
)

print(
    "=" * 90
)


# =============================================================================
# Step 13A. GSAT
# =============================================================================


GSAT_KEEP = [

    "Year",
    "Subject",
    "Human_Source",

    "GradeLevel",
    "ScoreRange",
    "IntervalType",

    "RawScore_Lower",
    "RawScore_Upper",

    "Comparison_Lower",
    "Comparison_Upper",

    "Count",
    "Count_Below_Band",

    "CumCount_LowToHigh",
    "Total_N",
]


missing_gsat_keep = [

    column

    for column
    in GSAT_KEEP

    if column not in gsat.columns
]


if missing_gsat_keep:

    raise KeyError(

        "\nGSAT 缺少 Step 13 所需欄位："

        f"\n{missing_gsat_keep}"
    )


gsat_human = gsat[

    GSAT_KEEP

].copy()


gsat_human = gsat_human.rename(

    columns={

        "Subject":
            "Human_Subject"

    }
)


# =============================================================================
# Step 13B. AST Mathematics A
# =============================================================================
#
# Step 11 已將：
#
#   2015–2021 BandType = score
#       → raw-score bands
#
#   2022–2024 BandType = level
#       → single-value Level
#
# 此處建立統一 IntervalType。
#
# =============================================================================


matha_for_merge = matha.copy()


# -----------------------------------------------------------------------------
# 建立 ScoreRange 顯示欄
# -----------------------------------------------------------------------------

if "BandValue" in matha_for_merge.columns:

    matha_for_merge[
        "ScoreRange"
    ] = (

        matha_for_merge[
            "BandValue"
        ]

        .astype(
            "string"
        )
    )

else:

    matha_for_merge[
        "ScoreRange"
    ] = (
        pd.NA
    )


# -----------------------------------------------------------------------------
# 建立 GradeLevel 顯示欄
#
# AST MathA 不使用 GSAT GradeLevel；
# 2022–2024 BandValue 本身是 Level，但為避免混淆，
# 此欄維持 NaN。
# -----------------------------------------------------------------------------

matha_for_merge[
    "GradeLevel"
] = np.nan


# -----------------------------------------------------------------------------
# 建立統一 IntervalType
# -----------------------------------------------------------------------------

def get_matha_interval_type(
    row
):

    lower = row[
        "Comparison_Lower"
    ]


    upper = row[
        "Comparison_Upper"
    ]


    # -------------------------------------------------------------------------
    # Level / single score
    # -------------------------------------------------------------------------

    if (
        pd.notna(
            lower
        )
        and
        pd.notna(
            upper
        )
        and
        np.isclose(
            float(
                lower
            ),
            float(
                upper
            )
        )
    ):

        return "Single"


    # -------------------------------------------------------------------------
    # 2015–2021 raw score ranges
    # -------------------------------------------------------------------------

    return "Closed Range"


matha_for_merge[
    "IntervalType"
] = matha_for_merge.apply(

    get_matha_interval_type,

    axis=1
)


MATHA_KEEP = [

    "Year",
    "Subject",
    "Human_Source",

    "GradeLevel",
    "ScoreRange",
    "IntervalType",

    "RawScore_Lower",
    "RawScore_Upper",

    "Comparison_Lower",
    "Comparison_Upper",

    "Count",
    "Count_Below_Band",

    "CumCount_LowToHigh",
    "Total_N",
]


missing_matha_keep = [

    column

    for column
    in MATHA_KEEP

    if column not in matha_for_merge.columns
]


if missing_matha_keep:

    raise KeyError(

        "\nAST Mathematics A 缺少 Step 13 所需欄位："

        f"\n{missing_matha_keep}"
    )


matha_human = matha_for_merge[

    MATHA_KEEP

].copy()


matha_human = matha_human.rename(

    columns={

        "Subject":
            "Human_Subject"

    }
)


# =============================================================================
# Step 13C. 合併
# =============================================================================


human_empirical = pd.concat(

    [

        gsat_human,
        matha_human

    ],

    ignore_index=True
)


# -----------------------------------------------------------------------------
# Year 標準化
# -----------------------------------------------------------------------------

human_empirical[
    "Year"
] = pd.to_numeric(

    human_empirical[
        "Year"
    ],

    errors="coerce"

).astype(
    "Int64"
)


# =============================================================================
# Step 13C-1. Retain only human distributions used in this study
# =============================================================================
#
# Study definition:
#   Chinese / English / Science / SocialStudies:
#       2015–2024
#
#   GSAT Mathematics:
#       2015–2021 -> Math
#       2022–2024 -> MathA
#
#   AST Mathematics A:
#       2015–2024 -> MathematicsA_AST
#
# MathB and any other distributions are not part of the analysis and are
# excluded here so that the analysis table, QC outputs, and public report
# contain only distributions actually used for model-to-human comparison.
# =============================================================================

study_human_mask = (

    (
        human_empirical[
            "Human_Subject"
        ].isin(
            [
                "Chinese",
                "English",
                "Science",
                "SocialStudies",
                "MathematicsA_AST",
            ]
        )

        &

        human_empirical[
            "Year"
        ].between(
            2015,
            2024,
            inclusive="both"
        )
    )

    |

    (
        (
            human_empirical[
                "Human_Subject"
            ]
            ==
            "Math"
        )

        &

        human_empirical[
            "Year"
        ].between(
            2015,
            2021,
            inclusive="both"
        )
    )

    |

    (
        (
            human_empirical[
                "Human_Subject"
            ]
            ==
            "MathA"
        )

        &

        human_empirical[
            "Year"
        ].between(
            2022,
            2024,
            inclusive="both"
        )
    )
)


excluded_human_distributions = (

    human_empirical.loc[
        ~study_human_mask,
        [
            "Human_Source",
            "Human_Subject",
            "Year",
        ]
    ]

    .drop_duplicates()

    .sort_values(
        [
            "Human_Source",
            "Human_Subject",
            "Year",
        ]
    )
)


if not excluded_human_distributions.empty:

    print(
        "\nHuman distributions excluded because they are outside the study definition:"
    )

    print(
        excluded_human_distributions.to_string(
            index=False
        )
    )


human_empirical = (

    human_empirical.loc[
        study_human_mask
    ]

    .copy()
)


# -----------------------------------------------------------------------------
# Human-distribution coverage QC
# -----------------------------------------------------------------------------

actual_human_pairs = set(

    zip(
        human_empirical[
            "Human_Subject"
        ].astype(str).tolist(),

        human_empirical[
            "Year"
        ].astype(int).tolist(),
    )
)


expected_human_pairs = {

    *[
        (subject, year)

        for subject in [
            "Chinese",
            "English",
            "Science",
            "SocialStudies",
            "MathematicsA_AST",
        ]

        for year in range(
            2015,
            2025
        )
    ],

    *[
        ("Math", year)

        for year in range(
            2015,
            2022
        )
    ],

    *[
        ("MathA", year)

        for year in range(
            2022,
            2025
        )
    ],
}


if actual_human_pairs != expected_human_pairs:

    missing_pairs = sorted(
        expected_human_pairs
        -
        actual_human_pairs
    )

    extra_pairs = sorted(
        actual_human_pairs
        -
        expected_human_pairs
    )

    raise ValueError(
        "\nHuman empirical distributions do not match the study definition."
        f"\nMissing subject-year pairs: {missing_pairs}"
        f"\nExtra subject-year pairs: {extra_pairs}"
    )


print(
    "\nHuman distribution study-scope QC: PASS"
)


# =============================================================================
# Step 13D. IntervalType 最終 QC
# =============================================================================


ALLOWED_INTERVAL_TYPES = {

    "Closed Range",
    "Left-Open Right-Closed",
    "Single",
}


found_interval_types = set(

    human_empirical[
        "IntervalType"
    ]
    .dropna()
    .unique()
    .tolist()
)


unknown_interval_types = (

    found_interval_types

    -

    ALLOWED_INTERVAL_TYPES
)


print(
    "\nHuman empirical IntervalType:"
)


print(

    human_empirical[
        "IntervalType"
    ]
    .value_counts(
        dropna=False
    )
)


if unknown_interval_types:

    raise ValueError(

        "\nHuman empirical data 發現未定義的 IntervalType："

        f"\n{sorted(unknown_interval_types)}"

        "\n程式停止，請先確認來源資料。"
    )


# =============================================================================
# Step 13E. 每個 Subject × Year 的基本 QC
# =============================================================================


human_distribution_qc = (

    human_empirical

    .groupby(

        [
            "Human_Source",
            "Human_Subject",
            "Year"
        ],

        as_index=False
    )

    .agg(

        Band_N=(
            "Count",
            "count"
        ),

        Human_Min=(
            "Comparison_Lower",
            "min"
        ),

        Human_Max=(
            "Comparison_Upper",
            "max"
        ),

        Sum_Count=(
            "Count",
            "sum"
        ),

        Total_N=(
            "Total_N",
            "max"
        ),
    )
)


human_distribution_qc[
    "Difference"
] = (

    human_distribution_qc[
        "Sum_Count"
    ]

    -

    human_distribution_qc[
        "Total_N"
    ]
)


# -----------------------------------------------------------------------------
# Chinese / English human comparison-scale QC
# -----------------------------------------------------------------------------

language_human_scale_qc = human_distribution_qc[

    (
        (
            human_distribution_qc[
                "Human_Subject"
            ]
            ==
            "Chinese"
        )
        &
        (
            human_distribution_qc[
                "Year"
            ].isin(
                [
                    2015,
                    2016,
                    2017
                ]
            )
        )
    )

    |

    (
        human_distribution_qc[
            "Human_Subject"
        ]
        ==
        "English"
    )

].copy()


for _, qc_row in language_human_scale_qc.iterrows():

    qc_subject = str(
        qc_row[
            "Human_Subject"
        ]
    )

    qc_year = int(
        qc_row[
            "Year"
        ]
    )

    qc_min = float(
        qc_row[
            "Human_Min"
        ]
    )

    qc_max = float(
        qc_row[
            "Human_Max"
        ]
    )


    if not (
        np.isclose(
            qc_min,
            0.0,
            atol=0.02
        )
        and
        np.isclose(
            qc_max,
            100.0,
            atol=0.02
        )
    ):

        raise ValueError(
            "\nLanguage human comparison scale is not 0–100."
            f"\nSubject: {qc_subject}"
            f"\nYear: {qc_year}"
            f"\nMin: {qc_min}"
            f"\nMax: {qc_max}"
        )


print(
    "\nChinese 2015–2017 / English 2015–2024 human 0–100 scale QC: PASS"
)


print(
    "\nHuman distribution QC:"
)


print(

    human_distribution_qc
    .to_string(
        index=False
    )
)


# -----------------------------------------------------------------------------
# 若 Sum Count != Total N，這裡直接警告，
# Step 14 還會再正式輸出 QC。
# -----------------------------------------------------------------------------

qc_difference_n = (

    human_distribution_qc[
        "Difference"
    ]
    !=
    0

).sum()


print(
    "\nHuman distributions with Count sum != Total N:",
    qc_difference_n
)


if qc_difference_n != 0:

    bad_human_count_qc = human_distribution_qc[
        human_distribution_qc[
            "Difference"
        ]
        !=
        0
    ].copy()

    raise ValueError(
        "\nHuman cumulative Count sum does not equal official Total N."
        f"\n{bad_human_count_qc.to_string(index=False)}"
    )


print(
    "Human cumulative-count QC: PASS"
)

# =============================================================================
# Step 14. Human cumulative QC
# =============================================================================

print("\n")
print("=" * 90)
print("Step 14. Human cumulative QC")
print("=" * 90)


qc_cumulative_rows = []


for (
    source,
    subject,
    year
), group in human_empirical.groupby(

    [
        "Human_Source",
        "Human_Subject",
        "Year"
    ]
):

    g = group.copy()


    total_count_sum = g[
        "Count"
    ].sum()


    total_n_official = g[
        "Total_N"
    ].max()


    difference = (
        total_count_sum
        -
        total_n_official
    )


    qc_cumulative_rows.append(

        {

            "Human_Source":
                source,

            "Human_Subject":
                subject,

            "Year":
                year,

            "Sum_Count":
                total_count_sum,

            "Official_Total_N":
                total_n_official,

            "Difference":
                difference,

        }
    )


qc_cumulative = pd.DataFrame(
    qc_cumulative_rows
)


print(
    "Human distributions:",
    len(
        qc_cumulative
    )
)


print(
    "Count sum != Total N:",
    (
        qc_cumulative[
            "Difference"
        ] != 0
    ).sum()
)

# =============================================================================
# Step 15. Empirical Percentile Function
# =============================================================================
#
# Primary empirical percentile method:
#
# 使用 CEEC 官方 empirical cumulative counts，
# 並於 score band 內進行 linear interpolation。
#
#
# Percentile =
#
# [N below band
#  + within-band fraction × N in band]
# ------------------------------------------------ × 100
#                    Total N
#
#
# Band membership 完全依官方 IntervalType：
#
#   1. Closed Range
#
#          lower <= score <= upper
#
#
#   2. Left-Open Right-Closed
#
#          lower < score <= upper
#
#
#   3. Single
#
#          score == lower == upper
#
#
# Single band 無法做 continuous linear interpolation，
# 因此使用 mid-rank：
#
#          within-band fraction = 0.5
#
#
# 對官方因小數顯示造成的微小 score gap，
# 僅允許 <= 0.02 分的 nearest-band correction，
# 並明確標記 Match_Status。
#
# =============================================================================


def empirical_percentile(
    human_data,
    human_subject,
    year,
    score
):

    # -------------------------------------------------------------------------
    # 預設輸出
    # -------------------------------------------------------------------------

    output = {

        "Match_Status":
            np.nan,

        "Band_IntervalType":
            np.nan,

        "Band_ScoreRange":
            np.nan,

        "Band_Lower_Raw":
            np.nan,

        "Band_Upper_Raw":
            np.nan,

        "Band_Lower_Comparison":
            np.nan,

        "Band_Upper_Comparison":
            np.nan,

        "Band_Count":
            np.nan,

        "Count_Below_Band":
            np.nan,

        "Total_N":
            np.nan,

        "Within_Band_Fraction":
            np.nan,

        "Empirical_CDF_Percentile":
            np.nan,

        "Empirical_TopTail_Percent":
            np.nan,
    }


    # -------------------------------------------------------------------------
    # Missing AI score
    # -------------------------------------------------------------------------

    if pd.isna(
        score
    ):

        output[
            "Match_Status"
        ] = "Missing_AI_Score"

        return output


    score = float(
        score
    )


    year = int(
        year
    )


    # -------------------------------------------------------------------------
    # 找同年度 + 同科目 human distribution
    # -------------------------------------------------------------------------

    subset = human_data[

        (
            human_data[
                "Human_Subject"
            ]
            ==
            human_subject
        )

        &

        (
            human_data[
                "Year"
            ]
            ==
            year
        )

    ].copy()


    if subset.empty:

        output[
            "Match_Status"
        ] = "No_Human_Distribution"

        return output


    # -------------------------------------------------------------------------
    # 必要欄位確認
    # -------------------------------------------------------------------------

    required_columns = [

        "IntervalType",
        "Comparison_Lower",
        "Comparison_Upper",
        "RawScore_Lower",
        "RawScore_Upper",
        "Count",
        "Count_Below_Band",
        "Total_N",
    ]


    missing_columns = [

        column

        for column
        in required_columns

        if column not in subset.columns
    ]


    if missing_columns:

        raise KeyError(

            "\nEmpirical percentile 缺少必要 human 欄位："

            f"\n{missing_columns}"
        )


    # -------------------------------------------------------------------------
    # 排序：低分 → 高分
    # -------------------------------------------------------------------------

    subset = (

        subset

        .sort_values(

            [
                "Comparison_Lower",
                "Comparison_Upper"
            ],

            ascending=[
                True,
                True
            ]
        )

        .reset_index(
            drop=True
        )
    )


    # -------------------------------------------------------------------------
    # 基本資訊
    # -------------------------------------------------------------------------

    total_n = float(

        subset[
            "Total_N"
        ]
        .max()
    )


    minimum_score = float(

        subset[
            "Comparison_Lower"
        ]
        .min()
    )


    maximum_score = float(

        subset[
            "Comparison_Upper"
        ]
        .max()
    )


    # -------------------------------------------------------------------------
    # AI score 低於最低範圍
    # -------------------------------------------------------------------------

    if score < minimum_score:

        output[
            "Match_Status"
        ] = "Below_Minimum"

        output[
            "Total_N"
        ] = total_n

        output[
            "Empirical_CDF_Percentile"
        ] = 0.0

        output[
            "Empirical_TopTail_Percent"
        ] = 100.0

        return output


    # -------------------------------------------------------------------------
    # AI score 高於最高範圍
    # -------------------------------------------------------------------------

    if score > maximum_score:

        output[
            "Match_Status"
        ] = "Above_Maximum"

        output[
            "Total_N"
        ] = total_n

        output[
            "Empirical_CDF_Percentile"
        ] = 100.0

        output[
            "Empirical_TopTail_Percent"
        ] = 0.0

        return output


    # =========================================================================
    # Step 15A. 依官方 IntervalType 做 band matching
    # =========================================================================


    MATCH_TOLERANCE = 1e-12


    match_mask = pd.Series(

        False,

        index=subset.index
    )


    for idx, band_row in subset.iterrows():

        lower = float(

            band_row[
                "Comparison_Lower"
            ]
        )


        upper = float(

            band_row[
                "Comparison_Upper"
            ]
        )


        interval_type = str(

            band_row[
                "IntervalType"
            ]

        ).strip()


        # ---------------------------------------------------------------------
        # Case 1：Single
        #
        # score == lower == upper
        # ---------------------------------------------------------------------

        if (
            interval_type
            ==
            "Single"
        ):

            is_match = np.isclose(

                score,

                lower,

                atol=
                    MATCH_TOLERANCE,

                rtol=
                    0.0
            )


        # ---------------------------------------------------------------------
        # Case 2：Left-Open Right-Closed
        #
        # lower < score <= upper
        #
        # 例如：
        #
        #   30.84 < X <= 37.00
        #
        # score = 37.00 → 此 band
        # score = 30.84 → 不屬於此 band
        # ---------------------------------------------------------------------

        elif (
            interval_type
            ==
            "Left-Open Right-Closed"
        ):

            lower_condition = (

                score

                >

                lower
                +
                MATCH_TOLERANCE
            )


            upper_condition = (

                score
                <
                upper

                or

                np.isclose(

                    score,

                    upper,

                    atol=
                        MATCH_TOLERANCE,

                    rtol=
                        0.0
                )
            )


            is_match = (

                lower_condition

                and

                upper_condition
            )


        # ---------------------------------------------------------------------
        # Case 3：Closed Range
        #
        # lower <= score <= upper
        # ---------------------------------------------------------------------

        elif (
            interval_type
            ==
            "Closed Range"
        ):

            lower_condition = (

                score
                >
                lower

                or

                np.isclose(

                    score,

                    lower,

                    atol=
                        MATCH_TOLERANCE,

                    rtol=
                        0.0
                )
            )


            upper_condition = (

                score
                <
                upper

                or

                np.isclose(

                    score,

                    upper,

                    atol=
                        MATCH_TOLERANCE,

                    rtol=
                        0.0
                )
            )


            is_match = (

                lower_condition

                and

                upper_condition
            )


        # ---------------------------------------------------------------------
        # 未定義 IntervalType：
        # 不自行推測
        # ---------------------------------------------------------------------

        else:

            raise ValueError(

                "\n發現未支援的 IntervalType。"

                f"\nHuman subject: "
                f"{human_subject}"

                f"\nYear: "
                f"{year}"

                f"\nIntervalType: "
                f"{interval_type}"

                f"\nBand: "
                f"{lower} - {upper}"
            )


        if is_match:

            match_mask.loc[
                idx
            ] = True


    matched = subset[

        match_mask

    ].copy()


    match_status = (
        "Matched"
    )


    # =========================================================================
    # Step 15B. 官方顯示小數造成的微小 gap
    # =========================================================================
    #
    # 例如舊年度 Closed Range：
    #
    #   59.01–64.90
    #   64.91–70.80
    #
    # 若 score 因轉換或 floating point 落在非常小的 gap，
    # 僅允許距離 <= 0.02 的 nearest-band correction。
    #
    # =========================================================================


    if matched.empty:

        subset[
            "_Distance"
        ] = np.where(

            score
            <
            subset[
                "Comparison_Lower"
            ],

            subset[
                "Comparison_Lower"
            ]
            -
            score,

            np.where(

                score
                >
                subset[
                    "Comparison_Upper"
                ],

                score
                -
                subset[
                    "Comparison_Upper"
                ],

                0.0
            )
        )


        minimum_distance = float(

            subset[
                "_Distance"
            ]
            .min()
        )


        nearest_candidates = subset[

            np.isclose(

                subset[
                    "_Distance"
                ],

                minimum_distance,

                atol=
                    MATCH_TOLERANCE,

                rtol=
                    0.0
            )

        ].copy()


        # ---------------------------------------------------------------------
        # 若有兩個同距離 band：
        #
        # 優先選 score 位於 upper boundary 的 lower band。
        #
        # 這可正確處理 left-open/right-closed shared boundary。
        # ---------------------------------------------------------------------

        if len(
            nearest_candidates
        ) > 1:

            upper_boundary_match = nearest_candidates[

                np.isclose(

                    nearest_candidates[
                        "Comparison_Upper"
                    ],

                    score,

                    atol=
                        MATCH_TOLERANCE,

                    rtol=
                        0.0
                )

            ].copy()


            if len(
                upper_boundary_match
            ) == 1:

                nearest_candidates = (
                    upper_boundary_match
                )


        # ---------------------------------------------------------------------
        # 若仍多筆：
        # 選 lower value 較小者
        # ---------------------------------------------------------------------

        nearest_candidates = (

            nearest_candidates

            .sort_values(

                [
                    "Comparison_Lower",
                    "Comparison_Upper"
                ],

                ascending=[
                    True,
                    True
                ]
            )
        )


        nearest_index = (

            nearest_candidates
            .index[
                0
            ]
        )


        nearest_distance = float(

            subset.loc[

                nearest_index,

                "_Distance"
            ]
        )


        if nearest_distance <= 0.02:

            matched = subset.loc[

                [
                    nearest_index
                ]

            ].copy()


            match_status = (
                "Nearest_Band_RoundingGap"
            )


        else:

            output[
                "Match_Status"
            ] = "No_Band_Match"

            output[
                "Total_N"
            ] = total_n

            return output


    # =========================================================================
    # Step 15C. 最終只能對應一個 band
    # =========================================================================


    if len(
        matched
    ) > 1:

        raise ValueError(

            "\n依官方 IntervalType 判定後仍對應到多個 human bands。"

            f"\nHuman subject: "
            f"{human_subject}"

            f"\nYear: "
            f"{year}"

            f"\nAI score: "
            f"{score}"

            f"\n\n{matched}"
        )


    # =========================================================================
    # Step 15D. 取出唯一 band
    # =========================================================================


    band = matched.iloc[
        0
    ]


    interval_type = str(

        band[
            "IntervalType"
        ]

    ).strip()


    if (
        "ScoreRange"
        in
        band.index
    ):

        score_range = band[
            "ScoreRange"
        ]

    else:

        score_range = np.nan


    raw_lower = float(

        band[
            "RawScore_Lower"
        ]
    )


    raw_upper = float(

        band[
            "RawScore_Upper"
        ]
    )


    lower = float(

        band[
            "Comparison_Lower"
        ]
    )


    upper = float(

        band[
            "Comparison_Upper"
        ]
    )


    band_count = float(

        band[
            "Count"
        ]
    )


    count_below = float(

        band[
            "Count_Below_Band"
        ]
    )


    total_n = float(

        band[
            "Total_N"
        ]
    )


    # =========================================================================
    # Step 15E. Within-band interpolation
    # =========================================================================


    if (
        interval_type
        ==
        "Single"

        or

        np.isclose(

            lower,

            upper,

            atol=
                MATCH_TOLERANCE,

            rtol=
                0.0
        )
    ):

        # ---------------------------------------------------------------------
        # Single-score band：
        #
        # 使用 mid-rank convention。
        # ---------------------------------------------------------------------

        within_fraction = 0.5


    else:

        # ---------------------------------------------------------------------
        # Continuous score band：
        #
        # linear interpolation
        #
        #   (score - lower)
        #   -----------------
        #   (upper - lower)
        #
        # 對 Left-Open Right-Closed：
        #
        #   score = upper → fraction = 1
        #
        # lower 本身不會進到此 band。
        # ---------------------------------------------------------------------

        within_fraction = (

            score
            -
            lower

        ) / (

            upper
            -
            lower
        )


        # ---------------------------------------------------------------------
        # Nearest_Band_RoundingGap 時可能略低於 0 或略高於 1，
        # 因此限制在 [0,1]。
        # ---------------------------------------------------------------------

        within_fraction = max(

            0.0,

            min(

                1.0,

                within_fraction
            )
        )


    # =========================================================================
    # Step 15F. Empirical CDF percentile
    # =========================================================================


    estimated_count_below = (

        count_below

        +

        within_fraction
        *
        band_count
    )


    empirical_cdf = (

        estimated_count_below

        /

        total_n

        *

        100.0
    )


    empirical_cdf = max(

        0.0,

        min(

            100.0,

            empirical_cdf
        )
    )


    # -------------------------------------------------------------------------
    # Top-tail percentage
    #
    # Conventional percentile：
    #   越高 = 表現越好
    #
    # Top-tail percentage：
    #   越低 = 排名前段
    # -------------------------------------------------------------------------

    empirical_top_tail = (

        100.0

        -

        empirical_cdf
    )


    # =========================================================================
    # Step 15G. 回傳
    # =========================================================================


    output.update(

        {

            "Match_Status":
                match_status,

            "Band_IntervalType":
                interval_type,

            "Band_ScoreRange":
                score_range,

            "Band_Lower_Raw":
                raw_lower,

            "Band_Upper_Raw":
                raw_upper,

            "Band_Lower_Comparison":
                lower,

            "Band_Upper_Comparison":
                upper,

            "Band_Count":
                band_count,

            "Count_Below_Band":
                count_below,

            "Total_N":
                total_n,

            "Within_Band_Fraction":
                within_fraction,

            "Empirical_CDF_Percentile":
                empirical_cdf,

            "Empirical_TopTail_Percent":
                empirical_top_tail,
        }
    )


    return output

# =============================================================================
# Step 16. 對全部 AI results 計算 empirical percentile
# =============================================================================

print("\n")
print("=" * 90)
print("Step 16. 計算 Empirical Percentile")
print("=" * 90)


empirical_rows = []


for _, row in ai_rawscore_summary.iterrows():

    year = int(
        row[
            "Year"
        ]
    )


    result = empirical_percentile(

        human_data=
            human_empirical,

        human_subject=
            row[
                "Human_Subject"
            ],

        year=
            year,

        score=
            row[
                "AI_ComparisonScore"
            ],
    )


    empirical_rows.append(
        result
    )


empirical_details = pd.DataFrame(
    empirical_rows
)


empirical_result = pd.concat(

    [

        ai_rawscore_summary
        .reset_index(
            drop=True
        ),

        empirical_details
        .reset_index(
            drop=True
        )

    ],

    axis=1
)


# =============================================================================
# Step 17. Empirical QC
# =============================================================================

print(
    "\n總 AI comparison 筆數：",
    len(
        empirical_result
    )
)


print(
    "\nMatch Status:"
)


print(

    empirical_result[
        "Match_Status"
    ]
    .value_counts(
        dropna=False
    )
)


qc_unmatched = empirical_result[

    empirical_result[
        "Match_Status"
    ].isin(

        [

            "No_Human_Distribution",
            "No_Band_Match",
            "Missing_AI_Score",

        ]
    )

].copy()


if len(
    qc_unmatched
) > 0:

    raise ValueError(
        "\nUnmatched AI comparison rows were found; analysis stopped."
        f"\n{qc_unmatched[['Analysis_Subject', 'Year', 'Model', 'AI_RawScore', 'AI_ComparisonScore', 'Human_Subject', 'Match_Status']].to_string(index=False)}"
    )


print(
    "Empirical percentile mapping QC: PASS"
)


# =============================================================================
# Step 18. Grouped human data → Normal approximation mean / SD
# =============================================================================
#
# 目的：
# 與原先 normal approximation method 做 sensitivity comparison。
#
# 各 band 使用 midpoint 代表。
#
# Weighted mean =
# sum(midpoint × count) / N
#
# Weighted population SD =
# sqrt[
#   sum(count × (midpoint - mean)^2) / N
# ]
#
# =============================================================================

print("\n")
print("=" * 90)
print("Step 18. 建立 Normal Approximation")
print("=" * 90)


human_normal = human_empirical.copy()


human_normal[
    "Band_Midpoint"
] = (

    human_normal[
        "Comparison_Lower"
    ]

    +

    human_normal[
        "Comparison_Upper"
    ]

) / 2.0


normal_summary_rows = []


for (
    human_subject,
    year
), group in human_normal.groupby(

    [
        "Human_Subject",
        "Year"
    ]
):

    g = group[

        group[
            "Band_Midpoint"
        ].notna()

        &

        group[
            "Count"
        ].notna()

    ].copy()


    if g.empty:

        continue


    weights = g[
        "Count"
    ].astype(
        float
    )


    midpoints = g[
        "Band_Midpoint"
    ].astype(
        float
    )


    total_n = weights.sum()


    grouped_mean = (

        weights
        *
        midpoints

    ).sum() / total_n


    grouped_variance = (

        weights

        *

        (
            midpoints
            -
            grouped_mean
        )
        ** 2

    ).sum() / total_n


    grouped_sd = math.sqrt(
        grouped_variance
    )


    normal_summary_rows.append(

        {

            "Human_Subject":
                human_subject,

            "Year":
                year,

            "Human_Total_N":
                total_n,

            "Grouped_Mean":
                grouped_mean,

            "Grouped_SD":
                grouped_sd,
        }
    )


normal_human_summary = pd.DataFrame(
    normal_summary_rows
)


normal_human_summary[
    "Year"
] = normal_human_summary[
    "Year"
].astype(
    "Int64"
)


# =============================================================================
# Step 19. Merge normal statistics
# =============================================================================

comparison = empirical_result.merge(

    normal_human_summary,

    on=[
        "Human_Subject",
        "Year"
    ],

    how="left",

    validate=
        "many_to_one"
)


# =============================================================================
# Step 20. Normal CDF
# =============================================================================

def normal_cdf(
    z
):

    if pd.isna(
        z
    ):

        return np.nan


    return (

        0.5

        *

        (
            1.0

            +

            math.erf(

                z
                /
                math.sqrt(
                    2.0
                )
            )
        )
    )


comparison[
    "Normal_Z"
] = (

    comparison[
        "AI_ComparisonScore"
    ]

    -

    comparison[
        "Grouped_Mean"
    ]

) / comparison[
    "Grouped_SD"
]


comparison[
    "Normal_CDF_Percentile"
] = (

    comparison[
        "Normal_Z"
    ]
    .apply(
        normal_cdf
    )

    *
    100.0
)


comparison[
    "Normal_TopTail_Percent"
] = (

    100.0

    -

    comparison[
        "Normal_CDF_Percentile"
    ]
)


# =============================================================================
# Step 21. Normal vs Empirical difference
# =============================================================================

comparison[
    "Difference_Empirical_minus_Normal"
] = (

    comparison[
        "Empirical_CDF_Percentile"
    ]

    -

    comparison[
        "Normal_CDF_Percentile"
    ]
)


comparison[
    "Absolute_Difference"
] = (

    comparison[
        "Difference_Empirical_minus_Normal"
    ]
    .abs()
)


# =============================================================================
# Step 22. Percentile threshold sensitivity
# =============================================================================
#
# 檢查主要解讀是否因方法改變。
#
# 例如：
# ≥50th
# ≥75th
# ≥90th
# ≥95th
#
# =============================================================================

THRESHOLDS = [
    50,
    75,
    90,
    95
]


for threshold in THRESHOLDS:

    comparison[
        f"Empirical_GE_{threshold}"
    ] = (

        comparison[
            "Empirical_CDF_Percentile"
        ]
        >=
        threshold
    )


    comparison[
        f"Normal_GE_{threshold}"
    ] = (

        comparison[
            "Normal_CDF_Percentile"
        ]
        >=
        threshold
    )


    comparison[
        f"Agreement_GE_{threshold}"
    ] = (

        comparison[
            f"Empirical_GE_{threshold}"
        ]

        ==

        comparison[
            f"Normal_GE_{threshold}"
        ]
    )


# =============================================================================
# Step 23. Percentile interpretation category
# =============================================================================

def percentile_category(
    percentile
):

    if pd.isna(
        percentile
    ):

        return np.nan


    if percentile >= 95:

        return ">=95th"


    if percentile >= 90:

        return "90th-<95th"


    if percentile >= 75:

        return "75th-<90th"


    if percentile >= 50:

        return "50th-<75th"


    return "<50th"


comparison[
    "Empirical_Percentile_Category"
] = (

    comparison[
        "Empirical_CDF_Percentile"
    ]
    .apply(
        percentile_category
    )
)


comparison[
    "Normal_Percentile_Category"
] = (

    comparison[
        "Normal_CDF_Percentile"
    ]
    .apply(
        percentile_category
    )
)


comparison[
    "Category_Agreement"
] = (

    comparison[
        "Empirical_Percentile_Category"
    ]

    ==

    comparison[
        "Normal_Percentile_Category"
    ]
)


# =============================================================================
# Step 24. Overall sensitivity summary
# =============================================================================

overall_summary = pd.DataFrame(

    {

        "Metric": [

            "N comparisons",

            "Mean empirical percentile",

            "Mean normal percentile",

            "Mean absolute difference",

            "Median absolute difference",

            "Maximum absolute difference",

            "Within 1 percentile point (%)",

            "Within 2 percentile points (%)",

            "Within 5 percentile points (%)",

            "Agreement at >=50th threshold (%)",

            "Agreement at >=75th threshold (%)",

            "Agreement at >=90th threshold (%)",

            "Agreement at >=95th threshold (%)",

            "Percentile-category agreement (%)",

        ],


        "Value": [

            len(
                comparison
            ),

            comparison[
                "Empirical_CDF_Percentile"
            ].mean(),

            comparison[
                "Normal_CDF_Percentile"
            ].mean(),

            comparison[
                "Absolute_Difference"
            ].mean(),

            comparison[
                "Absolute_Difference"
            ].median(),

            comparison[
                "Absolute_Difference"
            ].max(),

            (
                comparison[
                    "Absolute_Difference"
                ]
                <=
                1
            ).mean()
            *
            100,

            (
                comparison[
                    "Absolute_Difference"
                ]
                <=
                2
            ).mean()
            *
            100,

            (
                comparison[
                    "Absolute_Difference"
                ]
                <=
                5
            ).mean()
            *
            100,

            comparison[
                "Agreement_GE_50"
            ].mean()
            *
            100,

            comparison[
                "Agreement_GE_75"
            ].mean()
            *
            100,

            comparison[
                "Agreement_GE_90"
            ].mean()
            *
            100,

            comparison[
                "Agreement_GE_95"
            ].mean()
            *
            100,

            comparison[
                "Category_Agreement"
            ].mean()
            *
            100,

        ]
    }
)


overall_summary[
    "Value"
] = overall_summary[
    "Value"
].round(
    3
)


# =============================================================================
# Step 25. Summary by Subject
# =============================================================================

subject_summary = (

    comparison
    .groupby(
        "Analysis_Subject",
        as_index=False
    )
    .agg(

        N=(
            "Absolute_Difference",
            "count"
        ),

        Mean_Empirical_Percentile=(
            "Empirical_CDF_Percentile",
            "mean"
        ),

        Mean_Normal_Percentile=(
            "Normal_CDF_Percentile",
            "mean"
        ),

        Mean_Absolute_Difference=(
            "Absolute_Difference",
            "mean"
        ),

        Median_Absolute_Difference=(
            "Absolute_Difference",
            "median"
        ),

        Max_Absolute_Difference=(
            "Absolute_Difference",
            "max"
        ),

        Category_Agreement_Percent=(
            "Category_Agreement",
            lambda x:
                x.mean()
                *
                100
        ),

        Agreement_90th_Percent=(
            "Agreement_GE_90",
            lambda x:
                x.mean()
                *
                100
        ),

        Agreement_95th_Percent=(
            "Agreement_GE_95",
            lambda x:
                x.mean()
                *
                100
        ),
    )
)


# =============================================================================
# Step 26. Summary by Model
# =============================================================================

model_summary = (

    comparison
    .groupby(
        "Model",
        as_index=False
    )
    .agg(

        N=(
            "Absolute_Difference",
            "count"
        ),

        Mean_Empirical_Percentile=(
            "Empirical_CDF_Percentile",
            "mean"
        ),

        Mean_Normal_Percentile=(
            "Normal_CDF_Percentile",
            "mean"
        ),

        Mean_Absolute_Difference=(
            "Absolute_Difference",
            "mean"
        ),

        Median_Absolute_Difference=(
            "Absolute_Difference",
            "median"
        ),

        Max_Absolute_Difference=(
            "Absolute_Difference",
            "max"
        ),

        Category_Agreement_Percent=(
            "Category_Agreement",
            lambda x:
                x.mean()
                *
                100
        ),

        Agreement_90th_Percent=(
            "Agreement_GE_90",
            lambda x:
                x.mean()
                *
                100
        ),

        Agreement_95th_Percent=(
            "Agreement_GE_95",
            lambda x:
                x.mean()
                *
                100
        ),
    )
)


# =============================================================================
# Step 26A. Manuscript-ready Table 8
# =============================================================================
#
# Table 8 保留原稿的 human-cohort descriptive summary 結構：
#   Subject | Year | Human N | Mean | SD
#
# 注意：
#   Mean / SD 為 grouped-data midpoint estimates。
#   新版 primary percentile 不使用此 normal approximation；
#   Mean / SD 僅保留做 descriptive reporting 與 sensitivity analysis。
# =============================================================================

SUBJECT_DISPLAY_MAP = {
    "Chinese": "Chinese",
    "English": "English",
    "Mathematics": "GSAT Mathematics",
    "Science": "Natural Science",
    "Natural Science": "Natural Science",
    "SocialStudies": "Social Studies",
    "Social Studies": "Social Studies",
    "MathematicsA_AST": "AST Mathematics A",
    "AST Mathematics A": "AST Mathematics A",
    "GSAT Mathematics": "GSAT Mathematics",

    # human baseline 內部名稱
    "Math": "GSAT Mathematics",
    "MathA": "GSAT Mathematics",
}

SUBJECT_ORDER = [
    "Chinese",
    "English",
    "GSAT Mathematics",
    "Social Studies",
    "Natural Science",
    "AST Mathematics A",
]

MODEL_DISPLAY_MAP = {
    "GPT-4o": "GPT-4o",
    "o1-preview": "o1-preview",
    "o1-mini": "o1-mini",
    "GPT-5.4 Thinking": "GPT-5.4-Thinking",
    "Gemini 3 Flash Thinking": "Gemini-3-Flash-Thinking",
}

MODEL_ORDER = [
    "GPT-4o",
    "o1-preview",
    "o1-mini",
    "GPT-5.4-Thinking",
    "Gemini-3-Flash-Thinking",
]


table8 = normal_human_summary.copy()

# -----------------------------------------------------------------------------
# Table 8 再次確認僅保留本研究實際使用的 human distributions
# -----------------------------------------------------------------------------
# GSAT Mathematics 的官方科目名稱在 2022 年制度變更後分成 MathA / MathB。
# 本研究的 GSAT Mathematics 定義為：
#   2015–2021 → Math
#   2022–2024 → MathA
# MathB 不屬於本研究分析，因此必須在建立正式 Table 8 前排除。
#
# 其他研究科目：
#   Chinese / English / Science / SocialStudies → 2015–2024
#   MathematicsA_AST → 2015–2024
# -----------------------------------------------------------------------------

table8["Year"] = pd.to_numeric(
    table8["Year"],
    errors="coerce"
).astype("Int64")

_TABLE8_NON_MATH_SUBJECTS = {
    "Chinese",
    "English",
    "Science",
    "SocialStudies",
    "MathematicsA_AST",
}

_table8_keep_mask = (
    (
        table8["Human_Subject"].isin(
            _TABLE8_NON_MATH_SUBJECTS
        )
        &
        table8["Year"].between(
            2015,
            2024,
            inclusive="both"
        )
    )
    |
    (
        (table8["Human_Subject"] == "Math")
        &
        table8["Year"].between(
            2015,
            2021,
            inclusive="both"
        )
    )
    |
    (
        (table8["Human_Subject"] == "MathA")
        &
        table8["Year"].between(
            2022,
            2024,
            inclusive="both"
        )
    )
)

_table8_excluded = table8.loc[
    ~_table8_keep_mask,
    ["Human_Subject", "Year"]
].copy()

if not _table8_excluded.empty:
    print("\nTable 8 排除不屬於本研究的 human distributions：")
    print(
        _table8_excluded
        .drop_duplicates()
        .sort_values(["Human_Subject", "Year"])
        .to_string(index=False)
    )

table8 = table8.loc[
    _table8_keep_mask
].copy()

# -----------------------------------------------------------------------------
# 研究科目名稱轉成 manuscript 顯示名稱
# -----------------------------------------------------------------------------

table8["Subject"] = (
    table8["Human_Subject"]
    .map(SUBJECT_DISPLAY_MAP)
)

if table8["Subject"].isna().any():
    unknown_subjects = sorted(
        table8.loc[
            table8["Subject"].isna(),
            "Human_Subject"
        ]
        .astype(str)
        .unique()
        .tolist()
    )
    raise ValueError(
        "\nTable 8 發現未定義 Subject mapping："
        f"\n{unknown_subjects}"
    )

# -----------------------------------------------------------------------------
# 強制 QC：GSAT Mathematics 必須只有
#   2015–2021 = Math
#   2022–2024 = MathA
# -----------------------------------------------------------------------------

_table8_math_source_qc = table8.loc[
    table8["Human_Subject"].isin(["Math", "MathA"]),
    ["Human_Subject", "Year"]
].drop_duplicates().copy()

# Year 先轉成一般 Python int，避免 pandas Int64 / numpy int64
# 與一般 int 因 dtype 不同而被 DataFrame.equals() 誤判。
_table8_math_source_qc["Year"] = pd.to_numeric(
    _table8_math_source_qc["Year"],
    errors="raise"
).astype(int)

actual_math_pairs = set(
    zip(
        _table8_math_source_qc["Year"].tolist(),
        _table8_math_source_qc["Human_Subject"].astype(str).tolist(),
    )
)

expected_math_pairs = {
    *[(year, "Math") for year in range(2015, 2022)],
    *[(year, "MathA") for year in range(2022, 2025)],
}

if actual_math_pairs != expected_math_pairs:
    missing_pairs = sorted(expected_math_pairs - actual_math_pairs)
    extra_pairs = sorted(actual_math_pairs - expected_math_pairs)

    raise ValueError(
        "\nTable 8 GSAT Mathematics human baseline 年度來源錯誤。"
        "\n預期：2015–2021 使用 Math；2022–2024 使用 MathA。"
        f"\n缺少：{missing_pairs}"
        f"\n多出：{extra_pairs}"
        f"\n目前：\n{_table8_math_source_qc.sort_values(['Year', 'Human_Subject']).to_string(index=False)}"
    )

print(
    "Table 8 GSAT Mathematics human baseline 年度來源 QC：PASS"
)

# 正式欄位
table8 = table8[
    [
        "Subject",
        "Year",
        "Human_Total_N",
        "Grouped_Mean",
        "Grouped_SD",
    ]
].copy()

table8 = table8.rename(
    columns={
        "Human_Total_N": "Human N",
        "Grouped_Mean": "Mean",
        "Grouped_SD": "SD",
    }
)

# 固定排序
table8["_SubjectOrder"] = pd.Categorical(
    table8["Subject"],
    categories=SUBJECT_ORDER,
    ordered=True
)

table8 = (
    table8
    .sort_values(
        ["_SubjectOrder", "Year"]
    )
    .drop(columns=["_SubjectOrder"])
    .reset_index(drop=True)
)

# 數值格式
table8["Year"] = pd.to_numeric(
    table8["Year"],
    errors="raise"
).astype(int)

table8["Human N"] = pd.to_numeric(
    table8["Human N"],
    errors="raise"
).round(0).astype(int)

table8["Mean"] = pd.to_numeric(
    table8["Mean"],
    errors="raise"
).round(2)

table8["SD"] = pd.to_numeric(
    table8["SD"],
    errors="raise"
).round(2)

# Table 8 應為 6 subjects × 10 years = 60 rows
TABLE8_EXPECTED_ROWS = 60

if len(table8) != TABLE8_EXPECTED_ROWS:
    raise ValueError(
        "\nTable 8 列數不符合預期。"
        f"\n目前：{len(table8)}"
        f"\n預期：{TABLE8_EXPECTED_ROWS}"
    )

# 每科必須 10 年
_table8_year_qc = (
    table8
    .groupby("Subject")["Year"]
    .nunique()
)

if not (_table8_year_qc == 10).all():
    raise ValueError(
        "\nTable 8 有科目不是完整 2015–2024 共 10 年："
        f"\n{_table8_year_qc.to_dict()}"
    )


# =============================================================================
# Step 26B. Manuscript-ready Table 9 — empirical percentile
# =============================================================================
#
# Table 9 percentile uses the primary empirical method：
#   official CEEC cumulative counts
#   + within-band linear interpolation
#   + single-score band mid-rank
#
# 正式結構：
#   Subject | Model | Mean %ile | SD | Min | Max | N
#
# 統計單位：
#   每一 Subject × Model 的 2015–2024 annual empirical percentiles (n=10)
# =============================================================================

_table9_source = comparison.copy()

_table9_source["Subject"] = (
    _table9_source["Analysis_Subject"]
    .map(SUBJECT_DISPLAY_MAP)
)

_table9_source["Model_Display"] = (
    _table9_source["Model"]
    .map(MODEL_DISPLAY_MAP)
)

if _table9_source["Subject"].isna().any():
    unknown_subjects = sorted(
        _table9_source.loc[
            _table9_source["Subject"].isna(),
            "Analysis_Subject"
        ]
        .astype(str)
        .unique()
        .tolist()
    )
    raise ValueError(
        "\nTable 9 發現未定義 Subject mapping："
        f"\n{unknown_subjects}"
    )

if _table9_source["Model_Display"].isna().any():
    unknown_models = sorted(
        _table9_source.loc[
            _table9_source["Model_Display"].isna(),
            "Model"
        ]
        .astype(str)
        .unique()
        .tolist()
    )
    raise ValueError(
        "\nTable 9 發現未定義 Model mapping："
        f"\n{unknown_models}"
    )

# 只允許可以正式計算 empirical percentile 的 rows
_table9_unavailable = _table9_source[
    _table9_source["Empirical_CDF_Percentile"].isna()
].copy()

if not _table9_unavailable.empty:
    raise ValueError(
        "\nTable 9 存在缺少 empirical percentile 的資料，停止輸出正式表格。"
        f"\n缺少筆數：{len(_table9_unavailable)}"
        "\n請先檢查 09_QC_Unmatched。"
    )

# 每個 Subject × Model × Year 應只有一筆
_table9_duplicate = _table9_source.duplicated(
    subset=[
        "Subject",
        "Model_Display",
        "Year",
    ],
    keep=False
)

if _table9_duplicate.any():
    raise ValueError(
        "\nTable 9 發現 Subject × Model × Year 重複資料，停止。"
        f"\n重複筆數：{int(_table9_duplicate.sum())}"
    )

# 正式 Table 9
table9 = (
    _table9_source
    .groupby(
        ["Subject", "Model_Display"],
        as_index=False
    )
    .agg(
        Mean_Percentile=(
            "Empirical_CDF_Percentile",
            "mean"
        ),
        SD=(
            "Empirical_CDF_Percentile",
            "std"
        ),
        Min=(
            "Empirical_CDF_Percentile",
            "min"
        ),
        Max=(
            "Empirical_CDF_Percentile",
            "max"
        ),
        N=(
            "Empirical_CDF_Percentile",
            "count"
        ),
    )
)

table9 = table9.rename(
    columns={
        "Model_Display": "Model",
        "Mean_Percentile": "Mean %ile",
    }
)

# 固定投稿順序
table9["_SubjectOrder"] = pd.Categorical(
    table9["Subject"],
    categories=SUBJECT_ORDER,
    ordered=True
)

table9["_ModelOrder"] = pd.Categorical(
    table9["Model"],
    categories=MODEL_ORDER,
    ordered=True
)

table9 = (
    table9
    .sort_values(
        ["_SubjectOrder", "_ModelOrder"]
    )
    .drop(columns=["_SubjectOrder", "_ModelOrder"])
    .reset_index(drop=True)
)

# Table 9 應為 6 subjects × 5 models = 30 rows
TABLE9_EXPECTED_ROWS = 30

if len(table9) != TABLE9_EXPECTED_ROWS:
    raise ValueError(
        "\nTable 9 列數不符合預期。"
        f"\n目前：{len(table9)}"
        f"\n預期：{TABLE9_EXPECTED_ROWS}"
    )

# 每一格必須完整 10 年
if not (table9["N"] == 10).all():
    bad_n = table9.loc[
        table9["N"] != 10,
        ["Subject", "Model", "N"]
    ]
    raise ValueError(
        "\nTable 9 有 Subject × Model 未滿 10 年，停止輸出正式表格。"
        f"\n{bad_n.to_string(index=False)}"
    )

# 與 manuscript 現有表格一致：percentile 顯示至小數 1 位
for column in [
    "Mean %ile",
    "SD",
    "Min",
    "Max",
]:
    table9[column] = pd.to_numeric(
        table9[column],
        errors="raise"
    ).round(1)

table9["N"] = table9["N"].astype(int)


# =============================================================================
# Step 26C. Table 8 / Table 9 QC 顯示
# =============================================================================

print("\n")
print("=" * 90)
print("Step 26C. Manuscript-ready Table 8 / Table 9")
print("=" * 90)

print("\nTable 8 rows:", len(table8))
print("Table 9 rows:", len(table9))

print("\nTable 9 empirical percentile preview:")
print(
    table9.to_string(
        index=False
    )
)


# =============================================================================
# Step 27. Round summary
# =============================================================================

for dataframe in [

    subject_summary,
    model_summary,

]:

    numeric_columns = dataframe.select_dtypes(
        include=[
            np.number
        ]
    ).columns


    dataframe[
        numeric_columns
    ] = dataframe[
        numeric_columns
    ].round(
        3
    )


# =============================================================================
# Step 28. Largest differences
# =============================================================================

largest_difference = (

    comparison
    .sort_values(

        "Absolute_Difference",

        ascending=False
    )
    [
        [

            "Analysis_Subject",
            "Year",
            "Model",

            "AI_RawScore",
            "AI_ComparisonScore",

            "Human_Subject",

            "Empirical_CDF_Percentile",
            "Normal_CDF_Percentile",

            "Empirical_TopTail_Percent",
            "Normal_TopTail_Percent",

            "Difference_Empirical_minus_Normal",
            "Absolute_Difference",

            "Empirical_Percentile_Category",
            "Normal_Percentile_Category",

            "Category_Agreement",

            "Match_Status",

        ]
    ]
    .head(
        50
    )
)


# =============================================================================
# Step 29. Score-scale QC
# =============================================================================

score_scale_qc_rows = []


for (
    analysis_subject,
    year
), group in comparison.groupby(

    [
        "Analysis_Subject",
        "Year"
    ]
):

    human_min = group[
        "Band_Lower_Comparison"
    ].min()


    human_max = group[
        "Band_Upper_Comparison"
    ].max()


    ai_min = group[
        "AI_ComparisonScore"
    ].min()


    ai_max = group[
        "AI_ComparisonScore"
    ].max()


    score_scale_qc_rows.append(

        {

            "Analysis_Subject":
                analysis_subject,

            "Year":
                year,

            "AI_Min_ComparisonScore":
                ai_min,

            "AI_Max_ComparisonScore":
                ai_max,

            "Matched_Human_Band_Min":
                human_min,

            "Matched_Human_Band_Max":
                human_max,
        }
    )


score_scale_qc = pd.DataFrame(
    score_scale_qc_rows
)



# -----------------------------------------------------------------------------
# Explicit score-scale checks for the harmonized language subjects
# -----------------------------------------------------------------------------

for check_year in [
    2015,
    2016,
    2017
]:

    check = human_distribution_qc[
        (
            human_distribution_qc[
                "Human_Subject"
            ]
            ==
            "Chinese"
        )
        &
        (
            human_distribution_qc[
                "Year"
            ]
            ==
            check_year
        )
    ]

    if len(
        check
    ) != 1:

        raise ValueError(
            "\nChinese human scale QC row is missing or duplicated."
            f"\nYear: {check_year}"
        )

    if not (
        np.isclose(
            float(
                check.iloc[0][
                    "Human_Min"
                ]
            ),
            0.0,
            atol=0.02
        )
        and
        np.isclose(
            float(
                check.iloc[0][
                    "Human_Max"
                ]
            ),
            100.0,
            atol=0.02
        )
    ):

        raise ValueError(
            "\nChinese 2015–2017 human comparison scale is not 0–100."
            f"\nYear: {check_year}"
        )


for check_year in range(
    2015,
    2025
):

    check = human_distribution_qc[
        (
            human_distribution_qc[
                "Human_Subject"
            ]
            ==
            "English"
        )
        &
        (
            human_distribution_qc[
                "Year"
            ]
            ==
            check_year
        )
    ]

    if len(
        check
    ) != 1:

        raise ValueError(
            "\nEnglish human scale QC row is missing or duplicated."
            f"\nYear: {check_year}"
        )

    if not (
        np.isclose(
            float(
                check.iloc[0][
                    "Human_Min"
                ]
            ),
            0.0,
            atol=0.02
        )
        and
        np.isclose(
            float(
                check.iloc[0][
                    "Human_Max"
                ]
            ),
            100.0,
            atol=0.02
        )
    ):

        raise ValueError(
            "\nEnglish human comparison scale is not 0–100."
            f"\nYear: {check_year}"
        )


print(
    "\nLanguage score-scale QC: PASS"
)


# =============================================================================
# Step 30. Display QC
# =============================================================================

print("\n")
print("=" * 90)
print("Step 30. Sensitivity Analysis Results")
print("=" * 90)


print(
    "\n總 comparisons：",
    len(
        comparison
    )
)


print(
    "\nOverall Summary:"
)

print(
    overall_summary.to_string(
        index=False
    )
)


print(
    "\nBy Subject:"
)

print(
    subject_summary.to_string(
        index=False
    )
)


print(
    "\nBy Model:"
)

print(
    model_summary.to_string(
        index=False
    )
)


# =============================================================================
# Step 31. 數值四捨五入（輸出版）
# =============================================================================

comparison_export = comparison.copy()


round_columns = [

    "AI_RawScore",
    "AI_ComparisonScore",

    "Band_Lower_Raw",
    "Band_Upper_Raw",

    "Band_Lower_Comparison",
    "Band_Upper_Comparison",

    "Within_Band_Fraction",

    "Empirical_CDF_Percentile",
    "Empirical_TopTail_Percent",

    "Grouped_Mean",
    "Grouped_SD",

    "Normal_Z",
    "Normal_CDF_Percentile",
    "Normal_TopTail_Percent",

    "Difference_Empirical_minus_Normal",
    "Absolute_Difference",
]


for column in round_columns:

    if column in comparison_export.columns:

        comparison_export[
            column
        ] = comparison_export[
            column
        ].round(
            4
        )


# =============================================================================
# Step 32. Run information
# =============================================================================

run_info = pd.DataFrame(

    {

        "Item": [

            "Analysis",
            "Run time",

            "AI subjects",
            "Years",
            "Models",

            "Empirical method",

            "Within-band interpolation",

            "Single-score band rule",

            "Normal approximation",

            "Chinese 2015-2017 AI scale",
            "Chinese 2015-2017 human scale",

            "English 2015-2024 AI scale",
            "English 2015-2024 human scale",

            "GSAT Mathematics 2015-2021 human baseline",
            "GSAT Mathematics 2022-2024 human baseline",

            "AST Mathematics A 2015-2021 comparison scale",
            "AST Mathematics A 2022-2024 comparison scale",

            "Excel output file",
            "Word output file",
            "Manuscript Table 8/9 Excel",
            "Manuscript Table 8/9 Word",

        ],


        "Value": [

            "Empirical percentile sensitivity analysis",

            datetime.now().strftime(
                "%Y-%m-%d %H:%M:%S"
            ),

            "Chinese; English; Mathematics; Science; SocialStudies; MathematicsA_AST",

            "2015-2024",

            "; ".join(
                MODEL_COLUMNS.keys()
            ),

            "Official empirical cumulative counts",

            (
                "Linear interpolation within each official score band: "
                "(AI score - lower boundary) / "
                "(upper boundary - lower boundary)"
            ),

            "Mid-rank (0.5 of band count)",

            (
                "Grouped-data midpoint weighted mean and "
                "population SD followed by normal CDF; sensitivity analysis only"
            ),

            "AI raw score / 54 × 100",
            "official human boundary / 108 × 100",

            "AI raw score / 72 × 100",
            "official human distribution already on 0–100 scale",

            "Math",
            "MathA",

            "raw-score scale",
            "official CEEC 0–60 level scale",

            f"output/{OUTPUT_FILE.name}",

            f"output/{WORD_OUTPUT_FILE.name}",

            f"output/{MANUSCRIPT_TABLE_EXCEL.name}",

            f"output/{MANUSCRIPT_TABLE_WORD.name}",
        ]
    }
)


# =============================================================================
# Step 33. Export Excel
# =============================================================================

print("\n")
print("=" * 90)
print("Step 33. 輸出 Excel")
print("=" * 90)


with pd.ExcelWriter(

    OUTPUT_FILE,

    engine=
        "openpyxl"

) as writer:


    run_info.to_excel(

        writer,

        sheet_name=
            "00_RunInfo",

        index=False
    )


    ai_rawscore_summary.to_excel(

        writer,

        sheet_name=
            "01_AI_RawScore",

        index=False
    )


    human_empirical.to_excel(

        writer,

        sheet_name=
            "02_Human_Empirical",

        index=False
    )


    empirical_result.to_excel(

        writer,

        sheet_name=
            "03_Empirical_Percentile",

        index=False
    )


    comparison_export.to_excel(

        writer,

        sheet_name=
            "04_Normal_vs_Empirical",

        index=False
    )


    overall_summary.to_excel(

        writer,

        sheet_name=
            "05_Sensitivity_Overall",

        index=False
    )


    subject_summary.to_excel(

        writer,

        sheet_name=
            "06_By_Subject",

        index=False
    )


    model_summary.to_excel(

        writer,

        sheet_name=
            "07_By_Model",

        index=False
    )


    largest_difference.to_excel(

        writer,

        sheet_name=
            "08_Largest_Difference",

        index=False
    )


    qc_unmatched.to_excel(

        writer,

        sheet_name=
            "09_QC_Unmatched",

        index=False
    )


    normal_human_summary.to_excel(

        writer,

        sheet_name=
            "10_Human_NormalStats",

        index=False
    )


    qc_cumulative.to_excel(

        writer,

        sheet_name=
            "11_QC_Cumulative",

        index=False
    )


    score_scale_qc.to_excel(

        writer,

        sheet_name=
            "12_QC_ScoreScale",

        index=False
    )


    table8.to_excel(

        writer,

        sheet_name=
            "13_Table8_HumanDist",

        index=False
    )


    table9.to_excel(

        writer,

        sheet_name=
            "14_Table9_Empirical",

        index=False
    )


# =============================================================================
# Step 33A. Export manuscript-ready Excel (Table 8 + Table 9 only)
# =============================================================================

print("\n")
print("=" * 90)
print("Step 33A. 輸出 Manuscript-ready Table 8 / Table 9 Excel")
print("=" * 90)


def style_manuscript_worksheet(
    worksheet,
    table_number
):

    header_fill = PatternFill(
        "solid",
        fgColor="D9EAF7"
    )

    thin_gray = Side(
        style="thin",
        color="B7B7B7"
    )

    # Header
    for cell in worksheet[1]:
        cell.font = Font(
            name="Arial",
            size=10,
            bold=True
        )
        cell.fill = header_fill
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True
        )
        cell.border = Border(
            bottom=thin_gray
        )

    # Body
    for row in worksheet.iter_rows(
        min_row=2
    ):
        for cell in row:
            cell.font = Font(
                name="Arial",
                size=10
            )
            cell.alignment = Alignment(
                vertical="center"
            )

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions

    # Column widths
    if table_number == 8:
        widths = {
            "A": 24,
            "B": 10,
            "C": 14,
            "D": 12,
            "E": 12,
        }

        for row in range(2, worksheet.max_row + 1):
            worksheet[f"C{row}"].number_format = "#,##0"
            worksheet[f"D{row}"].number_format = "0.00"
            worksheet[f"E{row}"].number_format = "0.00"

    else:
        widths = {
            "A": 24,
            "B": 28,
            "C": 14,
            "D": 10,
            "E": 10,
            "F": 10,
            "G": 8,
        }

        for row in range(2, worksheet.max_row + 1):
            for col in ["C", "D", "E", "F"]:
                worksheet[f"{col}{row}"].number_format = "0.0"

            # 與原 manuscript note 一致：Mean percentile >= 90 加粗
            mean_cell = worksheet[f"C{row}"]
            if (
                mean_cell.value is not None
                and float(mean_cell.value) >= 90
            ):
                mean_cell.font = Font(
                    name="Arial",
                    size=10,
                    bold=True
                )

    for column_letter, width in widths.items():
        worksheet.column_dimensions[column_letter].width = width


with pd.ExcelWriter(
    MANUSCRIPT_TABLE_EXCEL,
    engine="openpyxl"
) as table_writer:

    table8.to_excel(
        table_writer,
        sheet_name="Table 8",
        index=False
    )

    table9.to_excel(
        table_writer,
        sheet_name="Table 9",
        index=False
    )

    style_manuscript_worksheet(
        table_writer.sheets["Table 8"],
        table_number=8
    )

    style_manuscript_worksheet(
        table_writer.sheets["Table 9"],
        table_number=9
    )


# =============================================================================
# Step 34. Export Word report
# =============================================================================
#
# 將 empirical percentile sensitivity analysis 的主要結果整理成 Word。
# Excel 保留完整明細；Word 提供方法、結果與 QC 摘要。
# =============================================================================

print("\n")
print("=" * 90)
print("Step 34. 輸出 Word")
print("=" * 90)


def format_word_value(value):

    if pd.isna(value):
        return ""

    if isinstance(value, (np.bool_, bool)):
        return "Yes" if bool(value) else "No"

    if isinstance(value, (np.integer, int)):
        return f"{int(value):,}"

    if isinstance(value, (np.floating, float)):
        return f"{float(value):.3f}"

    return str(value)


def add_dataframe_table(
    document,
    dataframe,
    title=None,
    max_rows=None
):

    if title:
        document.add_heading(
            title,
            level=2
        )

    df_word = dataframe.copy()

    if max_rows is not None:
        df_word = df_word.head(
            max_rows
        )

    if df_word.empty:
        document.add_paragraph(
            "No rows."
        )
        return

    table = document.add_table(
        rows=1,
        cols=len(df_word.columns)
    )

    table.style = "Table Grid"

    header_cells = table.rows[0].cells

    for col_index, column in enumerate(df_word.columns):
        header_cells[col_index].text = str(column)

    for _, data_row in df_word.iterrows():

        cells = table.add_row().cells

        for col_index, column in enumerate(df_word.columns):
            cells[col_index].text = format_word_value(
                data_row[column]
            )

    document.add_paragraph()


# -----------------------------------------------------------------------------
# 34A. 建立 Word
# -----------------------------------------------------------------------------

word_document = Document()

section = word_document.sections[0]
section.orientation = WD_ORIENT.LANDSCAPE
section.page_width, section.page_height = (
    section.page_height,
    section.page_width
)
section.top_margin = Inches(0.6)
section.bottom_margin = Inches(0.6)
section.left_margin = Inches(0.6)
section.right_margin = Inches(0.6)

word_document.styles["Normal"].font.name = "Arial"
word_document.styles["Normal"].font.size = Pt(9)

# -----------------------------------------------------------------------------
# 34B. 標題
# -----------------------------------------------------------------------------

title = word_document.add_heading(
    "Empirical Percentile Sensitivity Analysis",
    level=0
)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

subtitle = word_document.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle.add_run(
    "Taiwan GSAT (2015–2024) and AST Mathematics A"
).bold = True

word_document.add_paragraph(
    "Generated: "
    + datetime.now().strftime(
        "%Y-%m-%d %H:%M:%S"
    )
)

# -----------------------------------------------------------------------------
# 34C. Analysis objective
# -----------------------------------------------------------------------------

word_document.add_heading(
    "1. Analysis objective",
    level=1
)

word_document.add_paragraph(
    "Human-cohort percentile ranks are recalculated directly from the official "
    "CEEC empirical grade-band counts and cumulative distributions using an explicit "
    "within-band interpolation rule. The empirical results are then compared with "
    "the grouped-data normal-approximation method as a sensitivity analysis."
)

# -----------------------------------------------------------------------------
# 34D. 方法摘要
# -----------------------------------------------------------------------------

word_document.add_heading(
    "2. Analysis method",
    level=1
)

method_points = [
    (
        "Empirical percentile",
        "Official CEEC cumulative counts were used directly. For continuous score "
        "bands, the within-band fraction was computed by linear interpolation: "
        "(AI score - lower boundary) / (upper boundary - lower boundary)."
    ),
    (
        "Single-score band",
        "A mid-rank convention was used (0.5 of the band count)."
    ),
    (
        "Normal approximation",
        "Each official band was represented by its midpoint; the grouped-data "
        "weighted mean and population SD were calculated, followed by the normal CDF."
    ),
    (
        "Score-scale harmonization",
        "Chinese AI scores for 2015–2017 were converted from the 54-point objectively "
        "scoreable scale to 0–100 (raw/54×100), while the corresponding official human "
        "108-point score-band boundaries were converted to 0–100 (raw/108×100). "
        "English AI scores for 2015–2024 were converted from the 72-point objectively "
        "scoreable scale to 0–100 (raw/72×100); the official English human distribution "
        "was already on a 0–100 scale."
    ),
    (
        "Sensitivity checks",
        "Absolute percentile differences, percentile-category agreement, and agreement "
        "at the 50th, 75th, 90th, and 95th percentile thresholds were calculated."
    ),
]

for heading_text, body_text in method_points:
    paragraph = word_document.add_paragraph(
        style="List Bullet"
    )
    run = paragraph.add_run(
        heading_text + ": "
    )
    run.bold = True
    paragraph.add_run(
        body_text
    )

# -----------------------------------------------------------------------------
# 34E. 自動摘要主要結果
# -----------------------------------------------------------------------------

summary_lookup = dict(
    zip(
        overall_summary["Metric"],
        overall_summary["Value"]
    )
)

mean_abs_diff = summary_lookup.get(
    "Mean absolute difference",
    np.nan
)

median_abs_diff = summary_lookup.get(
    "Median absolute difference",
    np.nan
)

max_abs_diff = summary_lookup.get(
    "Maximum absolute difference",
    np.nan
)

agreement_90 = summary_lookup.get(
    "Agreement at >=90th threshold (%)",
    np.nan
)

agreement_95 = summary_lookup.get(
    "Agreement at >=95th threshold (%)",
    np.nan
)

category_agreement = summary_lookup.get(
    "Percentile-category agreement (%)",
    np.nan
)

word_document.add_heading(
    "3. Overall sensitivity result",
    level=1
)

result_paragraph = word_document.add_paragraph()
result_paragraph.add_run(
    "Across all model × subject × year comparisons, "
)
result_paragraph.add_run(
    f"the mean absolute empirical-versus-normal percentile difference was "
    f"{format_word_value(mean_abs_diff)} percentile points "
).bold = True
result_paragraph.add_run(
    f"(median {format_word_value(median_abs_diff)}; "
    f"maximum {format_word_value(max_abs_diff)}). "
)
result_paragraph.add_run(
    f"Agreement at the ≥90th-percentile threshold was "
    f"{format_word_value(agreement_90)}%, "
    f"agreement at the ≥95th-percentile threshold was "
    f"{format_word_value(agreement_95)}%, and overall percentile-category agreement "
    f"was {format_word_value(category_agreement)}%."
)

add_dataframe_table(
    word_document,
    overall_summary,
    title="Overall sensitivity summary"
)

add_dataframe_table(
    word_document,
    subject_summary,
    title="Sensitivity summary by subject"
)

add_dataframe_table(
    word_document,
    model_summary,
    title="Sensitivity summary by model"
)

# -----------------------------------------------------------------------------
# 34F. 最大差異 Top 20
# -----------------------------------------------------------------------------

largest_difference_word_columns = [
    "Analysis_Subject",
    "Year",
    "Model",
    "AI_ComparisonScore",
    "Empirical_CDF_Percentile",
    "Normal_CDF_Percentile",
    "Difference_Empirical_minus_Normal",
    "Absolute_Difference",
    "Empirical_Percentile_Category",
    "Normal_Percentile_Category",
    "Category_Agreement",
    "Match_Status",
]

largest_difference_word = largest_difference[
    [
        column
        for column in largest_difference_word_columns
        if column in largest_difference.columns
    ]
].copy()

add_dataframe_table(
    word_document,
    largest_difference_word,
    title="Largest empirical-versus-normal differences (Top 20)",
    max_rows=20
)

# -----------------------------------------------------------------------------
# 34G. QC
# -----------------------------------------------------------------------------

word_document.add_heading(
    "4. Quality-control summary",
    level=1
)

qc_paragraph = word_document.add_paragraph()
qc_paragraph.add_run(
    "Total AI comparisons: "
).bold = True
qc_paragraph.add_run(
    f"{len(comparison):,}"
)
qc_paragraph.add_run(
    " | Unmatched comparisons: "
).bold = True
qc_paragraph.add_run(
    f"{len(qc_unmatched):,}"
)

if len(qc_unmatched) == 0:
    word_document.add_paragraph(
        "All AI comparison rows were mapped to a corresponding human distribution/band "
        "under the implemented matching rules."
    )
else:
    qc_word_columns = [
        "Analysis_Subject",
        "Year",
        "Model",
        "AI_RawScore",
        "AI_ComparisonScore",
        "Human_Subject",
        "Match_Status",
    ]

    add_dataframe_table(
        word_document,
        qc_unmatched[
            [
                column
                for column in qc_word_columns
                if column in qc_unmatched.columns
            ]
        ],
        title="Unmatched rows",
        max_rows=50
    )

add_dataframe_table(
    word_document,
    qc_cumulative,
    title="Human cumulative-count QC"
)

# -----------------------------------------------------------------------------
# 34H. 輸出路徑
# -----------------------------------------------------------------------------

word_document.add_heading(
    "5. Output files",
    level=1
)

word_document.add_paragraph(
    f"Excel: output/{OUTPUT_FILE.name}"
)

word_document.add_paragraph(
    f"Word: output/{WORD_OUTPUT_FILE.name}"
)

word_document.add_paragraph(
    f"Manuscript tables (Excel): output/{MANUSCRIPT_TABLE_EXCEL.name}"
)

word_document.add_paragraph(
    f"Manuscript tables (Word): output/{MANUSCRIPT_TABLE_WORD.name}"
)

word_document.save(
    WORD_OUTPUT_FILE
)


# =============================================================================
# Step 34I. Export manuscript-ready Word — Table 8 + Table 9 only
# =============================================================================

print("\n")
print("=" * 90)
print("Step 34I. 輸出 Manuscript-ready Table 8 / Table 9 Word")
print("=" * 90)


def set_cell_text(
    cell,
    value,
    bold=False,
    align_center=False,
    font_size=8.5
):

    cell.text = ""
    paragraph = cell.paragraphs[0]

    if align_center:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run = paragraph.add_run(
        format_word_value(value)
    )
    run.font.name = "Arial"
    run.font.size = Pt(font_size)
    run.bold = bold


def add_formal_manuscript_table(
    document,
    dataframe,
    caption,
    note,
    table_number
):

    caption_paragraph = document.add_paragraph()
    caption_run = caption_paragraph.add_run(caption)
    caption_run.bold = True
    caption_run.font.name = "Arial"
    caption_run.font.size = Pt(10)

    table = document.add_table(
        rows=1,
        cols=len(dataframe.columns)
    )
    table.style = "Table Grid"
    table.autofit = True

    # Header
    for col_index, column in enumerate(dataframe.columns):
        set_cell_text(
            table.rows[0].cells[col_index],
            column,
            bold=True,
            align_center=True,
            font_size=8.5
        )

    # Body
    for _, data_row in dataframe.iterrows():
        row_cells = table.add_row().cells

        for col_index, column in enumerate(dataframe.columns):
            value = data_row[column]

            # Word 顯示格式
            if table_number == 8:
                if column == "Human N":
                    display_value = f"{int(value):,}"
                elif column in ["Mean", "SD"]:
                    display_value = f"{float(value):.2f}"
                elif column == "Year":
                    display_value = str(int(value))
                else:
                    display_value = str(value)

                bold_value = False

            else:
                if column in ["Mean %ile", "SD", "Min", "Max"]:
                    display_value = f"{float(value):.1f}"
                elif column == "N":
                    display_value = str(int(value))
                else:
                    display_value = str(value)

                # Table 9：mean empirical percentile >= 90 以粗體標示
                bold_value = (
                    column == "Mean %ile"
                    and float(value) >= 90
                )

            set_cell_text(
                row_cells[col_index],
                display_value,
                bold=bold_value,
                align_center=(
                    column not in ["Subject", "Model"]
                ),
                font_size=8.0
            )

    note_paragraph = document.add_paragraph()
    note_run = note_paragraph.add_run(
        "Note: " + note
    )
    note_run.italic = True
    note_run.font.name = "Arial"
    note_run.font.size = Pt(8.5)

    document.add_paragraph()


manuscript_document = Document()

manuscript_section = manuscript_document.sections[0]
manuscript_section.orientation = WD_ORIENT.LANDSCAPE
manuscript_section.page_width, manuscript_section.page_height = (
    manuscript_section.page_height,
    manuscript_section.page_width
)
manuscript_section.top_margin = Inches(0.5)
manuscript_section.bottom_margin = Inches(0.5)
manuscript_section.left_margin = Inches(0.5)
manuscript_section.right_margin = Inches(0.5)

manuscript_document.styles["Normal"].font.name = "Arial"
manuscript_document.styles["Normal"].font.size = Pt(9)

# Table 8
add_formal_manuscript_table(
    manuscript_document,
    table8,
    caption=(
        "Table 8. Estimated Human-cohort Score Distributions by "
        "Subject × Year (official CEEC)."
    ),
    note=(
        "Human N was obtained from the official CEEC subject-by-year distributions. "
        "Mean and SD are grouped-data estimates calculated from official score-band "
        "midpoints weighted by examinee counts after scale harmonization where required. "
        "For Chinese in 2015–2017, AI objective-choice scores were converted from the "
        "54-point objectively scoreable scale to 0–100 (raw/54×100), and the corresponding "
        "official 108-point human score-band boundaries were converted to 0–100 "
        "(raw/108×100). For English in 2015–2024, AI objective scores were converted from "
        "the 72-point objectively scoreable scale to 0–100 (raw/72×100); the official "
        "English human distribution was already on a 0–100 scale. These grouped mean/SD "
        "estimates are retained for descriptive reporting and the normal-approximation "
        "sensitivity analysis; the primary percentile analysis uses the empirical CEEC "
        "cumulative counts directly. For 2022–2024 GSAT Mathematics, Mathematics A "
        "examinees are used."
    ),
    table_number=8
)

manuscript_document.add_page_break()

# Table 9
add_formal_manuscript_table(
    manuscript_document,
    table9,
    caption=(
        "Table 9. AI Model Mean Human-cohort Percentile Rank by "
        "Subject × Model Using the Empirical CEEC Distribution."
    ),
    note=(
        "Each annual percentile was calculated directly from the official CEEC "
        "empirical cumulative grade-band counts. Within a continuous score band, "
        "linear interpolation was applied as (AI score − lower boundary) / "
        "(upper boundary − lower boundary); single-score bands used the mid-rank "
        "convention. Before percentile mapping, Chinese AI objective-choice scores for "
        "2015–2017 were converted from 54 points to a 0–100 comparison scale "
        "(raw/54×100), and the corresponding official 108-point human score-band "
        "boundaries were converted to 0–100 (raw/108×100). English AI objective scores "
        "for 2015–2024 were converted from 72 points to 0–100 (raw/72×100), while the "
        "official English human distribution was already on a 0–100 scale. Mean, SD, "
        "Min, and Max summarize the 10 annual empirical percentiles from 2015–2024; "
        "N = number of years. Bold indicates mean empirical percentile ≥ 90."
    ),
    table_number=9
)

manuscript_document.save(
    MANUSCRIPT_TABLE_WORD
)


# =============================================================================
# Step 35. Final check
# =============================================================================

if not OUTPUT_FILE.exists():

    raise RuntimeError(
        "\nExcel 輸出失敗。"
    )

if not WORD_OUTPUT_FILE.exists():

    raise RuntimeError(
        "\nWord 輸出失敗。"
    )

if not MANUSCRIPT_TABLE_EXCEL.exists():

    raise RuntimeError(
        "\nManuscript Table 8/9 Excel 輸出失敗。"
    )

if not MANUSCRIPT_TABLE_WORD.exists():

    raise RuntimeError(
        "\nManuscript Table 8/9 Word 輸出失敗。"
    )

print("\n")
print("=" * 90)
print("分析完成")
print("=" * 90)

print(
    "\nAI comparisons：",
    len(
        comparison
    )
)

print(
    "\n無法 mapping 筆數：",
    len(
        qc_unmatched
    )
)

print(
    "\nOutput directory："
)
print(
    OUTPUT_DIR
)

print(
    "\nExcel："
)
print(
    OUTPUT_FILE
)

print(
    "\nWord："
)
print(
    WORD_OUTPUT_FILE
)

print(
    "\nManuscript-ready Table 8/9 Excel："
)
print(
    MANUSCRIPT_TABLE_EXCEL
)

print(
    "\nManuscript-ready Table 8/9 Word："
)
print(
    MANUSCRIPT_TABLE_WORD
)

print("\n")
print(
    "Excel 請先查看："
)
print(
    "05_Sensitivity_Overall"
)
print(
    "06_By_Subject"
)
print(
    "08_Largest_Difference"
)
print(
    "09_QC_Unmatched"
)
print(
    "13_Table8_HumanDist"
)
print(
    "14_Table9_Empirical"
)

print("\n")
print(
    "Empirical percentile sensitivity analysis completed."
)
