# -*- coding: utf-8 -*-
# =============================================================================
# paired_blocked_statistical_analysis.py
#
# Paired / Blocked Statistical Analysis
# Taiwan GSAT (2015–2024) and AST Mathematics A
#
# Version:
#   Zero hard-coded statistical results
#   + superscript significance star directly after p-value in Word
#   + NO separate "Sig." column
#
# =============================================================================
#
# 【統計設計】
#
# Block:
#     Year (2015–2024; n = 10)
#
# Repeated factor:
#     Model (5 models)
#
# Primary omnibus test:
#     Friedman test
#
# Effect size:
#     Kendall's W
#
# Post-hoc:
#     Pairwise two-sided Wilcoxon signed-rank test
#
# Multiple-comparison correction:
#     Holm correction within each subject
#
#
# 【重要】
#
# 本程式不手填任何統計結果：
#
#   - Mean
#   - SD
#   - Median
#   - Friedman chi-square
#   - Kendall's W
#   - Wilcoxon W
#   - raw p
#   - Holm-adjusted p
#
# # 均由 repository 中 6 份 analysis-ready Excel 自動計算。
#
#
# 【Word 顯著標示】
#
# 不再使用：
#
#     p        Sig.
#     0.0195    *
#
# 改成：
#
#     p
#     0.0195*
#
# Word 中 * 為真正 superscript。
#
#
# Table 2:
#     Shapiro-Wilk p < 0.05 → p 後面上標 *
#
# Table 3:
#     Holm-adjusted p < 0.05
#     且 corresponding Friedman omnibus significant
#     → Holm p 後面上標 *
#
#     Raw p 不加星號。
#
# Table 5:
#     Levene p < 0.05 → p 後面上標 *
#
# Table 6:
#     Friedman p < 0.05 → p 後面上標 *
#
# Table 7:
#     無顯著標記。
#
#
# Required packages:
#
# pip install pandas numpy scipy statsmodels openpyxl python-docx
#
# =============================================================================


# =============================================================================
# STEP 1. Import packages
# =============================================================================

from pathlib import Path
from datetime import datetime

import itertools
import platform
import re
import sys

import numpy as np
import pandas as pd

from scipy.stats import (
    shapiro,
    levene,
    friedmanchisquare,
    wilcoxon,
)

from statsmodels.stats.multitest import multipletests

from openpyxl import load_workbook

from openpyxl.styles import (
    Font,
    PatternFill,
    Alignment,
    Border,
    Side,
)

from openpyxl.utils import get_column_letter


from docx import Document

from docx.shared import (
    Inches,
    Pt,
)

from docx.enum.section import WD_ORIENT

from docx.enum.text import (
    WD_ALIGN_PARAGRAPH,
)

from docx.enum.table import (
    WD_TABLE_ALIGNMENT,
    WD_CELL_VERTICAL_ALIGNMENT,
)

from docx.oxml import OxmlElement
from docx.oxml.ns import qn


# =============================================================================
# STEP 2. Project paths
# =============================================================================
#
# Repository structure:
#
# GSAT-Scoring/
# ├─ analysis/
# │  └─ paired_blocked_statistical_analysis.py
# │
# ├─ supplementary-materials/
# │  └─ analysis-ready-data/
# │     ├─ GSAT_Chinese.xlsx
# │     ├─ GSAT_English.xlsx
# │     ├─ GSAT_Mathematics.xlsx
# │     ├─ GSAT_Social_Studies.xlsx
# │     ├─ GSAT_Science.xlsx
# │     └─ MathA.xlsx
# │
# └─ output/
#
# No computer-specific absolute path is used.
# The project root is resolved from this .py file itself.
# =============================================================================

SCRIPT_DIR = Path(__file__).resolve().parent
PROJECT_DIR = SCRIPT_DIR.parent

DATA_DIR = (
    PROJECT_DIR
    / "supplementary-materials"
    / "analysis-ready-data"
)

OUTPUT_DIR = PROJECT_DIR / "output"

OUTPUT_DIR.mkdir(
    parents=True,
    exist_ok=True,
)


# =============================================================================
# STEP 3. Input files
# =============================================================================

INPUT_FILES = {

    "Chinese":
        DATA_DIR
        / "GSAT_Chinese.xlsx",

    "English":
        DATA_DIR
        / "GSAT_English.xlsx",

    "GSAT Mathematics":
        DATA_DIR
        / "GSAT_Mathematics.xlsx",

    "Social Studies":
        DATA_DIR
        / "GSAT_Social_Studies.xlsx",

    "Natural Science":
        DATA_DIR
        / "GSAT_Science.xlsx",

    "AST Mathematics A":
        DATA_DIR
        / "MathA.xlsx",
}


# =============================================================================
# STEP 4. Output files
# =============================================================================

OUTPUT_EXCEL = (
    OUTPUT_DIR
    / "paired_blocked_statistical_analysis.xlsx"
)

OUTPUT_LOG = (
    OUTPUT_DIR
    / "paired_blocked_statistical_analysis_log.txt"
)

OUTPUT_ACCURACY_CSV = (
    OUTPUT_DIR
    / "annual_objective_accuracy.csv"
)

OUTPUT_FRIEDMAN_CSV = (
    OUTPUT_DIR
    / "friedman_results.csv"
)

OUTPUT_WILCOXON_CSV = (
    OUTPUT_DIR
    / "wilcoxon_holm_results.csv"
)

OUTPUT_WORD = (
    OUTPUT_DIR
    / "paired_blocked_statistical_tables.docx"
)


# =============================================================================
# STEP 5. Analysis settings
# =============================================================================

YEARS = list(
    range(
        2015,
        2025,
    )
)


MODELS = [
    "GPT-4o",
    "o1-preview",
    "o1-mini",
    "GPT-5.4-Thinking",
    "Gemini-3-Flash-Thinking",
]


ALPHA = 0.05


# =============================================================================
# STEP 6. Objectively scoreable annual maximum
#
# These are scoring definitions / denominators,
# NOT statistical results.
# =============================================================================

OBJECTIVE_MAX = {

    "Chinese": {

        2015: 54,
        2016: 54,
        2017: 54,

        2018: 100,
        2019: 100,
        2020: 100,
        2021: 100,
        2022: 100,
        2023: 100,
        2024: 100,
    },


    "English": {

        year: 72

        for year
        in YEARS
    },


    "GSAT Mathematics": {

        year: 100

        for year
        in YEARS
    },


    "Social Studies": {

        year: 144

        for year
        in YEARS
    },


    "Natural Science": {

        year: 128

        for year
        in YEARS
    },


    "AST Mathematics A": {

        year: 100

        for year
        in YEARS
    },
}


# =============================================================================
# STEP 7. Column aliases
# =============================================================================

YEAR_ALIASES = [
    "Year",
    "ExamYear",
    "Exam_Year",
]


ROW_TYPE_ALIASES = [
    "Row_Type",
    "RowType",
]


MODEL_ALIASES = {


    "GPT-4o": [

        "GPT4o_Score",
        "GPT-4o_Score",

        "GPT4o",
        "GPT-4o",

        "GPT4o_RawTotal",
        "GPT-4o_RawTotal",

        "GPT4o_FinalTotal",
        "GPT-4o_FinalTotal",
    ],


    "o1-preview": [

        "o1_preview_Score",
        "o1-preview_Score",
        "o1preview_Score",

        "o1_preview",
        "o1-preview",

        "o1_preview_RawTotal",
        "o1-preview_RawTotal",

        "o1_preview_FinalTotal",
        "o1-preview_FinalTotal",
    ],


    "o1-mini": [

        "o1_mini_Score",
        "o1-mini_Score",
        "o1mini_Score",

        "o1_mini",
        "o1-mini",

        "o1_mini_RawTotal",
        "o1-mini_RawTotal",

        "o1_mini_FinalTotal",
        "o1-mini_FinalTotal",
    ],


    "GPT-5.4-Thinking": [

        "GPT54Thinking_Score",
        "GPT-5.4 Thinking_Score",
        "GPT-5.4-Thinking_Score",

        "GPT54Thinking",
        "GPT-5.4 Thinking",

        "GPT54Thinking_RawTotal",
        "GPT-5.4 Thinking_RawTotal",

        "GPT54Thinking_FinalTotal",
        "GPT-5.4 Thinking_FinalTotal",
    ],


    "Gemini-3-Flash-Thinking": [

        "Gemini3FlashThinking_Score",

        "Gemini 3 Flash Thinking_Score",

        "Gemini 3 Flash (Thinking Mode)_Score",

        "Gemini-3-Flash-Thinking_Score",

        "Gemini3FlashThinking",

        "Gemini 3 Flash (Thinking Mode)",

        "Gemini3FlashThinking_RawTotal",

        "Gemini 3 Flash (Thinking Mode)_RawTotal",

        "Gemini3FlashThinking_FinalTotal",

        "Gemini 3 Flash (Thinking Mode)_FinalTotal",
    ],
}


# =============================================================================
# STEP 8. Normalize column name
# =============================================================================

def normalize_name(
    value
):

    if value is None:

        return ""


    value = (
        str(
            value
        )
        .strip()
        .lower()
    )


    value = re.sub(
        r"[^a-z0-9]+",
        "",
        value,
    )


    return value


# =============================================================================
# STEP 9. Find column from aliases
# =============================================================================

def find_column(
    dataframe,
    aliases,
):

    normalized_columns = {

        normalize_name(
            column
        ):
            column

        for column
        in dataframe.columns
    }


    for alias in aliases:

        normalized_alias = (
            normalize_name(
                alias
            )
        )


        if (
            normalized_alias
            in normalized_columns
        ):

            return (
                normalized_columns[
                    normalized_alias
                ]
            )


    return None


# =============================================================================
# STEP 10. Find model columns
# =============================================================================

def find_model_columns(
    dataframe
):

    result = {}


    for model in MODELS:

        column = find_column(
            dataframe,
            MODEL_ALIASES[
                model
            ],
        )


        if column is None:

            return None


        result[
            model
        ] = column


    return result


# =============================================================================
# STEP 11. Standardize annual table
# =============================================================================

def standardize_annual_table(
    dataframe,
    year_column,
    model_columns,
):

    result = dataframe[
        [
            year_column,
            *model_columns.values(),
        ]
    ].copy()


    result.columns = [
        "Year",
        *MODELS,
    ]


    result[
        "Year"
    ] = pd.to_numeric(
        result[
            "Year"
        ],
        errors="coerce",
    )


    for model in MODELS:

        result[
            model
        ] = pd.to_numeric(
            result[
                model
            ],
            errors="coerce",
        )


    return result


# =============================================================================
# STEP 12. Extract annual candidate from one sheet
# =============================================================================

def extract_candidate_from_sheet(
    dataframe,
    sheet_name,
):

    if dataframe is None:

        return None


    if dataframe.empty:

        return None


    year_column = find_column(
        dataframe,
        YEAR_ALIASES,
    )


    if year_column is None:

        return None


    model_columns = find_model_columns(
        dataframe
    )


    if model_columns is None:

        return None


    work = dataframe.copy()


    work[
        year_column
    ] = pd.to_numeric(
        work[
            year_column
        ],
        errors="coerce",
    )


    work = work.loc[

        work[
            year_column
        ].between(
            2015,
            2024,
            inclusive="both",
        )

    ].copy()


    if work.empty:

        return None


    # =========================================================================
    # Priority 300:
    # Explicit Calculated_Year_Total
    # =========================================================================

    row_type_column = find_column(
        work,
        ROW_TYPE_ALIASES,
    )


    if row_type_column is not None:


        row_type_text = (
            work[
                row_type_column
            ]
            .astype(str)
            .str.lower()
            .str.replace(
                r"[^a-z0-9]+",
                "",
                regex=True,
            )
        )


        year_total_mask = (
            row_type_text
            .str.contains(
                "calculatedyeartotal",
                na=False,
            )
        )


        year_total_rows = (
            work.loc[
                year_total_mask
            ]
            .copy()
        )


        if not year_total_rows.empty:


            annual = (
                standardize_annual_table(
                    dataframe=
                        year_total_rows,

                    year_column=
                        year_column,

                    model_columns=
                        model_columns,
                )
            )


            return {

                "sheet_name":
                    sheet_name,

                "method":
                    "Explicit Calculated_Year_Total rows",

                "priority":
                    300,

                "data":
                    annual,
            }


    # =========================================================================
    # Priority 200:
    # Exactly one annual row per year
    # =========================================================================

    unique_years = sorted(

        work[
            year_column
        ]
        .dropna()
        .astype(int)
        .unique()
        .tolist()
    )


    year_counts = (
        work[
            year_column
        ]
        .value_counts()
    )


    if (
        unique_years
        == YEARS
        and
        year_counts.max()
        == 1
    ):


        annual = (
            standardize_annual_table(
                dataframe=
                    work,

                year_column=
                    year_column,

                model_columns=
                    model_columns,
            )
        )


        return {

            "sheet_name":
                sheet_name,

            "method":
                "Exactly one annual row per year",

            "priority":
                200,

            "data":
                annual,
        }


    # =========================================================================
    # Priority 100:
    # Item-level fallback → sum within year
    # =========================================================================

    temp = (
        standardize_annual_table(
            dataframe=
                work,

            year_column=
                year_column,

            model_columns=
                model_columns,
        )
    )


    temp = temp.loc[

        temp[
            MODELS
        ]
        .notna()
        .any(
            axis=1
        )

    ].copy()


    if temp.empty:

        return None


    annual = (
        temp
        .groupby(
            "Year",
            as_index=False,
        )[
            MODELS
        ]
        .sum(
            min_count=1
        )
    )


    return {

        "sheet_name":
            sheet_name,

        "method":
            "Item-level rows summed within year",

        "priority":
            100,

        "data":
            annual,
    }


# =============================================================================
# STEP 13. Validate annual candidate
# =============================================================================

def validate_annual_candidate(
    subject,
    annual,
):

    annual = annual.copy()


    annual[
        "Year"
    ] = pd.to_numeric(
        annual[
            "Year"
        ],
        errors="coerce",
    )


    annual = annual.loc[

        annual[
            "Year"
        ].isin(
            YEARS
        )

    ].copy()


    annual = (
        annual
        .sort_values(
            "Year"
        )
        .reset_index(
            drop=True
        )
    )


    # -------------------------------------------------------------------------
    # Exactly 10 rows
    # -------------------------------------------------------------------------

    if len(
        annual
    ) != 10:


        raise ValueError(

            f"[{subject}] "
            f"年度列數不是 10。"
            f"目前={len(annual)}"
        )


    # -------------------------------------------------------------------------
    # Exactly 2015–2024
    # -------------------------------------------------------------------------

    actual_years = (
        annual[
            "Year"
        ]
        .astype(int)
        .tolist()
    )


    if actual_years != YEARS:


        raise ValueError(

            f"[{subject}] "
            f"年份不是完整 2015–2024："
            f"{actual_years}"
        )


    # -------------------------------------------------------------------------
    # Duplicate year
    # -------------------------------------------------------------------------

    duplicate_years = (
        annual.loc[
            annual[
                "Year"
            ].duplicated(
                keep=False
            ),
            "Year",
        ]
        .tolist()
    )


    if duplicate_years:


        raise ValueError(

            f"[{subject}] "
            f"存在 duplicate year："
            f"{duplicate_years}"
        )


    # -------------------------------------------------------------------------
    # Numeric model values
    # -------------------------------------------------------------------------

    for model in MODELS:


        annual[
            model
        ] = pd.to_numeric(
            annual[
                model
            ],
            errors="coerce",
        )


    # -------------------------------------------------------------------------
    # Missing
    # -------------------------------------------------------------------------

    missing_counts = (
        annual[
            MODELS
        ]
        .isna()
        .sum()
    )


    if (
        missing_counts.sum()
        > 0
    ):


        raise ValueError(

            f"[{subject}] "
            f"年度模型分數有缺值：\n"
            f"{missing_counts}"
        )


    annual[
        "Year"
    ] = (
        annual[
            "Year"
        ]
        .astype(int)
    )


    # -------------------------------------------------------------------------
    # Range check
    # -------------------------------------------------------------------------

    for row_index in annual.index:


        year = int(
            annual.loc[
                row_index,
                "Year",
            ]
        )


        if (
            year
            not in OBJECTIVE_MAX[
                subject
            ]
        ):


            raise KeyError(

                f"[{subject}] "
                f"找不到 {year} 年 objective maximum。"
            )


        maximum = float(
            OBJECTIVE_MAX[
                subject
            ][
                year
            ]
        )


        for model in MODELS:


            score = float(
                annual.loc[
                    row_index,
                    model,
                ]
            )


            if score < -1e-9:


                raise ValueError(

                    f"[{subject}] "
                    f"{year} "
                    f"{model} "
                    f"score < 0："
                    f"{score}"
                )


            if (
                score
                >
                maximum
                + 1e-8
            ):


                raise ValueError(

                    f"[{subject}] "
                    f"{year} "
                    f"{model} "
                    f"score={score} "
                    f"> objective maximum={maximum}"
                )


    return annual


# =============================================================================
# STEP 14. Compare annual tables
# =============================================================================

def annual_tables_equal(
    table_1,
    table_2,
):

    table_1 = (
        table_1
        .sort_values(
            "Year"
        )
        .reset_index(
            drop=True
        )
    )


    table_2 = (
        table_2
        .sort_values(
            "Year"
        )
        .reset_index(
            drop=True
        )
    )


    if (
        table_1[
            "Year"
        ].tolist()
        !=
        table_2[
            "Year"
        ].tolist()
    ):

        return False


    return np.allclose(
        table_1[
            MODELS
        ].to_numpy(
            dtype=float
        ),
        table_2[
            MODELS
        ].to_numpy(
            dtype=float
        ),
        rtol=0,
        atol=1e-9,
        equal_nan=False,
    )


# =============================================================================
# STEP 15. Read one subject workbook
# =============================================================================

def read_subject_file(
    subject,
    file_path,
):

    print(
        "\n"
        + "=" * 100
    )


    print(
        f"處理科目：{subject}"
    )


    print(
        f"檔案：{file_path}"
    )


    if not file_path.exists():


        raise FileNotFoundError(

            f"\n找不到檔案：\n"
            f"{file_path}"
        )


    excel_file = pd.ExcelFile(
        file_path
    )


    print(
        "Sheets：",
        excel_file.sheet_names
    )


    candidates = []


    # -------------------------------------------------------------------------
    # Scan all sheets
    # -------------------------------------------------------------------------

    for sheet_name in excel_file.sheet_names:


        try:


            dataframe = pd.read_excel(
                excel_file,
                sheet_name=
                    sheet_name,
            )


        except Exception as error:


            print(

                f"略過 Sheet："
                f"{sheet_name}"
                f"；原因："
                f"{error}"
            )


            continue


        candidate = (
            extract_candidate_from_sheet(
                dataframe=
                    dataframe,
                sheet_name=
                    sheet_name,
            )
        )


        if candidate is not None:


            candidates.append(
                candidate
            )


    if not candidates:


        raise ValueError(

            f"\n[{subject}] "
            f"所有 sheets 都找不到可辨識的年度分數資料。"
        )


    # -------------------------------------------------------------------------
    # Validate all candidates
    # -------------------------------------------------------------------------

    qc_records = []

    valid_candidates = []


    for candidate in candidates:


        try:


            validated_data = (
                validate_annual_candidate(
                    subject=
                        subject,
                    annual=
                        candidate[
                            "data"
                        ],
                )
            )


            qc_records.append({

                "Subject":
                    subject,

                "File":
                    (
                        "supplementary-materials/"
                        "analysis-ready-data/"
                        f"{file_path.name}"
                    ),

                "Sheet":
                    candidate[
                        "sheet_name"
                    ],

                "Method":
                    candidate[
                        "method"
                    ],

                "Priority":
                    candidate[
                        "priority"
                    ],

                "Valid":
                    True,

                "Annual_Rows":
                    len(
                        validated_data
                    ),

                "Year_Min":
                    int(
                        validated_data[
                            "Year"
                        ].min()
                    ),

                "Year_Max":
                    int(
                        validated_data[
                            "Year"
                        ].max()
                    ),

                "Missing_N":
                    int(
                        validated_data[
                            MODELS
                        ]
                        .isna()
                        .sum()
                        .sum()
                    ),

                "Error":
                    "",
            })


            candidate[
                "validated_data"
            ] = validated_data


            valid_candidates.append(
                candidate
            )


        except Exception as error:


            qc_records.append({

                "Subject":
                    subject,

                "File":
                    (
                        "supplementary-materials/"
                        "analysis-ready-data/"
                        f"{file_path.name}"
                    ),

                "Sheet":
                    candidate[
                        "sheet_name"
                    ],

                "Method":
                    candidate[
                        "method"
                    ],

                "Priority":
                    candidate[
                        "priority"
                    ],

                "Valid":
                    False,

                "Annual_Rows":
                    np.nan,

                "Year_Min":
                    np.nan,

                "Year_Max":
                    np.nan,

                "Missing_N":
                    np.nan,

                "Error":
                    str(
                        error
                    ),
            })


    if not valid_candidates:


        qc_dataframe = pd.DataFrame(
            qc_records
        )


        print(
            "\n候選資料 QC："
        )


        print(
            qc_dataframe.to_string(
                index=False
            )
        )


        raise ValueError(

            f"\n[{subject}] "
            f"沒有任何有效年度資料。"
        )


    # -------------------------------------------------------------------------
    # Highest priority
    # -------------------------------------------------------------------------

    highest_priority = max(

        candidate[
            "priority"
        ]

        for candidate
        in valid_candidates
    )


    top_candidates = [

        candidate

        for candidate
        in valid_candidates

        if (
            candidate[
                "priority"
            ]
            == highest_priority
        )
    ]


    # -------------------------------------------------------------------------
    # Multiple top candidates must agree
    # -------------------------------------------------------------------------

    if len(
        top_candidates
    ) > 1:


        reference = (
            top_candidates[
                0
            ][
                "validated_data"
            ]
        )


        conflicting = []


        for candidate in top_candidates[
            1:
        ]:


            same = annual_tables_equal(
                reference,
                candidate[
                    "validated_data"
                ],
            )


            if not same:


                conflicting.append(
                    candidate[
                        "sheet_name"
                    ]
                )


        if conflicting:


            raise ValueError(

                f"\n[{subject}] "
                f"發現多個同優先級年度資料來源，"
                f"但結果不一致。\n"
                f"程式不自動猜測。\n"
                f"Conflicting sheets："
                f"{conflicting}"
            )


    selected = (
        top_candidates[
            0
        ]
    )


    selected_sheet = (
        selected[
            "sheet_name"
        ]
    )


    selected_method = (
        selected[
            "method"
        ]
    )


    annual_raw = (
        selected[
            "validated_data"
        ]
        .copy()
    )


    # -------------------------------------------------------------------------
    # Candidate QC table
    # -------------------------------------------------------------------------

    qc_dataframe = pd.DataFrame(
        qc_records
    )


    qc_dataframe[
        "Selected"
    ] = False


    selected_mask = (

        (
            qc_dataframe[
                "Sheet"
            ]
            == selected_sheet
        )

        &

        (
            qc_dataframe[
                "Method"
            ]
            == selected_method
        )

        &

        (
            qc_dataframe[
                "Priority"
            ]
            == highest_priority
        )
    )


    if selected_mask.any():


        selected_index = (
            qc_dataframe.loc[
                selected_mask
            ]
            .index[
                0
            ]
        )


        qc_dataframe.loc[
            selected_index,
            "Selected",
        ] = True


    print(
        f"\n選定 Sheet："
        f"{selected_sheet}"
    )


    print(
        f"抽取方法："
        f"{selected_method}"
    )


    print(
        "\nAnnual RAW scores："
    )


    print(
        annual_raw.to_string(
            index=False
        )
    )


    return {

        "annual_raw":
            annual_raw,

        "selected_sheet":
            selected_sheet,

        "selected_method":
            selected_method,

        "candidate_qc":
            qc_dataframe,
    }


# =============================================================================
# STEP 16. Raw → normalized annual accuracy (%)
# =============================================================================

def raw_to_accuracy(
    subject,
    annual_raw,
):

    annual_accuracy = (
        annual_raw
        .copy()
    )


    # Ensure model score columns are floating point before assigning
    # normalized percentage values. This avoids pandas dtype warnings
    # without changing any analytical calculation.
    annual_accuracy[
        MODELS
    ] = annual_accuracy[
        MODELS
    ].astype(float)


    for row_index in annual_accuracy.index:


        year = int(
            annual_accuracy.loc[
                row_index,
                "Year",
            ]
        )


        maximum = float(
            OBJECTIVE_MAX[
                subject
            ][
                year
            ]
        )


        for model in MODELS:


            raw_score = float(
                annual_accuracy.loc[
                    row_index,
                    model,
                ]
            )


            annual_accuracy.loc[
                row_index,
                model,
            ] = (
                raw_score
                /
                maximum
                *
                100.0
            )


    return annual_accuracy


# =============================================================================
# STEP 17. p-value display
# =============================================================================

def format_p(
    value
):

    if (
        value is None
        or
        pd.isna(
            value
        )
    ):

        return ""


    if value < 0.001:

        return "<0.001"


    return (
        f"{value:.3f}"
    )


# =============================================================================
# STEP 19. Table 4 - RAW descriptive statistics
# =============================================================================

def run_descriptive_raw(
    annual_raw_all
):

    rows = []


    for subject in INPUT_FILES.keys():


        subject_data = (
            annual_raw_all.loc[
                annual_raw_all[
                    "Subject"
                ]
                == subject
            ]
            .sort_values(
                "Year"
            )
        )


        for model in MODELS:


            values = (
                subject_data[
                    model
                ]
                .astype(float)
            )


            rows.append({

                "Subject":
                    subject,

                "Model":
                    model,

                "Min":
                    float(
                        values.min()
                    ),

                "Max":
                    float(
                        values.max()
                    ),

                "Mean":
                    float(
                        values.mean()
                    ),

                "SD":
                    float(
                        values.std(
                            ddof=1
                        )
                    ),

                "Median":
                    float(
                        values.median()
                    ),

                "N":
                    int(
                        len(
                            values
                        )
                    ),
            })


    return pd.DataFrame(
        rows
    )


# =============================================================================
# STEP 20. Normalized accuracy descriptive statistics
# =============================================================================

def run_descriptive_accuracy(
    annual_accuracy_all
):

    rows = []


    for subject in INPUT_FILES.keys():


        subject_data = (
            annual_accuracy_all.loc[
                annual_accuracy_all[
                    "Subject"
                ]
                == subject
            ]
            .sort_values(
                "Year"
            )
        )


        for model in MODELS:


            values = (
                subject_data[
                    model
                ]
                .astype(float)
            )


            rows.append({

                "Subject":
                    subject,

                "Model":
                    model,

                "Min_Accuracy_Percent":
                    float(
                        values.min()
                    ),

                "Max_Accuracy_Percent":
                    float(
                        values.max()
                    ),

                "Mean_Accuracy_Percent":
                    float(
                        values.mean()
                    ),

                "SD_Accuracy_Percent":
                    float(
                        values.std(
                            ddof=1
                        )
                    ),

                "Median_Accuracy_Percent":
                    float(
                        values.median()
                    ),

                "N":
                    int(
                        len(
                            values
                        )
                    ),
            })


    return pd.DataFrame(
        rows
    )


# =============================================================================
# STEP 21. Table 2 - Shapiro-Wilk legacy diagnostic
#
# NO Sig column.
# =============================================================================

def run_shapiro_legacy(
    annual_raw_all
):

    rows = []


    for subject in INPUT_FILES.keys():


        subject_data = (
            annual_raw_all.loc[
                annual_raw_all[
                    "Subject"
                ]
                == subject
            ]
        )


        for model in MODELS:


            values = (
                subject_data[
                    model
                ]
                .astype(float)
                .to_numpy()
            )


            result = shapiro(
                values
            )


            rows.append({

                "Subject":
                    subject,

                "Model":
                    model,

                "W":
                    float(
                        result.statistic
                    ),

                "p":
                    float(
                        result.pvalue
                    ),

                "p_Display":
                    format_p(
                        result.pvalue
                    ),

                "Significant_p_lt_0_05":
                    bool(
                        result.pvalue
                        < ALPHA
                    ),
            })


    return pd.DataFrame(
        rows
    )


# =============================================================================
# STEP 22. Table 5 - Levene legacy diagnostic
#
# NO Sig column.
# =============================================================================

def run_levene_legacy(
    annual_raw_all
):

    rows = []


    for subject in INPUT_FILES.keys():


        subject_data = (
            annual_raw_all.loc[
                annual_raw_all[
                    "Subject"
                ]
                == subject
            ]
        )


        arrays = [

            subject_data[
                model
            ]
            .astype(float)
            .to_numpy()

            for model
            in MODELS
        ]


        result = levene(
            *arrays,
            center="mean",
        )


        total_n = sum(
            len(
                array
            )
            for array
            in arrays
        )


        df1 = (
            len(
                MODELS
            )
            - 1
        )


        df2 = (
            total_n
            - len(
                MODELS
            )
        )


        rows.append({

            "Subject":
                subject,

            "Levene_Statistic":
                float(
                    result.statistic
                ),

            "df1":
                int(
                    df1
                ),

            "df2":
                int(
                    df2
                ),

            "p":
                float(
                    result.pvalue
                ),

            "p_Display":
                format_p(
                    result.pvalue
                ),

            "Significant_p_lt_0_05":
                bool(
                    result.pvalue
                    < ALPHA
                ),
        })


    return pd.DataFrame(
        rows
    )


# =============================================================================
# STEP 23. Table 6 - Friedman
#
# NO Sig column.
# =============================================================================

def run_friedman(
    annual_accuracy_all
):

    rows = []


    for subject in INPUT_FILES.keys():


        subject_data = (
            annual_accuracy_all.loc[
                annual_accuracy_all[
                    "Subject"
                ]
                == subject
            ]
            .sort_values(
                "Year"
            )
            .reset_index(
                drop=True
            )
        )


        arrays = [

            subject_data[
                model
            ]
            .astype(float)
            .to_numpy()

            for model
            in MODELS
        ]


        result = friedmanchisquare(
            *arrays
        )


        n_blocks = len(
            subject_data
        )


        k_models = len(
            MODELS
        )


        degrees_of_freedom = (
            k_models
            - 1
        )


        kendalls_w = (
            float(
                result.statistic
            )
            /
            (
                n_blocks
                *
                (
                    k_models
                    - 1
                )
            )
        )


        rows.append({

            "Subject":
                subject,

            "N_Years":
                int(
                    n_blocks
                ),

            "N_Models":
                int(
                    k_models
                ),

            "Friedman_ChiSquare":
                float(
                    result.statistic
                ),

            "df":
                int(
                    degrees_of_freedom
                ),

            "p":
                float(
                    result.pvalue
                ),

            "p_Display":
                format_p(
                    result.pvalue
                ),

            "Kendalls_W":
                float(
                    kendalls_w
                ),

            "Omnibus_Significant":
                bool(
                    result.pvalue
                    < ALPHA
                ),
        })


    return pd.DataFrame(
        rows
    )


# =============================================================================
# STEP 24. Table 7 - Friedman mean ranks
# =============================================================================

def run_friedman_mean_ranks(
    annual_accuracy_all
):

    rows = []


    for subject in INPUT_FILES.keys():


        subject_data = (
            annual_accuracy_all.loc[
                annual_accuracy_all[
                    "Subject"
                ]
                == subject
            ]
            .sort_values(
                "Year"
            )
            .reset_index(
                drop=True
            )
        )


        scores = (
            subject_data[
                MODELS
            ]
            .astype(float)
        )


        # ---------------------------------------------------------------------
        # Within-year ranks:
        #
        # Rank 1 = lowest
        # Rank 5 = highest
        #
        # Ties = average rank
        # ---------------------------------------------------------------------

        ranks = (
            scores
            .rank(
                axis=1,
                method="average",
                ascending=True,
            )
        )


        for model in MODELS:


            rows.append({

                "Subject":
                    subject,

                "Model":
                    model,

                "Mean_Rank":
                    float(
                        ranks[
                            model
                        ].mean()
                    ),

                "Median_Rank":
                    float(
                        ranks[
                            model
                        ].median()
                    ),

                "Mean_Accuracy_Percent":
                    float(
                        scores[
                            model
                        ].mean()
                    ),
            })


    return pd.DataFrame(
        rows
    )


# =============================================================================
# STEP 25. Table 3 - Wilcoxon + Holm
#
# NO Sig column.
#
# Word:
# significant star will be attached to Holm p.
# =============================================================================

def run_wilcoxon_holm(
    annual_accuracy_all,
    friedman_results,
):

    result_list = []


    for subject in INPUT_FILES.keys():


        subject_data = (
            annual_accuracy_all.loc[
                annual_accuracy_all[
                    "Subject"
                ]
                == subject
            ]
            .sort_values(
                "Year"
            )
            .reset_index(
                drop=True
            )
        )


        omnibus_row = (
            friedman_results.loc[
                friedman_results[
                    "Subject"
                ]
                == subject
            ]
            .iloc[
                0
            ]
        )


        omnibus_p = float(
            omnibus_row[
                "p"
            ]
        )


        omnibus_significant = bool(
            omnibus_p
            < ALPHA
        )


        subject_rows = []


        for (
            model_1,
            model_2,
        ) in itertools.combinations(
            MODELS,
            2,
        ):


            x = (
                subject_data[
                    model_1
                ]
                .astype(float)
                .to_numpy()
            )


            y = (
                subject_data[
                    model_2
                ]
                .astype(float)
                .to_numpy()
            )


            difference = (
                x
                - y
            )


            model_1_wins = int(
                np.sum(
                    difference
                    > 0
                )
            )


            model_2_wins = int(
                np.sum(
                    difference
                    < 0
                )
            )


            ties = int(
                np.sum(
                    np.isclose(
                        difference,
                        0,
                    )
                )
            )


            nonzero_pairs = int(
                np.sum(
                    ~np.isclose(
                        difference,
                        0,
                    )
                )
            )


            if nonzero_pairs == 0:


                wilcoxon_w = 0.0

                p_raw = 1.0

                method_used = (
                    "All paired differences = 0"
                )


            else:


                wilcoxon_result = (
                    wilcoxon(
                        x,
                        y,
                        alternative="two-sided",
                        zero_method="wilcox",
                        correction=False,
                        method="auto",
                    )
                )


                wilcoxon_w = float(
                    wilcoxon_result.statistic
                )


                p_raw = float(
                    wilcoxon_result.pvalue
                )


                method_used = (
                    "Two-sided Wilcoxon signed-rank; "
                    "zero_method='wilcox'; "
                    "method='auto'"
                )


            subject_rows.append({

                "Subject":
                    subject,

                "Model_1":
                    model_1,

                "Model_2":
                    model_2,

                "Mean_Accuracy_Model_1":
                    float(
                        np.mean(
                            x
                        )
                    ),

                "Mean_Accuracy_Model_2":
                    float(
                        np.mean(
                            y
                        )
                    ),

                "Mean_Difference_pp":
                    float(
                        np.mean(
                            difference
                        )
                    ),

                "Median_Difference_pp":
                    float(
                        np.median(
                            difference
                        )
                    ),

                "Model_1_Wins":
                    model_1_wins,

                "Model_2_Wins":
                    model_2_wins,

                "Ties":
                    ties,

                "N_Nonzero_Pairs":
                    nonzero_pairs,

                "Wilcoxon_W":
                    wilcoxon_w,

                "p_raw":
                    p_raw,

                "Test_Method":
                    method_used,

                "Omnibus_Friedman_p":
                    omnibus_p,

                "Omnibus_Significant":
                    omnibus_significant,
            })


        subject_results = pd.DataFrame(
            subject_rows
        )


        # ---------------------------------------------------------------------
        # Holm correction within subject
        # ---------------------------------------------------------------------

        (
            reject,
            p_holm,
            _,
            _,
        ) = multipletests(

            subject_results[
                "p_raw"
            ]
            .to_numpy(),

            alpha=ALPHA,

            method="holm",
        )


        subject_results[
            "p_Holm"
        ] = p_holm


        subject_results[
            "p_raw_Display"
        ] = (
            subject_results[
                "p_raw"
            ]
            .apply(
                format_p
            )
        )


        subject_results[
            "p_Holm_Display"
        ] = (
            subject_results[
                "p_Holm"
            ]
            .apply(
                format_p
            )
        )


        subject_results[
            "Significant_Holm"
        ] = reject


        subject_results[
            "Interpret_as_Significant"
        ] = (
            subject_results[
                "Significant_Holm"
            ]
            &
            subject_results[
                "Omnibus_Significant"
            ]
        )


        # ---------------------------------------------------------------------
        # Descriptive direction
        # ---------------------------------------------------------------------

        def get_direction(
            row
        ):


            mean_difference = (
                row[
                    "Mean_Difference_pp"
                ]
            )


            if np.isclose(
                mean_difference,
                0,
            ):


                return (
                    "No mean difference"
                )


            if mean_difference > 0:


                return (
                    f"{row['Model_1']} > "
                    f"{row['Model_2']}"
                )


            return (
                f"{row['Model_2']} > "
                f"{row['Model_1']}"
            )


        subject_results[
            "Direction"
        ] = (
            subject_results
            .apply(
                get_direction,
                axis=1,
            )
        )


        result_list.append(
            subject_results
        )


    return pd.concat(
        result_list,
        ignore_index=True,
    )


# =============================================================================
# STEP 26. Excel formatting
# =============================================================================

def format_excel_workbook(
    excel_path
):

    workbook = load_workbook(
        excel_path
    )


    header_fill = PatternFill(
        fill_type="solid",
        fgColor="D9EAF7",
    )


    significant_fill = PatternFill(
        fill_type="solid",
        fgColor="FFF2CC",
    )


    selected_fill = PatternFill(
        fill_type="solid",
        fgColor="E2F0D9",
    )


    invalid_fill = PatternFill(
        fill_type="solid",
        fgColor="F4CCCC",
    )


    thin_side = Side(
        style="thin",
        color="D9D9D9",
    )


    border = Border(
        left=thin_side,
        right=thin_side,
        top=thin_side,
        bottom=thin_side,
    )


    for worksheet in workbook.worksheets:


        worksheet.freeze_panes = (
            "A2"
        )


        # ---------------------------------------------------------------------
        # Header
        # ---------------------------------------------------------------------

        for cell in worksheet[
            1
        ]:


            cell.font = Font(
                bold=True
            )


            cell.fill = (
                header_fill
            )


            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )


            cell.border = (
                border
            )


        # ---------------------------------------------------------------------
        # Body
        # ---------------------------------------------------------------------

        for row in worksheet.iter_rows():


            for cell in row:


                cell.border = (
                    border
                )


                cell.alignment = Alignment(
                    vertical="top",
                    wrap_text=True,
                )


        # ---------------------------------------------------------------------
        # Column widths
        # ---------------------------------------------------------------------

        for column_cells in worksheet.columns:


            column_letter = (
                get_column_letter(
                    column_cells[
                        0
                    ].column
                )
            )


            maximum_length = 0


            for cell in column_cells:


                if cell.value is None:

                    continue


                maximum_length = max(
                    maximum_length,
                    len(
                        str(
                            cell.value
                        )
                    ),
                )


            worksheet.column_dimensions[
                column_letter
            ].width = min(
                max(
                    maximum_length
                    + 2,
                    10,
                ),
                42,
            )


        header_map = {

            cell.value:
                cell.column

            for cell
            in worksheet[
                1
            ]
        }


        # ---------------------------------------------------------------------
        # Selected source
        # ---------------------------------------------------------------------

        if (
            "Selected"
            in header_map
        ):


            selected_column = (
                header_map[
                    "Selected"
                ]
            )


            for row_number in range(
                2,
                worksheet.max_row
                + 1,
            ):


                value = worksheet.cell(
                    row=row_number,
                    column=selected_column,
                ).value


                if value in [
                    True,
                    "TRUE",
                    1,
                ]:


                    for cell in worksheet[
                        row_number
                    ]:


                        cell.fill = (
                            selected_fill
                        )


        # ---------------------------------------------------------------------
        # Invalid candidate
        # ---------------------------------------------------------------------

        if (
            "Valid"
            in header_map
        ):


            valid_column = (
                header_map[
                    "Valid"
                ]
            )


            for row_number in range(
                2,
                worksheet.max_row
                + 1,
            ):


                value = worksheet.cell(
                    row=row_number,
                    column=valid_column,
                ).value


                if value in [
                    False,
                    "FALSE",
                    0,
                ]:


                    for cell in worksheet[
                        row_number
                    ]:


                        cell.fill = (
                            invalid_fill
                        )


        # ---------------------------------------------------------------------
        # Significant rows
        # ---------------------------------------------------------------------

        for significance_column in [

            "Significant_p_lt_0_05",

            "Omnibus_Significant",

            "Interpret_as_Significant",

        ]:


            if (
                significance_column
                not in header_map
            ):

                continue


            column_number = (
                header_map[
                    significance_column
                ]
            )


            for row_number in range(
                2,
                worksheet.max_row
                + 1,
            ):


                value = worksheet.cell(
                    row=row_number,
                    column=column_number,
                ).value


                if value in [
                    True,
                    "TRUE",
                    1,
                ]:


                    for cell in worksheet[
                        row_number
                    ]:


                        cell.fill = (
                            significant_fill
                        )


    workbook.save(
        excel_path
    )


# =============================================================================
# STEP 27. Word helper - repeat table header
# =============================================================================

def set_repeat_table_header(
    row
):

    table_row = (
        row._tr
    )


    table_properties = (
        table_row
        .get_or_add_trPr()
    )


    table_header = OxmlElement(
        "w:tblHeader"
    )


    table_header.set(
        qn(
            "w:val"
        ),
        "true",
    )


    table_properties.append(
        table_header
    )


# =============================================================================
# STEP 28. Word helper - fixed layout
# =============================================================================

def set_fixed_table_layout(
    table
):

    table_properties = (
        table._tbl.tblPr
    )


    layout = OxmlElement(
        "w:tblLayout"
    )


    layout.set(
        qn(
            "w:type"
        ),
        "fixed",
    )


    table_properties.append(
        layout
    )


# =============================================================================
# STEP 29. Word helper - normal cell
# =============================================================================

def set_word_cell(
    cell,
    value,
    font_size=8,
    bold=False,
):

    cell.text = ""


    paragraph = (
        cell.paragraphs[
            0
        ]
    )


    paragraph.alignment = (
        WD_ALIGN_PARAGRAPH.LEFT
    )


    run = paragraph.add_run(
        str(
            value
        )
    )


    run.font.name = (
        "Times New Roman"
    )


    run.font.size = Pt(
        font_size
    )


    run.bold = (
        bold
    )


    cell.vertical_alignment = (
        WD_CELL_VERTICAL_ALIGNMENT.CENTER
    )


# =============================================================================
# STEP 30. Word helper - p-value + superscript *
#
# Key modification:
#
# Normal:
#     0.0195
#
# Significant:
#     0.0195*
#
# The "*" is a separate Word run with superscript=True.
# =============================================================================

def set_word_pvalue_cell(
    cell,
    p_text,
    significant,
    font_size=8,
):

    cell.text = ""


    paragraph = (
        cell.paragraphs[
            0
        ]
    )


    paragraph.alignment = (
        WD_ALIGN_PARAGRAPH.LEFT
    )


    # -------------------------------------------------------------------------
    # p-value
    # -------------------------------------------------------------------------

    p_run = paragraph.add_run(
        str(
            p_text
        )
    )


    p_run.font.name = (
        "Times New Roman"
    )


    p_run.font.size = Pt(
        font_size
    )


    # -------------------------------------------------------------------------
    # Superscript significance star
    # -------------------------------------------------------------------------

    if bool(
        significant
    ):


        star_run = paragraph.add_run(
            "*"
        )


        star_run.font.name = (
            "Times New Roman"
        )


        star_run.font.size = Pt(
            font_size
        )


        star_run.font.superscript = (
            True
        )


    cell.vertical_alignment = (
        WD_CELL_VERTICAL_ALIGNMENT.CENTER
    )


# =============================================================================
# STEP 31. Word caption
# =============================================================================

def add_word_caption(
    document,
    text
):

    paragraph = (
        document.add_paragraph()
    )


    paragraph.paragraph_format.space_before = Pt(
        8
    )


    paragraph.paragraph_format.space_after = Pt(
        4
    )


    run = paragraph.add_run(
        text
    )


    run.bold = True


    run.font.name = (
        "Times New Roman"
    )


    run.font.size = Pt(
        10
    )


# =============================================================================
# STEP 32. Word table note
# =============================================================================

def add_word_note(
    document,
    text
):

    paragraph = (
        document.add_paragraph()
    )


    paragraph.paragraph_format.space_before = Pt(
        3
    )


    paragraph.paragraph_format.space_after = Pt(
        8
    )


    run = paragraph.add_run(
        "Note: "
        + text
    )


    run.font.name = (
        "Times New Roman"
    )


    run.font.size = Pt(
        8
    )


# =============================================================================
# STEP 33. Word value formatter
# =============================================================================

def word_value(
    value,
    decimals=3,
):

    if (
        value is None
        or
        pd.isna(
            value
        )
    ):

        return ""


    if isinstance(
        value,
        (
            bool,
            np.bool_,
        ),
    ):


        return (
            "Yes"
            if value
            else "No"
        )


    if isinstance(
        value,
        (
            int,
            np.integer,
        ),
    ):


        return str(
            int(
                value
            )
        )


    if isinstance(
        value,
        (
            float,
            np.floating,
        ),
    ):


        return (
            f"{float(value):.{decimals}f}"
        )


    return str(
        value
    )


# =============================================================================
# STEP 34. Generic dataframe → Word table
#
# star_rules example:
#
# {
#     "p_Display": "Omnibus_Significant"
# }
#
# means:
#
# when writing p_Display cell,
# add superscript * if corresponding
# Omnibus_Significant == True.
# =============================================================================

def add_dataframe_to_word(
    document,
    dataframe,
    display_columns,
    column_labels,
    decimals_map=None,
    font_size=7.5,
    star_rules=None,
):

    if decimals_map is None:

        decimals_map = {}


    if star_rules is None:

        star_rules = {}


    table = document.add_table(
        rows=1,
        cols=len(
            display_columns
        ),
    )


    table.alignment = (
        WD_TABLE_ALIGNMENT.CENTER
    )


    table.style = (
        "Table Grid"
    )


    table.autofit = (
        False
    )


    set_fixed_table_layout(
        table
    )


    # -------------------------------------------------------------------------
    # Header
    # -------------------------------------------------------------------------

    header_cells = (
        table.rows[
            0
        ].cells
    )


    for index, label in enumerate(
        column_labels
    ):


        set_word_cell(
            header_cells[
                index
            ],
            label,
            font_size=
                font_size,
            bold=True,
        )


    set_repeat_table_header(
        table.rows[
            0
        ]
    )


    # -------------------------------------------------------------------------
    # Body
    # -------------------------------------------------------------------------

    for _, record in dataframe.iterrows():


        cells = (
            table.add_row()
            .cells
        )


        for index, column in enumerate(
            display_columns
        ):


            value = (
                record[
                    column
                ]
            )


            # -----------------------------------------------------------------
            # p-value cell requiring superscript star
            # -----------------------------------------------------------------

            if column in star_rules:


                significance_column = (
                    star_rules[
                        column
                    ]
                )


                significant = bool(
                    record[
                        significance_column
                    ]
                )


                set_word_pvalue_cell(
                    cells[
                        index
                    ],
                    p_text=
                        value,
                    significant=
                        significant,
                    font_size=
                        font_size,
                )


                continue


            # -----------------------------------------------------------------
            # Text columns
            # -----------------------------------------------------------------

            if column in [

                "Subject",
                "Model",
                "Model_1",
                "Model_2",
                "p_Display",
                "p_raw_Display",
                "p_Holm_Display",
                "Direction",

            ]:


                display_value = str(
                    value
                )


            else:


                decimals = (
                    decimals_map.get(
                        column,
                        3,
                    )
                )


                display_value = (
                    word_value(
                        value,
                        decimals=
                            decimals,
                    )
                )


            set_word_cell(
                cells[
                    index
                ],
                display_value,
                font_size=
                    font_size,
            )


    return table


# =============================================================================
# STEP 35. Create Word output
#
# NO Sig. columns.
# Stars directly after p-value.
# =============================================================================

def create_word_output(
    output_path,
    shapiro_table,
    wilcoxon_table,
    descriptive_raw,
    levene_table,
    friedman_table,
    mean_rank_table,
):

    document = Document()


    # -------------------------------------------------------------------------
    # Landscape
    # -------------------------------------------------------------------------

    section = (
        document.sections[
            0
        ]
    )


    section.orientation = (
        WD_ORIENT.LANDSCAPE
    )


    section.page_width = Inches(
        11
    )


    section.page_height = Inches(
        8.5
    )


    section.top_margin = Inches(
        0.5
    )


    section.bottom_margin = Inches(
        0.5
    )


    section.left_margin = Inches(
        0.5
    )


    section.right_margin = Inches(
        0.5
    )


    # -------------------------------------------------------------------------
    # Normal style
    # -------------------------------------------------------------------------

    normal_style = (
        document.styles[
            "Normal"
        ]
    )


    normal_style.font.name = (
        "Times New Roman"
    )


    normal_style.font.size = Pt(
        9
    )


    # -------------------------------------------------------------------------
    # Title
    # -------------------------------------------------------------------------

    title = (
        document.add_paragraph()
    )


    title.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )


    title_run = title.add_run(

        "Paired / Blocked Statistical Analysis\n"
        "Taiwan GSAT (2015–2024) and AST Mathematics A"
    )


    title_run.bold = True


    title_run.font.name = (
        "Times New Roman"
    )


    title_run.font.size = Pt(
        14
    )


    # -------------------------------------------------------------------------
    # Subtitle
    # -------------------------------------------------------------------------

    subtitle = (
        document.add_paragraph()
    )


    subtitle.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )


    subtitle_run = subtitle.add_run(

        "Matched block = examination year (2015–2024); "
        "repeated factor = model. "
        "Friedman and pairwise Wilcoxon analyses use "
        "annual normalized objective-choice accuracy (%)."
    )


    subtitle_run.font.name = (
        "Times New Roman"
    )


    subtitle_run.font.size = Pt(
        9
    )


    # =========================================================================
    # TABLE 2
    #
    # p* directly in p column
    # NO Sig. column
    # =========================================================================

    add_word_caption(
        document,
        "Table 2. Shapiro–Wilk Normality of Annual Raw Scores by Subject × Model.",
    )


    table2_display = (
        shapiro_table[
            [
                "Subject",
                "Model",
                "W",
                "p_Display",
                "Significant_p_lt_0_05",
            ]
        ]
        .copy()
    )


    add_dataframe_to_word(
        document=
            document,

        dataframe=
            table2_display,

        display_columns=[
            "Subject",
            "Model",
            "W",
            "p_Display",
        ],

        column_labels=[
            "Subject",
            "Model",
            "W",
            "p",
        ],

        decimals_map={
            "W": 3,
        },

        font_size=8,

        star_rules={
            "p_Display":
                "Significant_p_lt_0_05",
        },
    )


    add_word_note(

        document,

        "Shapiro–Wilk tests were calculated on the 10 annual raw-score values "
        "for each subject × model series. This table is retained as a legacy "
        "diagnostic for reproducibility; normality is not an assumption of the "
        "Friedman matched-samples analysis. "
        "* p < 0.05.",
    )


    document.add_page_break()


    # =========================================================================
    # TABLE 3
    #
    # Holm p* directly in Holm p column.
    #
    # Raw p has NO star.
    #
    # NO Sig. column.
    # =========================================================================

    add_word_caption(
        document,
        "Table 3. Pairwise Wilcoxon Signed-Rank Comparisons with Holm Adjustment, by Subject.",
    )


    table3_display = (
        wilcoxon_table[
            [
                "Subject",
                "Model_1",
                "Model_2",
                "Mean_Difference_pp",
                "Wilcoxon_W",
                "p_raw_Display",
                "p_Holm_Display",
                "Interpret_as_Significant",
            ]
        ]
        .copy()
    )


    add_dataframe_to_word(
        document=
            document,

        dataframe=
            table3_display,

        display_columns=[
            "Subject",
            "Model_1",
            "Model_2",
            "Mean_Difference_pp",
            "Wilcoxon_W",
            "p_raw_Display",
            "p_Holm_Display",
        ],

        column_labels=[
            "Subject",
            "Model 1",
            "Model 2",
            "Mean diff. (pp)",
            "W",
            "Raw p",
            "Holm p",
        ],

        decimals_map={
            "Mean_Difference_pp":
                2,

            "Wilcoxon_W":
                1,
        },

        font_size=7.0,

        star_rules={
            "p_Holm_Display":
                "Interpret_as_Significant",
        },
    )


    add_word_note(

        document,

        "Two-sided Wilcoxon signed-rank tests were applied to the 10 paired "
        "annual normalized objective-choice accuracy values for each model pair. "
        "Holm adjustment was applied across the 10 pairwise comparisons within "
        "each subject. Mean diff. (pp) = Model 1 minus Model 2 in percentage "
        "points. * Holm-adjusted p < 0.05 following a significant subject-level "
        "Friedman omnibus test.",
    )


    document.add_page_break()


    # =========================================================================
    # TABLE 4
    #
    # No significance marker.
    # =========================================================================

    add_word_caption(
        document,
        "Table 4. Descriptive Statistics of Annual Total Scores Across GSAT and AST Subjects (Five Models).",
    )


    table4_display = (
        descriptive_raw[
            [
                "Subject",
                "Model",
                "Min",
                "Max",
                "Mean",
                "SD",
                "Median",
                "N",
            ]
        ]
        .copy()
    )


    add_dataframe_to_word(
        document=
            document,

        dataframe=
            table4_display,

        display_columns=[
            "Subject",
            "Model",
            "Min",
            "Max",
            "Mean",
            "SD",
            "Median",
            "N",
        ],

        column_labels=[
            "Subject",
            "Model",
            "Min",
            "Max",
            "Mean",
            "SD",
            "Median",
            "N",
        ],

        decimals_map={
            "Min":
                2,

            "Max":
                2,

            "Mean":
                2,

            "SD":
                2,

            "Median":
                2,

            "N":
                0,
        },

        font_size=7.5,
    )


    add_word_note(

        document,

        "Values are annual total raw scores. Descriptive statistics are "
        "calculated across the 10 annual examinations from 2015 through 2024. "
        "SD = standard deviation across annual totals.",
    )


    document.add_page_break()


    # =========================================================================
    # TABLE 5
    #
    # p* directly in p column
    # NO Sig. column
    # =========================================================================

    add_word_caption(
        document,
        "Table 5. Levene's Test of Homogeneity of Variances Across the Five Models, by Subject.",
    )


    table5_display = (
        levene_table[
            [
                "Subject",
                "Levene_Statistic",
                "df1",
                "df2",
                "p_Display",
                "Significant_p_lt_0_05",
            ]
        ]
        .copy()
    )


    add_dataframe_to_word(
        document=
            document,

        dataframe=
            table5_display,

        display_columns=[
            "Subject",
            "Levene_Statistic",
            "df1",
            "df2",
            "p_Display",
        ],

        column_labels=[
            "Subject",
            "Levene statistic",
            "df1",
            "df2",
            "p",
        ],

        decimals_map={
            "Levene_Statistic":
                3,

            "df1":
                0,

            "df2":
                0,
        },

        font_size=8,

        star_rules={
            "p_Display":
                "Significant_p_lt_0_05",
        },
    )


    add_word_note(

        document,

        "Levene's test is reproduced from the annual raw-score data for "
        "transparency and audit purposes. Homogeneity of variance is not "
        "an assumption of the Friedman matched-samples analysis. "
        "* p < 0.05.",
    )


    document.add_page_break()


    # =========================================================================
    # TABLE 6
    #
    # Friedman p* directly in p column
    # NO Sig. column
    # =========================================================================

    add_word_caption(
        document,
        "Table 6. Friedman Test for Paired Annual Model Performance, by Subject.",
    )


    table6_display = (
        friedman_table[
            [
                "Subject",
                "N_Years",
                "Friedman_ChiSquare",
                "df",
                "p_Display",
                "Kendalls_W",
                "Omnibus_Significant",
            ]
        ]
        .copy()
    )


    add_dataframe_to_word(
        document=
            document,

        dataframe=
            table6_display,

        display_columns=[
            "Subject",
            "N_Years",
            "Friedman_ChiSquare",
            "df",
            "p_Display",
            "Kendalls_W",
        ],

        column_labels=[
            "Subject",
            "N years",
            "Friedman χ²",
            "df",
            "p",
            "Kendall's W",
        ],

        decimals_map={
            "N_Years":
                0,

            "Friedman_ChiSquare":
                3,

            "df":
                0,

            "Kendalls_W":
                3,
        },

        font_size=8,

        star_rules={
            "p_Display":
                "Omnibus_Significant",
        },
    )


    add_word_note(

        document,

        "The Friedman test treats examination year as the matched block and "
        "model as the repeated factor. Each subject contains 10 annual blocks "
        "(2015–2024) and five models. The test uses annual normalized "
        "objective-choice accuracy (%). Kendall's W is reported as the "
        "effect-size measure. * p < 0.05.",
    )


    document.add_page_break()


    # =========================================================================
    # TABLE 7
    #
    # No p-value / no star.
    # =========================================================================

    add_word_caption(
        document,
        "Table 7. Mean Ranks from the Friedman Analysis, by Subject and Model.",
    )


    table7_display = (
        mean_rank_table[
            [
                "Subject",
                "Model",
                "Mean_Rank",
                "Median_Rank",
                "Mean_Accuracy_Percent",
            ]
        ]
        .copy()
    )


    add_dataframe_to_word(
        document=
            document,

        dataframe=
            table7_display,

        display_columns=[
            "Subject",
            "Model",
            "Mean_Rank",
            "Median_Rank",
            "Mean_Accuracy_Percent",
        ],

        column_labels=[
            "Subject",
            "Model",
            "Mean rank",
            "Median rank",
            "Mean accuracy (%)",
        ],

        decimals_map={
            "Mean_Rank":
                2,

            "Median_Rank":
                2,

            "Mean_Accuracy_Percent":
                2,
        },

        font_size=8,
    )


    add_word_note(

        document,

        "Ranks were calculated within each examination year and then averaged "
        "across the 10 matched annual blocks. Rank 1 indicates the lowest "
        "annual model performance and rank 5 the highest. Tied values receive "
        "average ranks.",
    )


    # -------------------------------------------------------------------------
    # Save Word
    # -------------------------------------------------------------------------

    document.save(
        output_path
    )


# =============================================================================
# STEP 36. START
# =============================================================================

print(
    "\n"
    + "=" * 100
)


print(
    "Paired / Blocked Statistical Analysis"
)


print(
    "Paired / Blocked Statistical Analysis"
)


print(
    "=" * 100
)


print(
    "\n版本：零手填分析結果"
)


print(
    "Word：p-value 後直接加 superscript *"
)


print(
    "Word：不再產生 Sig. 欄位"
)


print(
    "\nPrimary analysis："
)


print(
    "  Friedman test"
)


print(
    "  Kendall's W"
)


print(
    "  Paired Wilcoxon signed-rank"
)


print(
    "  Holm correction"
)


print(
    "\nMatched block：Year (2015–2024)"
)


print(
    "Repeated factor：Model (5 models)"
)


print(
    "\nInferential outcome：Annual objective-choice accuracy (%)"
)


print(
    "\nOutput folder："
)


print(
    OUTPUT_DIR
)


# =============================================================================
# STEP 37. Check six input files
# =============================================================================

missing_files = [

    str(
        path
    )

    for path
    in INPUT_FILES.values()

    if not path.exists()
]


if missing_files:


    raise FileNotFoundError(

        "\n以下輸入檔案不存在：\n\n"

        +

        "\n".join(
            missing_files
        )
    )


print(
    "\nPASS：六個輸入檔案皆存在。"
)


# =============================================================================
# STEP 38. Read six subjects
# =============================================================================

annual_raw_list = []

annual_accuracy_list = []

candidate_qc_list = []

selected_source_rows = []


for (
    subject,
    file_path,
) in INPUT_FILES.items():


    result = (
        read_subject_file(
            subject=
                subject,
            file_path=
                file_path,
        )
    )


    # -------------------------------------------------------------------------
    # Raw
    # -------------------------------------------------------------------------

    annual_raw = (
        result[
            "annual_raw"
        ]
        .copy()
    )


    annual_raw.insert(
        0,
        "Subject",
        subject,
    )


    annual_raw_list.append(
        annual_raw
    )


    # -------------------------------------------------------------------------
    # Normalized accuracy
    # -------------------------------------------------------------------------

    annual_accuracy = (
        raw_to_accuracy(
            subject=
                subject,
            annual_raw=
                result[
                    "annual_raw"
                ],
        )
    )


    annual_accuracy.insert(
        0,
        "Subject",
        subject,
    )


    annual_accuracy_list.append(
        annual_accuracy
    )


    # -------------------------------------------------------------------------
    # Candidate QC
    # -------------------------------------------------------------------------

    candidate_qc_list.append(
        result[
            "candidate_qc"
        ]
    )


    # -------------------------------------------------------------------------
    # Selected source
    # -------------------------------------------------------------------------

    selected_source_rows.append({

        "Subject":
            subject,

        "File":
            (
                "supplementary-materials/"
                "analysis-ready-data/"
                f"{file_path.name}"
            ),

        "Selected_Sheet":
            result[
                "selected_sheet"
            ],

        "Extraction_Method":
            result[
                "selected_method"
            ],
    })


# =============================================================================
# STEP 39. Combine six subjects
# =============================================================================

annual_raw_all = pd.concat(
    annual_raw_list,
    ignore_index=True,
)


annual_accuracy_all = pd.concat(
    annual_accuracy_list,
    ignore_index=True,
)


candidate_qc_all = pd.concat(
    candidate_qc_list,
    ignore_index=True,
)


selected_sources = pd.DataFrame(
    selected_source_rows
)


# =============================================================================
# STEP 40. Scoring-scale table
# =============================================================================

scoring_scale_rows = []


for subject in INPUT_FILES.keys():


    for year in YEARS:


        scoring_scale_rows.append({

            "Subject":
                subject,

            "Year":
                year,

            "Objective_Max":
                OBJECTIVE_MAX[
                    subject
                ][
                    year
                ],
        })


scoring_scale = pd.DataFrame(
    scoring_scale_rows
)


# =============================================================================
# STEP 41. Merge Objective_Max
# =============================================================================

annual_raw_all = (
    annual_raw_all
    .merge(
        scoring_scale,
        on=[
            "Subject",
            "Year",
        ],
        how="left",
        validate="one_to_one",
    )
)


annual_accuracy_all = (
    annual_accuracy_all
    .merge(
        scoring_scale,
        on=[
            "Subject",
            "Year",
        ],
        how="left",
        validate="one_to_one",
    )
)


# =============================================================================
# STEP 42. Global QC
# =============================================================================

print(
    "\n"
    + "=" * 100
)


print(
    "GLOBAL DATA QC"
)


print(
    "=" * 100
)


# -------------------------------------------------------------------------
# Six subjects
# -------------------------------------------------------------------------

n_subjects = (
    annual_raw_all[
        "Subject"
    ]
    .nunique()
)


if n_subjects != 6:


    raise ValueError(

        f"科目數不是 6，"
        f"目前={n_subjects}"
    )


# -------------------------------------------------------------------------
# Ten years per subject
# -------------------------------------------------------------------------

year_counts = (
    annual_raw_all
    .groupby(
        "Subject"
    )[
        "Year"
    ]
    .nunique()
)


if not (
    year_counts
    == 10
).all():


    raise ValueError(

        "\n某科並非完整 10 年：\n"

        +

        year_counts.to_string()
    )


# -------------------------------------------------------------------------
# Duplicate Subject + Year
# -------------------------------------------------------------------------

duplicate_subject_year = (
    annual_raw_all
    .duplicated(
        subset=[
            "Subject",
            "Year",
        ],
        keep=False,
    )
)


if duplicate_subject_year.any():


    raise ValueError(

        "\n存在 duplicate Subject + Year：\n"

        +

        annual_raw_all.loc[
            duplicate_subject_year
        ].to_string(
            index=False
        )
    )


# -------------------------------------------------------------------------
# Missing raw
# -------------------------------------------------------------------------

missing_raw = int(
    annual_raw_all[
        MODELS
    ]
    .isna()
    .sum()
    .sum()
)


if missing_raw != 0:


    raise ValueError(

        f"Raw annual score 有 "
        f"{missing_raw} 個缺值。"
    )


# -------------------------------------------------------------------------
# Missing accuracy
# -------------------------------------------------------------------------

missing_accuracy = int(
    annual_accuracy_all[
        MODELS
    ]
    .isna()
    .sum()
    .sum()
)


if missing_accuracy != 0:


    raise ValueError(

        f"Normalized accuracy 有 "
        f"{missing_accuracy} 個缺值。"
    )


# -------------------------------------------------------------------------
# Accuracy range
# -------------------------------------------------------------------------

accuracy_values = (
    annual_accuracy_all[
        MODELS
    ]
    .to_numpy(
        dtype=float
    )
)


if (
    np.nanmin(
        accuracy_values
    )
    < -1e-8
):


    raise ValueError(
        "Normalized accuracy 出現 <0 的值。"
    )


if (
    np.nanmax(
        accuracy_values
    )
    > 100.0000001
):


    raise ValueError(
        "Normalized accuracy 出現 >100 的值。"
    )


print(
    "PASS：6 subjects"
)


print(
    "PASS：每科 2015–2024 共 10 年"
)


print(
    "PASS：Subject + Year 無 duplicate"
)


print(
    "PASS：RAW annual score 無缺值"
)


print(
    "PASS：normalized accuracy 無缺值"
)


print(
    "PASS：normalized accuracy 全部介於 0–100%"
)


# =============================================================================
# STEP 43. Long-format normalized accuracy
# =============================================================================

annual_accuracy_long = (
    annual_accuracy_all[
        [
            "Subject",
            "Year",
            "Objective_Max",
            *MODELS,
        ]
    ]
    .melt(
        id_vars=[
            "Subject",
            "Year",
            "Objective_Max",
        ],
        value_vars=
            MODELS,
        var_name=
            "Model",
        value_name=
            "Objective_Choice_Accuracy_Percent",
    )
)


# =============================================================================
# STEP 44. Run analyses
# =============================================================================

table2_shapiro_legacy = (
    run_shapiro_legacy(
        annual_raw_all
    )
)


table4_descriptive_raw = (
    run_descriptive_raw(
        annual_raw_all
    )
)


table5_levene_legacy = (
    run_levene_legacy(
        annual_raw_all
    )
)


table6_friedman = (
    run_friedman(
        annual_accuracy_all
    )
)


table7_mean_ranks = (
    run_friedman_mean_ranks(
        annual_accuracy_all
    )
)


table3_wilcoxon = (
    run_wilcoxon_holm(
        annual_accuracy_all=
            annual_accuracy_all,
        friedman_results=
            table6_friedman,
    )
)


descriptive_accuracy = (
    run_descriptive_accuracy(
        annual_accuracy_all
    )
)


# =============================================================================
# STEP 45. Significant pairs only
# =============================================================================

significant_pairs = (
    table3_wilcoxon.loc[
        table3_wilcoxon[
            "Interpret_as_Significant"
        ]
    ]
    .copy()
    .reset_index(
        drop=True
    )
)


# =============================================================================
# STEP 46. Metadata
# =============================================================================

metadata_rows = [


    [
        "Script",
        "paired_blocked_statistical_analysis.py",
    ],


    [
        "Purpose",
        "Paired / blocked statistical analysis",
    ],


    [
        "Results hard-coded",
        "No",
    ],


    [
        "Matched block",
        "Year (2015–2024)",
    ],


    [
        "N matched blocks",
        "10 per subject",
    ],


    [
        "Repeated factor",
        "Model (5 models)",
    ],


    [
        "Primary inferential outcome",
        "Annual normalized objective-choice accuracy (%)",
    ],


    [
        "Normalization",
        "annual raw score / objectively scoreable annual maximum × 100",
    ],


    [
        "Primary omnibus test",
        "Friedman test",
    ],


    [
        "Effect size",
        "Kendall's W = Friedman chi-square / [n × (k − 1)]",
    ],


    [
        "Post-hoc",
        "Two-sided paired Wilcoxon signed-rank test",
    ],


    [
        "Multiplicity correction",
        "Holm correction across 10 pairwise comparisons within each subject",
    ],


    [
        "Word significance format",
        "Superscript * directly after the relevant p-value; no separate Sig. column",
    ],


    [
        "Alpha",
        "0.05, two-sided",
    ],


    [
        "Table 4 scale",
        "Annual raw total scores",
    ],


    [
        "Table 2",
        "Legacy Shapiro-Wilk diagnostic; not an assumption of Friedman",
    ],


    [
        "Table 5",
        "Legacy Levene diagnostic; not an assumption of Friedman",
    ],


    [
        "Generated",
        datetime.now().strftime(
            "%Y-%m-%d %H:%M:%S"
        ),
    ],


    [
        "Python",
        sys.version.replace(
            "\n",
            " "
        ),
    ],


    [
        "Platform",
        platform.platform(),
    ],


    [
        "pandas",
        pd.__version__,
    ],


    [
        "numpy",
        np.__version__,
    ],
]


try:

    import scipy


    metadata_rows.append(
        [
            "scipy",
            scipy.__version__,
        ]
    )


except Exception:

    pass


try:

    import statsmodels


    metadata_rows.append(
        [
            "statsmodels",
            statsmodels.__version__,
        ]
    )


except Exception:

    pass


metadata = pd.DataFrame(
    metadata_rows,
    columns=[
        "Item",
        "Value",
    ],
)


# =============================================================================
# STEP 47. Export CSV
# =============================================================================

annual_accuracy_long.to_csv(
    OUTPUT_ACCURACY_CSV,
    index=False,
    encoding="utf-8-sig",
)


table6_friedman.to_csv(
    OUTPUT_FRIEDMAN_CSV,
    index=False,
    encoding="utf-8-sig",
)


table3_wilcoxon.to_csv(
    OUTPUT_WILCOXON_CSV,
    index=False,
    encoding="utf-8-sig",
)


# =============================================================================
# STEP 48. Export Excel
# =============================================================================

with pd.ExcelWriter(
    OUTPUT_EXCEL,
    engine="openpyxl",
) as writer:


    metadata.to_excel(
        writer,
        sheet_name=
            "00_README",
        index=False,
    )


    scoring_scale.to_excel(
        writer,
        sheet_name=
            "01_ScoringScale",
        index=False,
    )


    selected_sources.to_excel(
        writer,
        sheet_name=
            "02_SelectedSources",
        index=False,
    )


    candidate_qc_all.to_excel(
        writer,
        sheet_name=
            "03_Candidate_QC",
        index=False,
    )


    annual_raw_all.to_excel(
        writer,
        sheet_name=
            "04_AnnualRaw_Wide",
        index=False,
    )


    annual_accuracy_all.to_excel(
        writer,
        sheet_name=
            "05_AnnualAccuracy_Wide",
        index=False,
    )


    annual_accuracy_long.to_excel(
        writer,
        sheet_name=
            "06_AnnualAccuracy_Long",
        index=False,
    )


    table2_shapiro_legacy.to_excel(
        writer,
        sheet_name=
            "07_Table2_Shapiro_Legacy",
        index=False,
    )


    table3_wilcoxon.to_excel(
        writer,
        sheet_name=
            "08_Table3_Wilcoxon_Holm",
        index=False,
    )


    table4_descriptive_raw.to_excel(
        writer,
        sheet_name=
            "09_Table4_Descriptive_Raw",
        index=False,
    )


    table5_levene_legacy.to_excel(
        writer,
        sheet_name=
            "10_Table5_Levene_Legacy",
        index=False,
    )


    table6_friedman.to_excel(
        writer,
        sheet_name=
            "11_Table6_Friedman",
        index=False,
    )


    table7_mean_ranks.to_excel(
        writer,
        sheet_name=
            "12_Table7_MeanRanks",
        index=False,
    )


    significant_pairs.to_excel(
        writer,
        sheet_name=
            "13_SignificantPairs",
        index=False,
    )


    descriptive_accuracy.to_excel(
        writer,
        sheet_name=
            "14_Descriptive_Accuracy",
        index=False,
    )


# =============================================================================
# STEP 49. Format Excel
# =============================================================================

format_excel_workbook(
    OUTPUT_EXCEL
)


# =============================================================================
# STEP 50. Export Word
# =============================================================================

create_word_output(
    output_path=
        OUTPUT_WORD,

    shapiro_table=
        table2_shapiro_legacy,

    wilcoxon_table=
        table3_wilcoxon,

    descriptive_raw=
        table4_descriptive_raw,

    levene_table=
        table5_levene_legacy,

    friedman_table=
        table6_friedman,

    mean_rank_table=
        table7_mean_ranks,
)


# =============================================================================
# STEP 51. Create log
# =============================================================================

log_lines = []


log_lines.append(
    "=" * 100
)


log_lines.append(
    "Paired / Blocked Statistical Analysis"
)


log_lines.append(
    "Paired / Blocked Statistical Analysis"
)


log_lines.append(
    "=" * 100
)


log_lines.append("")


log_lines.append(
    "ZERO HARD-CODED ANALYSIS RESULTS"
)


log_lines.append(
    "All descriptive and inferential results were calculated from the input Excel files."
)


log_lines.append("")


log_lines.append(
    "WORD SIGNIFICANCE FORMAT"
)


log_lines.append(
    "No separate Sig. column."
)


log_lines.append(
    "Significant results are marked by a superscript * directly after the relevant p-value."
)


log_lines.append("")


log_lines.append(
    "ANALYSIS DESIGN"
)


log_lines.append(
    "Block = Year (2015–2024; n=10)"
)


log_lines.append(
    "Repeated factor = Model (5 models)"
)


log_lines.append(
    "Outcome = annual normalized objective-choice accuracy (%)"
)


log_lines.append(
    "Omnibus = Friedman test"
)


log_lines.append(
    "Effect size = Kendall's W"
)


log_lines.append(
    "Post-hoc = paired Wilcoxon signed-rank test"
)


log_lines.append(
    "Multiplicity = Holm correction within each subject"
)


log_lines.append("")


# -----------------------------------------------------------------------------
# Selected sources
# -----------------------------------------------------------------------------

log_lines.append(
    "=" * 100
)


log_lines.append(
    "SELECTED INPUT SOURCES"
)


log_lines.append(
    "=" * 100
)


log_lines.append(
    selected_sources.to_string(
        index=False
    )
)


log_lines.append("")


# -----------------------------------------------------------------------------
# Friedman
# -----------------------------------------------------------------------------

log_lines.append(
    "=" * 100
)


log_lines.append(
    "TABLE 6 – FRIEDMAN RESULTS"
)


log_lines.append(
    "=" * 100
)


log_lines.append(
    table6_friedman[
        [
            "Subject",
            "N_Years",
            "Friedman_ChiSquare",
            "df",
            "p_Display",
            "Kendalls_W",
        ]
    ].to_string(
        index=False
    )
)


log_lines.append("")


# -----------------------------------------------------------------------------
# Mean ranks
# -----------------------------------------------------------------------------

log_lines.append(
    "=" * 100
)


log_lines.append(
    "TABLE 7 – FRIEDMAN MEAN RANKS"
)


log_lines.append(
    "=" * 100
)


log_lines.append(
    table7_mean_ranks[
        [
            "Subject",
            "Model",
            "Mean_Rank",
            "Median_Rank",
            "Mean_Accuracy_Percent",
        ]
    ].to_string(
        index=False
    )
)


log_lines.append("")


# -----------------------------------------------------------------------------
# Significant pairwise
# -----------------------------------------------------------------------------

log_lines.append(
    "=" * 100
)


log_lines.append(
    "HOLM-SIGNIFICANT WILCOXON PAIRS"
)


log_lines.append(
    "=" * 100
)


if significant_pairs.empty:


    log_lines.append(
        "No Holm-adjusted pairwise comparison was significant."
    )


else:


    log_lines.append(
        significant_pairs[
            [
                "Subject",
                "Model_1",
                "Model_2",
                "Mean_Difference_pp",
                "Median_Difference_pp",
                "Model_1_Wins",
                "Model_2_Wins",
                "Ties",
                "Wilcoxon_W",
                "p_raw_Display",
                "p_Holm_Display",
                "Direction",
            ]
        ].to_string(
            index=False
        )
    )


log_lines.append("")


log_lines.append(
    "=" * 100
)


log_lines.append(
    "OUTPUT FILES"
)


log_lines.append(
    "=" * 100
)


log_lines.append(
    f"Excel: output/{OUTPUT_EXCEL.name}"
)


log_lines.append(
    f"Word: output/{OUTPUT_WORD.name}"
)


log_lines.append(
    f"Annual Accuracy CSV: output/{OUTPUT_ACCURACY_CSV.name}"
)


log_lines.append(
    f"Friedman CSV: output/{OUTPUT_FRIEDMAN_CSV.name}"
)


log_lines.append(
    f"Wilcoxon-Holm CSV: output/{OUTPUT_WILCOXON_CSV.name}"
)


with open(
    OUTPUT_LOG,
    "w",
    encoding="utf-8",
) as file:


    file.write(
        "\n".join(
            log_lines
        )
    )


# =============================================================================
# STEP 52. Console - Friedman
# =============================================================================

print(
    "\n"
    + "=" * 100
)


print(
    "TABLE 6 – FRIEDMAN TEST"
)


print(
    "=" * 100
)


print(
    table6_friedman[
        [
            "Subject",
            "N_Years",
            "Friedman_ChiSquare",
            "df",
            "p_Display",
            "Kendalls_W",
        ]
    ].to_string(
        index=False
    )
)


# =============================================================================
# STEP 53. Console - mean ranks
# =============================================================================

print(
    "\n"
    + "=" * 100
)


print(
    "TABLE 7 – FRIEDMAN MEAN RANKS"
)


print(
    "=" * 100
)


print(
    table7_mean_ranks[
        [
            "Subject",
            "Model",
            "Mean_Rank",
            "Median_Rank",
            "Mean_Accuracy_Percent",
        ]
    ].to_string(
        index=False
    )
)


# =============================================================================
# STEP 54. Console - significant Wilcoxon pairs
# =============================================================================

print(
    "\n"
    + "=" * 100
)


print(
    "HOLM-SIGNIFICANT WILCOXON PAIRS"
)


print(
    "=" * 100
)


if significant_pairs.empty:


    print(
        "沒有 Holm-adjusted p < 0.05 "
        "且 Friedman omnibus significant 的 pairwise comparison。"
    )


else:


    print(
        significant_pairs[
            [
                "Subject",
                "Model_1",
                "Model_2",
                "Mean_Difference_pp",
                "Wilcoxon_W",
                "p_raw_Display",
                "p_Holm_Display",
                "Direction",
            ]
        ].to_string(
            index=False
        )
    )


# =============================================================================
# FINAL SUMMARY
# =============================================================================

print(
    "\n"
    + "=" * 100
)

print(
    "Paired / blocked statistical analysis completed."
)

print(
    "=" * 100
)

print(
    "\nOutput directory:"
)

print(
    OUTPUT_DIR
)

print(
    "\nGenerated files:"
)

for generated_file in [
    OUTPUT_EXCEL,
    OUTPUT_LOG,
    OUTPUT_ACCURACY_CSV,
    OUTPUT_FRIEDMAN_CSV,
    OUTPUT_WILCOXON_CSV,
    OUTPUT_WORD,
]:
    print(
        "  -",
        generated_file.name
    )

print(
    "\nWord significance format:"
)

print(
    "  Table 2: superscript * after p when Shapiro-Wilk p < 0.05"
)

print(
    "  Table 3: superscript * after Holm-adjusted p when significant "
    "following a significant Friedman omnibus test"
)

print(
    "  Table 5: superscript * after p when Levene p < 0.05"
)

print(
    "  Table 6: superscript * after p when Friedman p < 0.05"
)

print(
    "  Table 7: no significance marker"
)

print(
    "\nDone."
)
